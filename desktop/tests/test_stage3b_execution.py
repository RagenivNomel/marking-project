"""Stage 3B tests at the application-to-production orchestration seam."""
import csv
from dataclasses import replace
import gc
import json
from pathlib import Path
import shutil
import threading
import time
import unittest
import uuid
import sys
from unittest.mock import patch

from openpyxl import Workbook, load_workbook
from pypdf import PdfReader, PdfWriter
from grading.schemas import CRITERIA, Identity, ROOT, ValidationError, Validator

DESKTOP_DEPS = ROOT / ".desktop-deps"
if DESKTOP_DEPS.is_dir():
    sys.path.insert(0, str(DESKTOP_DEPS))
from PySide6.QtCore import QCoreApplication, QEventLoop, QTimer

from application import Action, AssignmentSource, WorkflowController
from application.workflows.task_storage import task_output_paths
from desktop.projection import project
from desktop.teacher_flow import derive_teacher_flow
from desktop.bridge import DesktopBridge
from grading.sol_grader import CredentialUnavailable, SolGrader
from tests.test_real_batch_integration import RecordingTransport, working_state
from workflow.calibration_pipeline import CalibrationPipeline
from workflow.pipeline import Pipeline, fingerprint, job_key
from excel.schema import AUDIT_HEADERS, AUDIT_SHEET, HEADERS, SHEET
from excel.workbook import ExcelStore


def _valid_response(identity):
    return {
        "reading_quality": {
            "complete_essay_legible": True,
            "uncertainties": [],
            "materially_affects_grade": False,
        },
        "grading": {
            **identity,
            "criteria": {
                name: {"rating": "可以更进一步", "short_comment": f"围绕{name}补充具体细节。"}
                for name in CRITERIA
            },
            "teacher_comment": "The composition is complete; add more specific details to the central paragraph.",
        },
        "evidence": [{"judgment": "Clear focus", "page_numbers": [1], "rationale": "The main action is concentrated in the middle section."}],
    }


class Stage3BApplicationExecutionTests(unittest.TestCase):
    def setUp(self):
        base = ROOT / "test-workspace-temp"
        base.mkdir(parents=True, exist_ok=True)
        self.root = base / f"stage3b-execution-{uuid.uuid4().hex}"
        self.root.mkdir()
        self.addCleanup(self._cleanup_temp)
        self.pile = self.root / "pile"
        self.pile.mkdir()
        self.identities = [
            {"class_name": "207", "student_id": "01", "student_name": "Alice Chen"},
            {"class_name": "207", "student_id": "02", "student_name": "Ben Lim"},
        ]
        self._make_split()
        self.roster = self.root / "roster.xlsx"
        self._make_roster()
        self.decisions = self.root / "identity_decisions.json"
        self._make_decisions()
        self.source = AssignmentSource(
            split_pile=self.pile,
            roster=self.roster,
            identity_decisions=self.decisions,
            config_dir=ROOT / "config",
        )
        self.source_hashes = {path: path.read_bytes() for path in self.pile.glob("*.pdf")}
        self.transport = RecordingTransport([Identity.from_dict(item) for item in self.identities])
        self.grader_instances = []
        self.controller = WorkflowController(
            project_dir=self.root,
            pipeline_factory=self._pipeline_factory,
        )

    def _cleanup_temp(self):
        gc.collect()
        for _ in range(8):
            try:
                shutil.rmtree(self.root)
                return
            except PermissionError:
                time.sleep(0.05)
                gc.collect()
        shutil.rmtree(self.root, ignore_errors=True)

    def _make_split(self):
        continuous = PdfWriter()
        continuous.add_blank_page(width=612, height=792)
        second = continuous.add_blank_page(width=612, height=792)
        second.rotate(90)
        with (self.pile / "_continuous.pdf").open("wb") as handle:
            continuous.write(handle)
        pages = PdfReader(self.pile / "_continuous.pdf").pages
        rows = []
        for index, page in enumerate(pages, 1):
            path = self.pile / f"essay_{index:03}.pdf"
            writer = PdfWriter()
            writer.add_page(page)
            with path.open("wb") as handle:
                writer.write(handle)
            rows.append({
                "start_page_idx": index - 1,
                "pages": 1,
                "raw_ocr_name": self.identities[index - 1]["student_name"],
                "name_preview": "",
            })
        with (self.pile / "_manifest.csv").open("w", newline="", encoding="utf-8") as handle:
            writer = csv.DictWriter(handle, fieldnames=("start_page_idx", "pages", "raw_ocr_name", "name_preview"))
            writer.writeheader()
            writer.writerows(rows)

    def _make_roster(self):
        book = Workbook()
        sheet = book.active
        sheet.title = "作文诊断输入"
        sheet.append(["班号", "学生姓名", "班级"])
        for identity in self.identities:
            sheet.append([identity["student_id"], identity["student_name"], identity["class_name"]])
        book.save(self.roster)
        book.close()

    def _make_decisions(self):
        decisions = []
        for index, identity in enumerate(self.identities, 1):
            decisions.append({
                "source_pdf": f"essay_{index:03}.pdf",
                "class_name": identity["class_name"],
                "student_id": identity["student_id"],
                "match_status": "STRONG_ROSTER_MATCH",
                "evidence": "Teacher confirmed for Stage 3B test.",
            })
        self.decisions.write_text(json.dumps(decisions, ensure_ascii=False), encoding="utf-8")

    def _pipeline_factory(self, project_dir, batch_id, config_dir, *, workbook_path=None):
        def calibration_factory(project_dir, model_profile=None):
            profile = json.loads((Path(config_dir) / "calibration_model.json").read_text(encoding="utf-8"))["profiles"][model_profile]
            return CalibrationPipeline(
                project_dir,
                config_dir=config_dir,
                model_profile=model_profile,
                grader_factory=lambda: self._new_grader(profile),
            )
        return Pipeline(
            project_dir,
            batch_id,
            config_dir=config_dir,
            allow_real=True,
            real_pipeline_factory=calibration_factory,
            workbook_path=workbook_path,
        )

    def _new_grader(self, profile):
        grader = SolGrader(
            model=profile["model"],
            reasoning_effort=profile["reasoning_effort"],
            timeout_seconds=profile["request_timeout_seconds"],
            transport=self.transport,
        )
        self.grader_instances.append(grader)
        return grader

    def _add_interrupted_retry_jobs(self, workbook, count=3):
        batch_id = self.controller.workflow.marking_batch_id(self.source)
        batch_dir = self.root / "jobs" / batch_id
        paths = []
        job_ids = []
        for index in range(count):
            identity = Identity(
                student_id=f"9{index + 1}",
                student_name=f"Retry Fixture {index + 1}",
                class_name="207",
            )
            key = job_key(identity)
            source_digest = f"{index + 1:064x}"
            parent = batch_dir / key
            child = self.root / "jobs" / f"calibration_{batch_id}_{key}"
            parent.mkdir(parents=True)
            child.mkdir(parents=True)
            checkpoint = {
                "schema_version": 2,
                "mode": "REAL_PDF_BATCH",
                "job_id": key,
                "identity": identity.to_dict(),
                "input_digest": f"{index + 11:064x}",
                "source_sha256": source_digest,
                "source_page_count": 1,
                "model_profile": "luna_xhigh",
                "execution_namespace": batch_id,
                "workbook": str(workbook),
                "state": "TRANSCRIBED",
                "history": ["NEW", "SCANNED", "IDENTIFIED", "TRANSCRIBED"],
                "last_error": None,
            }
            child_checkpoint = {
                **checkpoint,
                "mode": "CALIBRATION",
                "job_id": child.name,
                "live_request_attempts": 1,
            }
            (parent / "student_record.json").write_text(
                json.dumps(checkpoint, ensure_ascii=False), encoding="utf-8"
            )
            (child / "student_record.json").write_text(
                json.dumps(child_checkpoint, ensure_ascii=False), encoding="utf-8"
            )
            (parent / "real_grading_bridge.json").write_text(
                json.dumps({
                    "calibration_job": str(child),
                    "execution_namespace": batch_id,
                    "calibration_attempt_namespace": batch_id,
                    "status": "IN_PROGRESS",
                }),
                encoding="utf-8",
            )
            (parent / "real_grading_attempts.json").write_text(
                json.dumps({"schema_version": 1, "attempts": [{"status": "INTERRUPTED"}]}),
                encoding="utf-8",
            )
            paths.extend((parent, child))
            job_ids.append(key)
        return tuple(job_ids), tuple(paths)

    @staticmethod
    def _delete_workbook_result(workbook, job_id):
        editable = load_workbook(workbook)
        try:
            audit = editable[AUDIT_SHEET]
            audit_row = next(
                row[0].row for row in audit.iter_rows(min_row=2)
                if row[0].value == job_id
            )
            result_row = audit.cell(audit_row, 4).value
            editable[SHEET].delete_rows(result_row, 1)
            audit.delete_rows(audit_row, 1)
            editable.save(workbook)
        finally:
            editable.close()

    def test_skipped_second_essay_is_saved_aside_and_the_rest_is_marked(self):
        from desktop.identity_confirmation import SKIPPED_FOLDER, build_identity_review, save_identity_confirmations
        # Essay 2 is Alice's second essay: confirm essay 1 and skip essay 2.
        decisions = json.loads(self.decisions.read_text(encoding="utf-8"))[:1]
        self.decisions.write_text(json.dumps(decisions, ensure_ascii=False), encoding="utf-8")
        source, saved = save_identity_confirmations(self.source, [{"sourcePdf": "essay_002.pdf", "skip": True}])
        self.assertEqual(saved, 2)
        copy = self.root / SKIPPED_FOLDER / "pile-作文002.pdf"
        self.assertEqual(copy.read_bytes(), (self.pile / "essay_002.pdf").read_bytes())

        inspection = self.controller.inspect(source)
        self.assertEqual((inspection.summary["submissions"], inspection.summary["skipped"]), (1, 1))
        self.assertEqual(inspection.skipped, (str(copy),))
        self.assertFalse(build_identity_review(source, inspection)["visible"])
        flow = derive_teacher_flow(inspection)
        self.assertEqual((flow["step"], flow["primaryAction"]), ("ready_to_mark", "RUN_MARKING"))
        self.assertIn("另有1份已跳过", project(inspection, source)["subtitle"])
        self.assertEqual(project(inspection, source)["skippedEssays"],
                         [{"path": str(copy), "label": "批改跳过的作文 002"}])

        result = self.controller.execute(Action.RUN_MARKING, source)
        self.assertEqual([item["identity"]["student_id"] for item in result["result"]["validated"]], ["01"])
        self.assertEqual(len(self.transport.calls), 1)
        after = self.controller.inspect(source)
        self.assertEqual((after.summary["committed_results"], after.summary["skipped"]), (1, 1))

    def test_unconfirmed_assignments_cannot_start_marking(self):
        source = AssignmentSource(
            split_pile=self.pile,
            roster=self.roster,
            config_dir=ROOT / "config",
        )
        with self.assertRaises(ValidationError):
            self.controller.execute(Action.RUN_MARKING, source)
        self.assertEqual(self.transport.calls, [])
        batch = self.controller.workflow.marking_batch_id(self.source)
        self.assertFalse((self.root / "output" / batch / "results.xlsx").exists())

    def test_fresh_and_damaged_completed_assignments_have_distinct_states(self):
        # The expected results workbook is an output, not a prerequisite.
        batch_id = self.controller.workflow.marking_batch_id(self.source)
        fresh_source = self.controller.bind_source(self.source)
        self.assertIsNone(fresh_source.workbook)
        task_paths = task_output_paths(self.source)
        self.assertEqual(task_paths.results_directory.parent, self.pile.resolve() / "Results")
        self.assertEqual(fresh_source.results_directory, task_paths.results_directory)
        self.assertEqual(fresh_source.feedback_cards_directory, task_paths.feedback_cards_directory)
        self.assertFalse(fresh_source.results_directory.exists())
        fresh = self.controller.inspect(fresh_source)
        fresh_flow = derive_teacher_flow(project(fresh, fresh_source, language="en"), "en")
        self.assertEqual(fresh.summary["marking_available"], 2)
        self.assertEqual(fresh.summary["grading_results_available"], 0)
        self.assertNotIn("WORKBOOK_UNREADABLE", [a.code for a in fresh.attention])
        self.assertEqual(fresh_flow["step"], "ready_to_mark")
        self.assertTrue(fresh_flow["primaryActionEnabled"])
        self.assertEqual(self.transport.calls, [])

        # A local stub at the external transport boundary creates real
        # production checkpoints and Excel rows without contacting a model.
        outcome = self.controller.execute(Action.RUN_MARKING, self.source)
        self.assertEqual(len(self.transport.calls), 2)
        workbook = outcome["source"].workbook
        self.assertTrue(workbook.is_file())
        self.assertEqual(workbook, task_paths.workbook)
        self.assertEqual(list(self.root.rglob("results.xlsx")), [workbook])
        self.assertFalse((self.root / "output" / batch_id / "results.xlsx").exists())
        complete = self.controller.inspect(self.source)
        complete_source = self.controller.bind_source(self.source)
        complete_flow = derive_teacher_flow(project(complete, complete_source, language="en"), "en")
        self.assertEqual(complete.summary["grading_results_available"], 2)
        self.assertEqual(complete.summary["awaiting_review"], 2)
        self.assertEqual(complete_flow["step"], "review")
        self.assertEqual(complete_flow["headline"], "Teacher review is ready ✓")

        workbook.write_bytes(b"not an Excel workbook")
        damaged_source = self.controller.bind_source(self.source)
        self.assertEqual(damaged_source.workbook, workbook)
        damaged = self.controller.inspect(damaged_source)
        damaged_flow = derive_teacher_flow(project(damaged, damaged_source, language="en"), "en")
        self.assertIn("WORKBOOK_UNREADABLE", [a.code for a in damaged.attention])
        self.assertGreater(damaged.summary["blocking_errors"], 0)
        self.assertEqual(damaged_flow["step"], "resolve_attention")
        self.assertNotIn("Teacher review is ready", damaged_flow["headline"])
        self.assertFalse(damaged_flow["primaryActionEnabled"])
        self.assertFalse(any(a.action == Action.RUN_MARKING for a in damaged.available_actions))

        # With the workbook gone nothing is saved, so every essay is not done.
        workbook.unlink()
        missing_source = self.controller.bind_source(self.source)
        self.assertIsNone(missing_source.workbook)
        missing = self.controller.inspect(missing_source)
        self.assertEqual(missing.summary["marking_available"], 2)
        self.assertEqual(missing.summary["committed_results"], 0)
        self.assertEqual(len(self.transport.calls), 2)

    def test_preflight_blocked_jobs_remain_visible_and_unfinished(self):
        app = QCoreApplication.instance() or QCoreApplication([])

        class UnavailableTransport:
            calls = 0

            def ensure_ready(self):
                raise CredentialUnavailable("Local model access is unavailable")

            def __call__(self, payload, timeout_seconds):
                self.calls += 1
                raise AssertionError("No model request should occur after failed preflight")

        self.transport = UnavailableTransport()
        before = self.controller.inspect(self.source)
        bridge = DesktopBridge(controller=self.controller)
        bridge.setLanguage("en")
        bridge._source = self.source
        bridge._inspection = before
        bridge._state = project(before, self.source, language="en")
        try:
            bridge.startMarking()
            deadline = time.monotonic() + 5
            while bridge.busy and time.monotonic() < deadline:
                app.processEvents(QEventLoop.AllEvents, 20)
                time.sleep(0.002)
            app.processEvents(QEventLoop.AllEvents, 20)
            state = bridge.state
            self.assertFalse(bridge.busy)
            self.assertEqual(self.transport.calls, 0)
            self.assertEqual(len(state["markingResult"]["failed"]), 2)
            self.assertEqual(state["summary"]["grading_results_available"], 0)
            self.assertEqual(state["summary"]["blocking_errors"], 0)
            self.assertEqual(state["teacherFlow"]["step"], "marking_start_failed")
            self.assertEqual(state["teacherFlow"]["headline"], "Marking could not start")
            self.assertEqual(state["teacherFlow"]["primaryActionLabel"], "Continue marking")
            self.assertTrue(state["teacherFlow"]["primaryActionEnabled"])
            self.assertIn("Your essays are still ready", state["notice"])
            self.assertNotIn("Teacher review is ready", state["notice"])
            self.assertEqual(len(state["attention"]), 1)
            self.assertIn("MARKING_STARTUP_FAILED", state["attention"][0]["details"])
            self.assertIn("Your essays are still ready", state["attention"][0]["message"])
            self.assertFalse(any("WORKBOOK_UNREADABLE" in a["details"] for a in state["attention"]))
            self.assertIsNone(bridge._source.workbook)
            batch = self.controller.workflow.marking_batch_id(self.source)
            self.assertFalse((self.root / "output" / batch / "results.xlsx").exists())

            expected = task_output_paths(self.source).workbook
            stale = replace(self.source, workbook=expected, job_roots=(self.root / "jobs" / batch,))
            rebound = self.controller.bind_source(stale)
            self.assertIsNone(rebound.workbook)
            reopened = self.controller.inspect(stale)
            self.assertEqual(reopened.summary["marking_available"], 2)
            # A failed attempt leaves nothing behind: reopened, the essays are
            # simply not done and ready to mark.
            self.assertEqual(derive_teacher_flow(project(reopened, rebound, language="en"), "en")["step"],
                             "ready_to_mark")
            for identity in self.identities:
                self.assertEqual(working_state(self.root / "jobs" / batch / job_key(Identity.from_dict(identity))), [])
        finally:
            bridge.shutdown()

    def test_cross_project_workbook_and_normal_pile_are_rejected_before_merging(self):
        other = self.root / "other-project" / "results.xlsx"
        other.parent.mkdir()
        Workbook().save(other)
        before = self.controller.inspect(self.source)
        self.assertEqual(before.summary["submissions"], 2)
        mixed = self.controller.inspect(replace(self.source, workbook=other))
        self.assertEqual(mixed.summary["submissions"], 0)
        self.assertEqual(len(mixed.attention), 1)
        self.assertEqual(mixed.attention[0].code, "SOURCE_ASSOCIATION_MISMATCH")
        self.assertIn("different tasks", project(mixed, language="en")["attention"][0]["message"])
        self.assertEqual({path: path.read_bytes() for path in self.source_hashes}, self.source_hashes)

    def test_worker_start_exception_remains_visible_without_saved_job_evidence(self):
        app = QCoreApplication.instance() or QCoreApplication([])
        actual = self.controller

        class StartupFailure:
            def inspect(self, source):
                return actual.inspect(source)

            def execute(self, action, source, **kwargs):
                raise RuntimeError("simulated worker startup failure")

        before = actual.inspect(self.source)
        bridge = DesktopBridge(controller=StartupFailure())
        bridge.setLanguage("en")
        bridge._source = self.source
        bridge._inspection = before
        bridge._state = project(before, self.source, language="en")
        try:
            bridge.startMarking()
            deadline = time.monotonic() + 5
            while bridge.busy and time.monotonic() < deadline:
                app.processEvents(QEventLoop.AllEvents, 20)
                time.sleep(0.002)
            app.processEvents(QEventLoop.AllEvents, 20)
            state = bridge.state
            self.assertEqual(state["teacherFlow"]["step"], "marking_start_failed")
            self.assertEqual(state["teacherFlow"]["primaryActionLabel"], "Continue marking")
            self.assertTrue(state["teacherFlow"]["primaryActionEnabled"])
            self.assertEqual(state["summary"]["grading_results_available"], 0)
            self.assertEqual(state["attention"][0]["details"].splitlines()[0], "MARKING_STARTUP_FAILED")
            self.assertFalse((self.root / "jobs").exists())
            self.assertFalse((self.root / "output").exists())
            self.assertEqual(len(self.transport.calls), 0)
        finally:
            bridge.shutdown()

    def test_controller_runs_independent_students_and_resumes_without_regrading(self):
        before = self.controller.inspect(self.source)
        self.assertEqual(before.summary["submissions"], 2)
        self.assertEqual(len(before.available_actions), 1)
        self.assertEqual(before.available_actions[0].action, Action.RUN_MARKING)
        self.assertTrue(before.available_actions[0].execution_wired)

        stop_after_one = threading.Event()
        progress = []

        def interrupt_after_first_saved(item):
            progress.append(dict(item))
            if item.get("completed") == 1:
                stop_after_one.set()

        with patch("workflow.pipeline.MAX_CONCURRENT_GRADERS", 1):
            first = self.controller.execute(
                Action.RUN_MARKING,
                self.source,
                progress_callback=interrupt_after_first_saved,
                cancelled=stop_after_one.is_set,
            )
        self.assertTrue(first["result"]["interrupted"])
        self.assertEqual(len(first["result"]["validated"]), 1)
        self.assertEqual(len(self.transport.calls), 1)
        self.assertTrue(progress)
        self.assertEqual(max(item["completed"] for item in progress), 1)
        self.assertTrue(progress[-1]["interrupted"])
        self.assertEqual(progress[-1]["active"], 0)
        self.assertEqual(progress[-1]["current"], "")
        self.assertFalse(any(item["current"].endswith(" · Ben Lim") for item in progress))

        reopened = self.controller.inspect(self.source)
        self.assertEqual(reopened.summary["grading_results_available"], 1)
        self.assertEqual(reopened.summary["awaiting_review"], 1)
        self.assertEqual(reopened.summary["blocking_errors"], 0)
        self.assertEqual(reopened.summary["marking_available"], 1)
        self.assertTrue(any(Action.RUN_MARKING in item.next_actions for item in reopened.submissions))
        resumed_flow = derive_teacher_flow(project(reopened, self.controller.bind_source(self.source), language="en"), "en")
        self.assertEqual(resumed_flow["step"], "resume_marking")
        self.assertEqual(resumed_flow["primaryActionLabel"], "Continue marking")
        self.assertTrue(resumed_flow["primaryActionEnabled"])

        with patch("workflow.pipeline.MAX_CONCURRENT_GRADERS", 1):
            resumed = self.controller.execute(Action.RUN_MARKING, self.source)
        self.assertFalse(resumed["result"]["failed"])
        self.assertEqual(len(self.transport.calls), 2)
        self.assertEqual(len(self.grader_instances), 2)
        self.assertEqual(len({id(grader) for grader in self.grader_instances}), 2)
        for index, (payload, _timeout) in enumerate(self.transport.calls):
            prompt = payload["input"][0]["content"][0]["text"]
            selected = self.identities[index]["student_name"]
            other = self.identities[1 - index]["student_name"]
            self.assertIn(selected, prompt)
            self.assertNotIn(other, prompt)
            encoded_pdf = payload["input"][0]["content"][1]["file_data"].split(",", 1)[1]
            import base64
            self.assertEqual(base64.b64decode(encoded_pdf), (self.pile / f"essay_{index + 1:03}.pdf").read_bytes())
            self.assertEqual(payload["model"], "gpt-5.6-luna")
            self.assertEqual(payload["reasoning"]["effort"], "xhigh")

        workbook = resumed["source"].workbook
        self.assertTrue(workbook.is_file())
        book = load_workbook(workbook, read_only=True, data_only=True)
        try:
            rows = list(book[SHEET].iter_rows(min_row=2, values_only=True))
            self.assertEqual(len(rows), 2)
            self.assertEqual([row[1] for row in rows], ["01", "02"])
            self.assertTrue(all(row[4] is None and row[5] is None and row[6] is None for row in rows))
            self.assertEqual(book[AUDIT_SHEET].max_row, 3)
        finally:
            book.close()

        reopened_final = self.controller.inspect(self.source)
        self.assertEqual(reopened_final.summary["grading_results_available"], 2)
        self.assertEqual(reopened_final.summary["awaiting_review"], 2)
        self.assertEqual(reopened_final.summary["blocking_errors"], 0)
        flow = derive_teacher_flow(project(reopened_final, self.controller.bind_source(self.source), language="en"), "en")
        self.assertEqual(flow["headline"], "Teacher review is ready ✓")
        self.assertFalse(any(self.root.rglob("student_card.png")))
        self.assertEqual(self.source_hashes, {path: path.read_bytes() for path in self.pile.glob("*.pdf")})

        # A direct accidental repeat sees durable completions; it cannot reset
        # review status or overwrite teacher-entered wording.
        identity = self.identities[0]
        student_key = job_key(Identity.from_dict(identity))
        from excel.workbook import ExcelStore
        store = ExcelStore(workbook, Validator(ROOT / "config"))
        store.set_status(student_key, Identity.from_dict(identity), "VALIDATED", "APPROVED")
        editable = load_workbook(workbook)
        result_sheet = editable[SHEET]
        row_number = next(row[0].row for row in editable[AUDIT_SHEET].iter_rows(min_row=2) if row[0].value == student_key)
        result_sheet.cell(row_number, 24).value = "Teacher-edited feedback kept on resume."
        editable.save(workbook)
        editable.close()
        self.controller.execute(Action.RUN_MARKING, self.source)
        checked = load_workbook(workbook, read_only=True, data_only=True)
        try:
            row_number = next(row[3].value for row in checked[AUDIT_SHEET].iter_rows(min_row=2) if row[0].value == student_key)
            self.assertEqual(checked[SHEET].cell(row_number, 25).value, "APPROVED")
            self.assertEqual(checked[SHEET].cell(row_number, 24).value, "Teacher-edited feedback kept on resume.")
        finally:
            checked.close()
        self.assertEqual(len(self.transport.calls), 2)

    def test_unfinished_attempt_is_included_by_the_single_continue_operation(self):
        self.transport = RecordingTransport([Identity.from_dict(item) for item in self.identities], failure_index=0)
        self.grader_instances.clear()
        partial = self.controller.execute(Action.RUN_MARKING, self.source)
        self.assertEqual(len(partial["result"]["failed"]), 1)
        self.assertEqual(len(partial["result"]["validated"]), 1)
        inspection = self.controller.inspect(self.source)
        self.assertEqual(inspection.summary["marking_available"], 1)
        failed = [item for item in inspection.submissions if item.student_id == "01"]
        self.assertEqual(len(failed), 1)
        state = project(inspection, self.controller.bind_source(self.source), language="en")
        flow = derive_teacher_flow(state, "en")
        self.assertEqual(flow["primaryAction"], Action.RUN_MARKING.value)
        self.assertEqual(flow["primaryActionLabel"], "Continue marking")

        self.transport = RecordingTransport([Identity.from_dict(item) for item in self.identities])
        retried = self.controller.execute(Action.RUN_MARKING, self.source)
        self.assertFalse(retried["result"]["failed"])
        self.assertEqual(retried["result"]["selected"], 1)
        self.assertEqual(len(self.transport.calls), 1)
        retry_prompt = self.transport.calls[0][0]["input"][0]["content"][0]["text"]
        self.assertIn(self.identities[0]["student_name"], retry_prompt)
        final = self.controller.inspect(self.source)
        self.assertEqual(final.summary["grading_results_available"], 2)
        self.assertEqual(final.summary["marking_available"], 0)

    def test_post_response_failure_remains_ordinary_unfinished_work(self):
        self.transport = RecordingTransport(
            [Identity.from_dict(item) for item in self.identities], malformed_index=0
        )
        result = self.controller.execute(Action.RUN_MARKING, self.source)
        self.assertEqual(len(result["result"]["failed"]), 1)
        inspection = self.controller.inspect(self.source)
        self.assertEqual(inspection.summary["blocking_errors"], 0)
        self.assertEqual(inspection.summary["marking_available"], 1)
        state = project(inspection, self.controller.bind_source(self.source), language="en")
        self.transport = RecordingTransport(
            [Identity.from_dict(item) for item in self.identities]
        )
        resumed = self.controller.execute(Action.RUN_MARKING, self.source)
        self.assertEqual(resumed["result"]["selected"], 1)
        self.assertEqual(len(self.transport.calls), 1)
        self.assertEqual(self.controller.inspect(self.source).summary["marking_available"], 0)

    def test_workbook_save_failure_keeps_essay_unfinished_for_ordinary_marking(self):
        original_save = ExcelStore._save
        failed_once = False

        def fail_second_student_save(store, book):
            nonlocal failed_once
            rows = list(book[SHEET].iter_rows(min_row=2, values_only=True))
            if not failed_once and rows and rows[-1][1] == "02":
                failed_once = True
                raise OSError("synthetic authoritative workbook write failure")
            return original_save(store, book)

        with patch.object(ExcelStore, "_save", fail_second_student_save):
            first = self.controller.execute(Action.RUN_MARKING, self.source)
        self.assertTrue(failed_once)
        self.assertEqual(len(first["result"]["failed"]), 1)
        self.assertEqual(len(self.transport.calls), 2)

        failed_identity = Identity.from_dict(self.identities[1])
        failed_job = (
            self.root / "jobs" / self.controller.workflow.marking_batch_id(self.source)
            / job_key(failed_identity)
        )
        self.assertEqual(working_state(failed_job), [])

        inspection = self.controller.inspect(self.source)
        failed_entry = next(item for item in inspection.submissions if item.student_id == "02")
        self.assertTrue(any(Action.RUN_MARKING in item.next_actions for item in inspection.submissions))
        self.assertEqual(inspection.summary["marking_available"], 1)
        self.assertFalse(hasattr(failed_entry, "persistence_recoverable"))

        self.transport = RecordingTransport([Identity.from_dict(item) for item in self.identities])
        resumed = self.controller.execute(Action.RUN_MARKING, self.source)
        self.assertEqual(resumed["result"]["selected"], 1)
        self.assertEqual(len(self.transport.calls), 1)
        final = self.controller.inspect(self.source)
        self.assertEqual(final.summary["grading_results_available"], 2)
        self.assertEqual(final.summary["marking_available"], 0)

    def test_validated_artifacts_without_authoritative_row_are_not_done(self):
        with patch("workflow.pipeline.MAX_CONCURRENT_GRADERS", 1):
            first = self.controller.execute(Action.RUN_MARKING, self.source)
        workbook = first["source"].workbook
        target_key = job_key(Identity.from_dict(self.identities[1]))
        self._delete_workbook_result(workbook, target_key)
        self.transport.calls.clear()

        reopened = self.controller.inspect(self.source)
        target = next(item for item in reopened.submissions if item.submission_id == target_key)
        self.assertFalse(target.workbook_valid)
        self.assertEqual(reopened.summary["grading_results_available"], 1)
        self.assertIn(Action.RUN_MARKING, target.next_actions)
        self.assertEqual(reopened.summary["marking_available"], 1)
        self.assertEqual(reopened.summary["blocking_errors"], 0)
        flow = derive_teacher_flow(project(reopened, self.controller.bind_source(self.source), language="en"), "en")
        self.assertEqual(flow["primaryAction"], Action.RUN_MARKING.value)
        self.assertEqual(flow["primaryActionLabel"], "Continue marking")

        self.transport = RecordingTransport([Identity.from_dict(item) for item in self.identities])
        with patch("workflow.pipeline.MAX_CONCURRENT_GRADERS", 1):
            resumed = self.controller.execute(Action.RUN_MARKING, self.source)
        self.assertEqual(resumed["result"]["selected"], 1)
        self.assertEqual(len(self.transport.calls), 1)
        self.assertEqual(self.controller.inspect(self.source).summary["grading_results_available"], 2)

    def test_persistence_receipt_does_not_define_done(self):
        with patch("workflow.pipeline.MAX_CONCURRENT_GRADERS", 1):
            first = self.controller.execute(Action.RUN_MARKING, self.source)
        workbook = first["source"].workbook
        target_key = job_key(Identity.from_dict(self.identities[1]))
        target_job = (
            self.root / "jobs" / self.controller.workflow.marking_batch_id(self.source) / target_key
        )
        self._delete_workbook_result(workbook, target_key)
        marker_path = target_job / "workbook_persistence.json"
        marker = {"schema_version": 1, "job_id": target_key, "status": "PENDING",
                  "workbook": str(workbook)}
        target_job.mkdir(parents=True, exist_ok=True)

        variants = [
            {**marker, "status": "SAVED"},
            {**marker, "status": "UNKNOWN"},
            {**marker, "schema_version": 99},
        ]
        for variant in variants:
            with self.subTest(marker=variant):
                marker_path.write_text(json.dumps(variant), encoding="utf-8")
                inspection = self.controller.inspect(self.source)
                target = next(item for item in inspection.submissions if item.submission_id == target_key)
                self.assertIn(Action.RUN_MARKING, target.next_actions)
                self.assertEqual(inspection.summary["blocking_errors"], 0)

        marker_path.write_text("{", encoding="utf-8")
        malformed = self.controller.inspect(self.source)
        target = next(item for item in malformed.submissions if item.submission_id == target_key)
        self.assertIn(Action.RUN_MARKING, target.next_actions)
        self.assertEqual(malformed.summary["blocking_errors"], 0)

    def test_missing_or_invalid_validated_result_remains_not_done(self):
        with patch("workflow.pipeline.MAX_CONCURRENT_GRADERS", 1):
            first = self.controller.execute(Action.RUN_MARKING, self.source)
        workbook = first["source"].workbook
        target_key = job_key(Identity.from_dict(self.identities[1]))
        target_job = (
            self.root / "jobs" / self.controller.workflow.marking_batch_id(self.source) / target_key
        )
        self._delete_workbook_result(workbook, target_key)
        target_job.mkdir(parents=True, exist_ok=True)
        validated_path = target_job / "validated_result.json"
        validated_bytes = json.dumps(_valid_response(self.identities[1])["grading"]).encode("utf-8")

        missing = self.controller.inspect(self.source)
        missing_target = next(item for item in missing.submissions if item.submission_id == target_key)
        self.assertIn(Action.RUN_MARKING, missing_target.next_actions)
        self.assertEqual(missing.summary["blocking_errors"], 0)

        validated_path.write_bytes(validated_bytes[:-2] + b"invalid\n")
        invalid = self.controller.inspect(self.source)
        invalid_target = next(item for item in invalid.submissions if item.submission_id == target_key)
        self.assertIn(Action.RUN_MARKING, invalid_target.next_actions)
        self.assertEqual(invalid.summary["blocking_errors"], 0)

    def test_workbook_row_without_audit_blocks_duplicate_risk(self):
        target_identity = Identity.from_dict(self.identities[1])
        target_key = job_key(target_identity)
        batch_id = self.controller.workflow.marking_batch_id(self.source)
        workbook = task_output_paths(self.source).workbook
        target_job = (
            self.root / "jobs" / batch_id / target_key
        )
        target_job.mkdir(parents=True)
        source_pdf = self.pile / "essay_002.pdf"
        source_bytes = source_pdf.read_bytes()
        import hashlib
        source_digest = hashlib.sha256(source_bytes).hexdigest()
        digest = fingerprint({
            "identity": target_identity.to_dict(),
            "source_sha256": source_digest,
            "essay_question": None,
            "model_profile": "luna_xhigh",
        })
        grading = _valid_response(target_identity.to_dict())["grading"]
        (target_job / "source.pdf").write_bytes(source_bytes)
        for name, value in {
            "source.json": {
                "mode": "real-pdf", "source_original_path": str(source_pdf),
                "source_sha256": source_digest, "page_count": 1,
            },
            "identity.json": {"method": "explicit-batch-input", **target_identity.to_dict()},
            "essay.json": {"method": "direct_pdf", "page_count": 1, "essay_question": None},
            "grading_result.json": grading,
            "validated_result.json": grading,
            "student_record.json": {
                "schema_version": 2, "mode": "REAL_PDF_BATCH", "job_id": target_key,
                "identity": target_identity.to_dict(), "input_digest": digest,
                "source_sha256": source_digest, "source_page_count": 1,
                "model_profile": "luna_xhigh", "execution_namespace": batch_id,
                "workbook": str(workbook), "state": "VALIDATED",
                "history": ["NEW", "SCANNED", "IDENTIFIED", "TRANSCRIBED", "GRADED", "VALIDATED"],
                "last_error": None,
            },
        }.items():
            (target_job / name).write_text(
                json.dumps(value, ensure_ascii=False), encoding="utf-8"
            )

        workbook.parent.mkdir(parents=True)
        editable = Workbook()
        sheet = editable.active
        sheet.title = SHEET
        sheet.append(HEADERS)
        values = [target_identity.class_name, target_identity.student_id,
                  target_identity.student_name, "", None, None, None]
        for name in CRITERIA:
            values.extend((grading["criteria"][name]["rating"],
                           grading["criteria"][name]["short_comment"]))
        values.extend((grading["teacher_comment"], "PENDING"))
        sheet.append(values)
        audit = editable.create_sheet(AUDIT_SHEET)
        audit.append(AUDIT_HEADERS)
        audit.sheet_state = "hidden"
        editable.save(workbook)
        editable.close()

        reopened = self.controller.inspect(self.source)
        target = next(item for item in reopened.submissions if item.submission_id == target_key)
        self.assertNotIn(Action.RUN_MARKING, target.next_actions)
        self.assertGreater(reopened.summary["blocking_errors"], 0)
        self.assertIn("WORKBOOK_ROW_INVALID", {item.code for item in target.attention})

    def test_pending_persistence_receipt_does_not_reopen_committed_result(self):
        with patch("workflow.pipeline.MAX_CONCURRENT_GRADERS", 1):
            first = self.controller.execute(Action.RUN_MARKING, self.source)
        workbook = first["source"].workbook
        identity = Identity.from_dict(self.identities[1])
        key = job_key(identity)
        marker_path = (
            self.root / "jobs" / self.controller.workflow.marking_batch_id(self.source)
            / key / "workbook_persistence.json"
        )
        marker_path.parent.mkdir(parents=True, exist_ok=True)
        marker = {"schema_version": 1, "job_id": key, "status": "PENDING", "workbook": str(workbook)}
        marker_path.write_text(json.dumps(marker, ensure_ascii=False), encoding="utf-8")
        workbook_before = workbook.read_bytes()
        calls_before = len(self.transport.calls)

        reopened = self.controller.inspect(self.source)
        target = next(item for item in reopened.submissions if item.submission_id == key)
        self.assertTrue(target.workbook_valid)
        self.assertNotIn(Action.RUN_MARKING, target.next_actions)
        self.assertEqual(len(self.transport.calls), calls_before)
        self.assertEqual(workbook.read_bytes(), workbook_before)
        saved_marker = json.loads(marker_path.read_text(encoding="utf-8"))
        self.assertEqual(saved_marker["status"], "PENDING")
        self.assertEqual(self.controller.inspect(self.source).summary["marking_available"], 0)

    def test_committed_rows_are_excluded_from_direct_marking_selection(self):
        with patch("workflow.pipeline.MAX_CONCURRENT_GRADERS", 1):
            self.controller.execute(Action.RUN_MARKING, self.source)
        target_key = job_key(Identity.from_dict(self.identities[1]))
        marker_path = (
            self.root / "jobs" / self.controller.workflow.marking_batch_id(self.source)
            / target_key / "workbook_persistence.json"
        )
        marker_path.parent.mkdir(parents=True, exist_ok=True)
        marker = {"schema_version": 1, "job_id": target_key, "status": "PENDING"}
        marker_path.write_text(json.dumps(marker), encoding="utf-8")

        selected = []

        def capture_run(pipeline, records, source_dir, **kwargs):
            selected.extend(job_key(Identity.from_dict({
                key: record[key] for key in ("class_name", "student_id", "student_name")
            })) for record in records)
            return {
                "selected": len(records), "validated": [], "failed": [],
                "already_complete": 0, "workbook": str(pipeline.excel.path),
            }

        with patch.object(Pipeline, "run_real_batch", capture_run):
            self.controller.execute(Action.RUN_MARKING, self.source)
        self.assertNotIn(target_key, selected)

    def test_desktop_runs_off_gui_thread_and_rejects_duplicate_clicks(self):
        app = QCoreApplication.instance() or QCoreApplication([])
        base = RecordingTransport([Identity.from_dict(item) for item in self.identities])
        caller_threads = []

        class SlowTransport:
            def ensure_ready(self):
                return None

            def __call__(self, payload, timeout_seconds):
                caller_threads.append(threading.get_ident())
                time.sleep(0.12)
                return base(payload, timeout_seconds)

        self.transport = SlowTransport()
        inspection = self.controller.inspect(self.source)
        state = project(inspection, self.source, language="en")
        bridge = DesktopBridge(controller=self.controller)
        bridge.setLanguage("en")
        bridge._source = self.source
        bridge._inspection = inspection
        bridge._state = state
        bridge._real_state = state
        heartbeat = []
        QTimer.singleShot(30, lambda: heartbeat.append("gui-responsive"))
        gui_thread = threading.get_ident()

        bridge.startMarking()
        bridge.startMarking()
        deadline = time.monotonic() + 5
        while bridge.busy and time.monotonic() < deadline:
            app.processEvents(QEventLoop.AllEvents, 20)
            time.sleep(0.002)
        app.processEvents(QEventLoop.AllEvents, 20)
        try:
            self.assertFalse(bridge.busy)
            self.assertEqual(len(base.calls), 2)
            self.assertTrue(heartbeat)
            self.assertEqual(len(caller_threads), 2)
            self.assertTrue(all(thread_id != gui_thread for thread_id in caller_threads))
            self.assertEqual(bridge.state["teacherFlow"]["step"], "review")
            self.assertEqual(bridge.state["progress"]["completed"], 2)
        finally:
            bridge.shutdown()

    def test_graceful_desktop_close_saves_current_essay_and_stops_before_next(self):
        app = QCoreApplication.instance() or QCoreApplication([])
        base = RecordingTransport([Identity.from_dict(item) for item in self.identities])
        call_started = threading.Event()

        class SlowTransport:
            def ensure_ready(self):
                return None

            def __call__(self, payload, timeout_seconds):
                call_started.set()
                time.sleep(0.12)
                return base(payload, timeout_seconds)

        self.transport = SlowTransport()
        inspection = self.controller.inspect(self.source)
        bridge = DesktopBridge(controller=self.controller)
        bridge.setLanguage("en")
        bridge._source = self.source
        bridge._inspection = inspection
        bridge._state = project(inspection, self.source, language="en")
        with patch("workflow.pipeline.MAX_CONCURRENT_GRADERS", 1):
            bridge.startMarking()
            self.assertTrue(call_started.wait(4), "The first isolated grading call did not start")
            bridge.shutdown()
        app.processEvents(QEventLoop.AllEvents, 20)
        try:
            self.assertEqual(len(base.calls), 1)
            reopened = self.controller.inspect(self.source)
            self.assertEqual(reopened.summary["grading_results_available"], 1)
            self.assertEqual(reopened.summary["awaiting_review"], 1)
            self.assertEqual(reopened.summary["marking_available"], 1)
            flow = derive_teacher_flow(project(reopened, self.controller.bind_source(self.source), language="en"), "en")
            self.assertEqual(flow["step"], "resume_marking")
            self.assertEqual(flow["primaryActionLabel"], "Continue marking")
            self.assertEqual(flow["headline"], "1 of 2 saved · 1 still to mark")
        finally:
            bridge.shutdown()

    def test_cancel_marking_button_request_stops_queue_and_keeps_saved_result(self):
        app = QCoreApplication.instance() or QCoreApplication([])
        base = RecordingTransport([Identity.from_dict(item) for item in self.identities])
        call_started = threading.Event()
        release_call = threading.Event()

        class HeldTransport:
            def ensure_ready(self):
                return None

            def __call__(self, payload, timeout_seconds):
                call_started.set()
                if not release_call.wait(5):
                    raise TimeoutError("test grading call was not released")
                return base(payload, timeout_seconds)

        self.transport = HeldTransport()
        inspection = self.controller.inspect(self.source)
        bridge = DesktopBridge(controller=self.controller)
        bridge.setLanguage("en")
        bridge._source = self.source
        bridge._inspection = inspection
        bridge._state = project(inspection, self.source, language="en")
        try:
            with patch("workflow.pipeline.MAX_CONCURRENT_GRADERS", 1):
                bridge.startMarking()
                self.assertTrue(call_started.wait(4), "The first isolated grading call did not start")
                bridge.cancelMarking()
                bridge.cancelMarking()
                self.assertTrue(bridge.state["progress"]["cancelRequested"])
                self.assertIn("Finishing", bridge.state["progress"]["saved"])
                release_call.set()
                deadline = time.monotonic() + 5
                while bridge.busy and time.monotonic() < deadline:
                    app.processEvents(QEventLoop.AllEvents, 20)
                    time.sleep(0.002)
                app.processEvents(QEventLoop.AllEvents, 20)
            self.assertFalse(bridge.busy)
            self.assertEqual(len(base.calls), 1)
            self.assertTrue(bridge.state["markingResult"]["interrupted"])
            self.assertEqual(self.controller.inspect(self.source).summary["grading_results_available"], 1)
        finally:
            release_call.set()
            bridge.shutdown()


if __name__ == "__main__":
    unittest.main()
