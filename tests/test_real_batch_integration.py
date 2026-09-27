import base64
import copy
import json
from pathlib import Path
import unittest

from openpyxl import load_workbook
from pypdf import PdfWriter

from excel.schema import AUDIT_SHEET, HEADERS, SHEET
from grading.schemas import CRITERIA, Identity
from grading.sol_grader import SolGrader
from tests.local_temp import local_test_directory
from workflow.calibration_pipeline import CalibrationPipeline
from workflow.pipeline import Pipeline, job_key


def valid_response(identity):
    return {
        "reading_quality": {
            "complete_essay_legible": True,
            "uncertainties": [],
            "materially_affects_grade": False,
        },
        "grading": {
            **identity.to_dict(),
            "criteria": {
                name: {"rating": "可以更进一步", "short_comment": f"围绕{name}补充具体细节。"}
                for name in CRITERIA
            },
            "teacher_comment": "文章内容完整，关键段落仍可补充更具体的细节。",
        },
        "evidence": [{"judgment": "重点突出", "page_numbers": [1], "rationale": "中段的主要行动集中呈现。"}],
    }


def working_state(essay_dir):
    """Every file in an essay folder except Stage 3C card output."""
    essay_dir = Path(essay_dir)
    if not essay_dir.is_dir():
        return []
    return sorted(
        path.relative_to(essay_dir).as_posix()
        for path in essay_dir.rglob("*")
        if path.is_file() and path.relative_to(essay_dir).parts[0] != "output"
    )


def raw_response(value):
    return {"id": "batch-test", "model": "gpt-5.6-luna", "status": "completed",
            "output_text": json.dumps(value, ensure_ascii=False) if isinstance(value, dict) else value}


class RecordingTransport:
    def __init__(self, identities, failure_index=None, malformed_index=None, validation_index=None):
        self.identities = identities
        self.failure_index = failure_index
        self.malformed_index = malformed_index
        self.validation_index = validation_index
        self.calls = []

    def ensure_ready(self):
        return None

    def __call__(self, payload, timeout_seconds):
        self.calls.append((copy.deepcopy(payload), timeout_seconds))
        prompt = payload["input"][0]["content"][0]["text"]
        identity = next(item for item in self.identities if item.student_name in prompt)
        index = self.identities.index(identity)
        if index == self.failure_index:
            raise RuntimeError("synthetic grader failure")
        if index == self.malformed_index:
            return raw_response("not-json")
        value = valid_response(identity)
        if index == self.validation_index:
            value["grading"]["criteria"][CRITERIA[0]]["rating"] = "优秀"
        return raw_response(value)


class RealBatchIntegrationTests(unittest.TestCase):
    def setUp(self):
        self.temp = local_test_directory("real-batch")
        self.root = Path(self.temp.__enter__())
        self.addCleanup(self.temp.__exit__, None, None, None)
        self.source_dir = self.root / "pile"
        self.source_dir.mkdir()
        self.identities = [
            Identity("01", "甲学生", "207"),
            Identity("02", "乙学生", "207"),
            Identity("03", "丙学生", "207"),
        ]
        self.records = []
        for sequence, identity in enumerate(self.identities, 1):
            path = self.source_dir / f"submission_{sequence:03}.pdf"
            writer = PdfWriter()
            writer.add_blank_page(width=612, height=792)
            writer.add_metadata({"/StudentMarker": identity.student_name})
            with path.open("wb") as handle:
                writer.write(handle)
            self.records.append({
                "sequence": sequence,
                "source_pdf": path.name,
                **identity.to_dict(),
                "match_status": "STRONG_ROSTER_MATCH",
            })

    def pipeline_factory(self, transport):
        def factory(project_dir, model_profile=None):
            config_dir = Path(__file__).resolve().parents[1] / "config"
            profile = json.loads((config_dir / "calibration_model.json").read_text(encoding="utf-8"))["profiles"][model_profile]
            return CalibrationPipeline(
                project_dir,
                config_dir=config_dir,
                model_profile=model_profile,
                grader_factory=lambda: SolGrader(
                    model=profile["model"],
                    reasoning_effort=profile["reasoning_effort"],
                    transport=transport,
                ),
            )
        return factory

    def make_pipeline(self, transport, batch="real-batch"):
        return Pipeline(
            self.root,
            batch,
            allow_real=True,
            real_pipeline_factory=self.pipeline_factory(transport),
        )

    def rewrite_sources(self, marker):
        for record, identity in zip(self.records, self.identities):
            path = self.source_dir / record["source_pdf"]
            writer = PdfWriter()
            writer.add_blank_page(width=612, height=792)
            writer.add_metadata({"/StudentMarker": f"{marker}-{identity.student_name}"})
            with path.open("wb") as handle:
                writer.write(handle)

    def assert_recorded_inputs_are_isolated(self, transport, identities=None):
        identities = identities or self.identities
        self.assertEqual(len(transport.calls), len(identities))
        for payload, _timeout in transport.calls:
            prompt = payload["input"][0]["content"][0]["text"]
            identity = next(item for item in identities if item.student_name in prompt)
            self.assertIn(identity.student_name, prompt)
            for other in identities:
                if other != identity:
                    self.assertNotIn(other.student_name, prompt)
            file_data = payload["input"][0]["content"][1]["file_data"]
            self.assertTrue(file_data.startswith("data:application/pdf;base64,"))
            record = next(item for item in self.records if item["student_id"] == identity.student_id)
            expected = (self.source_dir / record["source_pdf"]).read_bytes()
            self.assertEqual(base64.b64decode(file_data.split(",", 1)[1]), expected)

    def test_multiple_students_limit_isolated_and_consolidated(self):
        transport = RecordingTransport(self.identities)
        pipeline = self.make_pipeline(transport)
        result = pipeline.run_real_batch(
            self.records,
            self.source_dir,
            limit=2,
            essay_question="Q2 题目",
            model_profile="luna_xhigh",
        )
        self.assertEqual(result["selected"], 2)
        self.assertEqual(len(result["validated"]), 2)
        self.assertEqual(result["failed"], [])
        self.assertEqual(len(transport.calls), 2)
        for identity in self.identities[:2]:
            self.assertEqual(working_state(self.root / "jobs" / "real-batch" / job_key(identity)), [])
        book = load_workbook(result["workbook"], read_only=True, data_only=True)
        try:
            self.assertEqual([cell.value for cell in book[SHEET][1]], HEADERS)
            rows = list(book[SHEET].iter_rows(min_row=2, values_only=True))
            self.assertEqual(len(rows), 2)
            self.assertEqual([row[1] for row in rows], ["01", "02"])
            self.assertTrue(all(row[4] is None and row[5] is None and row[6] is None for row in rows))
            self.assertNotIn("下次提升重点", HEADERS)
            self.assertEqual(book[AUDIT_SHEET].max_row, 3)
        finally:
            book.close()
        self.assert_recorded_inputs_are_isolated(transport, self.identities[:2])

    def test_failure_isolation_excludes_failed_student_and_continues(self):
        for label, kwargs in (("exception", {"failure_index": 1}), ("malformed", {"malformed_index": 1}),
                              ("validation", {"validation_index": 1})):
            with self.subTest(label=label):
                if label != "exception":
                    self.rewrite_sources(label)
                transport = RecordingTransport(self.identities, **kwargs)
                pipeline = self.make_pipeline(transport, f"real-batch-{label}")
                result = pipeline.run_real_batch(self.records, self.source_dir, model_profile="luna_xhigh")
                self.assertEqual(result["selected"], 3)
                self.assertEqual(len(result["validated"]), 2)
                self.assertEqual(len(result["failed"]), 1)
                self.assertEqual(len(transport.calls), 3)
                book = load_workbook(result["workbook"], read_only=True, data_only=True)
                try:
                    rows = list(book[SHEET].iter_rows(min_row=2, values_only=True))
                    self.assertEqual([row[1] for row in rows], ["01", "03"])
                finally:
                    book.close()
                # A failed essay is not done and keeps nothing to resume from.
                self.assertEqual(working_state(self.root / "jobs" / f"real-batch-{label}" / job_key(self.identities[1])), [])


if __name__ == "__main__":
    unittest.main()
