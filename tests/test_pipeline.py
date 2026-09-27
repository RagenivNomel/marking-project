import copy
from pathlib import Path
import unittest
from unittest.mock import Mock, patch
from openpyxl import load_workbook, Workbook
from excel.schema import AUDIT_HEADERS, AUDIT_SHEET, HEADERS, SHEET
from excel.workbook import APPROVAL_STATUSES, APPROVAL_VALIDATION_RANGE, ExcelStore
from grading.grader import GradingInput
from excel.workbook import read_roster
from grading.mock_grader import MockGrader
from grading.schemas import CRITERIA, Identity, ROOT, ValidationError
from scanning.identity import confirm_identity
from workflow.pipeline import Pipeline, job_key
from workflow.state import State
from workflow.storage import batch_lock, read_json
from tests.local_temp import local_test_directory


class PipelineTests(unittest.TestCase):
    def setUp(self):
        self.temp = local_test_directory(self._testMethodName)
        self.root = Path(self.temp.__enter__())
        self.addCleanup(self.temp.__exit__, None, None, None)
        self.source = read_json(ROOT / "examples/mock_student.json")
        self.identity = Identity.from_dict(self.source["identity"])
        self.pipeline = Pipeline(self.root)
        self.job = self.pipeline.jobs_dir / job_key(self.identity)

    def edit_excel(self, cell, value):
        book = load_workbook(self.pipeline.excel.path)
        try:
            book[SHEET][cell] = value
            book.save(self.pipeline.excel.path)
        finally:
            book.close()

    def test_end_to_end_and_complete_restart_does_no_work(self):
        result = self.pipeline.run_mock(self.source)
        self.assertEqual(result["state"], "COMPLETE")
        checkpoint = read_json(self.job / "student_record.json")
        self.assertEqual(checkpoint["history"], [s.value for s in State])
        approved = self.pipeline.excel.get(job_key(self.identity), self.identity, approved=True)
        self.assertEqual(read_json(result["render"])["approved_excel_values"], approved.to_dict())
        before = self.pipeline.excel.path.read_bytes()
        with patch.object(MockGrader, "grade", side_effect=AssertionError("Should not regrade")), patch.object(self.pipeline.renderer, "render", side_effect=AssertionError("Should not rerender")):
            self.assertEqual(self.pipeline.run_mock(self.source), result)
        self.assertEqual(self.pipeline.excel.path.read_bytes(), before)

    def test_wide_teacher_table_keeps_scores_blank_and_audit_outside_main_sheet(self):
        source = {**self.source, "topic": "Q2 虚构题目"}
        result = self.pipeline.run_mock(source)
        book = load_workbook(result["workbook"])
        try:
            self.assertEqual([cell.value for cell in book[SHEET][1]], HEADERS)
            row = dict(zip(HEADERS, [cell.value for cell in book[SHEET][2]]))
            self.assertEqual(row["题目"], "Q2 虚构题目")
            self.assertIsNone(row["内容分"])
            self.assertIsNone(row["语文与结构分"])
            self.assertIsNone(row["总分"])
            self.assertNotIn("任务ID", row)
            self.assertNotIn("流程状态", row)
            self.assertEqual(book[AUDIT_SHEET].sheet_state, "hidden")
            self.assertEqual([cell.value for cell in book[AUDIT_SHEET][1]], AUDIT_HEADERS)
            book[SHEET]["E2"] = 18
            book[SHEET]["F2"] = 27
            book[SHEET]["G2"] = 45
            book.save(result["workbook"])
        finally:
            book.close()
        record = self.pipeline.excel.get(job_key(self.identity), self.identity, approved=True)
        self.assertEqual((record.content_score, record.language_structure_score, record.total_score), (18, 27, 45))

    def test_authoritative_teacher_scores_can_be_imported_without_grader_scores(self):
        result = self.pipeline.validator.validate(
            MockGrader().grade(GradingInput(self.identity, self.source["essay"])),
            self.identity,
        )
        store = ExcelStore(self.root / "teacher-score-import" / "results.xlsx", self.pipeline.validator)
        store.ensure_draft(
            result,
            self.identity,
            "authoritative-score-import",
            "source-digest",
            "Q2",
            {"content_score": 22, "language_structure_score": 21, "total_score": 43},
        )
        record = store.get("authoritative-score-import", self.identity)
        self.assertEqual((record.topic, record.content_score, record.language_structure_score, record.total_score), ("Q2", 22, 21, 43))

    def test_workbook_keeps_one_row_per_student(self):
        result = self.pipeline.validator.validate(
            MockGrader().grade(GradingInput(self.identity, self.source["essay"])),
            self.identity,
        )
        store = ExcelStore(self.root / "one-row-per-student" / "results.xlsx", self.pipeline.validator)
        store.ensure_draft(result, self.identity, "first-submission", "source-digest")
        store.ensure_draft(result, self.identity, "first-submission", "source-digest")  # resume: no-op
        with self.assertRaisesRegex(ValidationError, "lacks matching audit evidence"):
            store.ensure_draft(result, self.identity, "second-submission", "other-digest")
        book = load_workbook(store.path)
        try:
            self.assertEqual(book[SHEET].max_row, 2)
        finally:
            book.close()

    def test_production_workbook_has_approval_dropdown_and_preserves_approved_value(self):
        validated = self.pipeline.validator.validate(
            MockGrader().grade(GradingInput(self.identity, self.source["essay"])),
            self.identity,
        )
        draft_store = ExcelStore(self.root / "approval-default" / "results.xlsx", self.pipeline.validator)
        draft_store.ensure_draft(validated, self.identity, "approval-default", "source-digest")
        draft_book = load_workbook(draft_store.path)
        try:
            self.assertEqual(draft_book[SHEET]["Y2"].value, "PENDING")
        finally:
            draft_book.close()

        result = self.pipeline.run_mock(self.source)
        book = load_workbook(result["workbook"])
        try:
            validations = list(book[SHEET].data_validations.dataValidation)
            approval = [item for item in validations if str(item.sqref) == APPROVAL_VALIDATION_RANGE]
            self.assertEqual(len(approval), 1)
            self.assertEqual(approval[0].formula1, '"' + ','.join(APPROVAL_STATUSES) + '"')
            rating_ranges = {str(item.sqref) for item in validations if item.formula1 == '"' + ','.join(("做得很好", "可以更进一步", "待加强")) + '"'}
            self.assertEqual(
                rating_ranges,
                {f"{column}2:{column}1048576" for column in ("H", "J", "L", "N", "P", "R", "T", "V")},
            )
            self.assertEqual(book[SHEET]["Y2"].value, "APPROVED")
            book[SHEET]["Y2"] = "APPROVED"
            book.save(result["workbook"])
        finally:
            book.close()
        self.pipeline.excel.ensure_validation_rules()
        reopened = load_workbook(result["workbook"], data_only=True)
        try:
            self.assertEqual(reopened[SHEET]["Y2"].value, "APPROVED")
            self.assertEqual(len([item for item in reopened[SHEET].data_validations.dataValidation if str(item.sqref) == APPROVAL_VALIDATION_RANGE]), 1)
            self.assertEqual(reopened[SHEET].max_column, 25)
        finally:
            reopened.close()

    def test_every_required_failure_is_end_to_end_and_never_writes_excel_or_renders(self):
        def invalid_rating(result):
            result["criteria"][next(iter(result["criteria"]))]["rating"] = "优秀"

        def missing_criterion(result):
            result["criteria"].pop(next(iter(result["criteria"])))

        def empty_comment(result):
            result["criteria"][next(iter(result["criteria"]))]["short_comment"] = " \n"

        def over_limit(result):
            result["criteria"][next(iter(result["criteria"]))]["short_comment"] = "长" * 91

        def identity_mismatch(result):
            result["student_name"] = "另一个学生"

        for name, mutate in (
            ("invalid-rating", invalid_rating),
            ("missing-criterion", missing_criterion),
            ("empty-comment", empty_comment),
            ("over-limit", over_limit),
            ("identity-mismatch", identity_mismatch),
        ):
            with self.subTest(failure=name), local_test_directory(name) as directory:
                class MutatingGrader:
                    def grade(inner_self, request):
                        raw = MockGrader().grade(request)
                        mutate(raw)
                        return raw

                renderer = Mock()
                pipeline = Pipeline(Path(directory), name, grader_factory=MutatingGrader, renderer=renderer)
                with self.assertRaises(ValidationError):
                    pipeline.run_mock(self.source)
                job = pipeline.jobs_dir / job_key(self.identity)
                self.assertFalse(pipeline.excel.path.exists())
                self.assertFalse((job / "output").exists())
                renderer.render.assert_not_called()
                checkpoint = read_json(job / "student_record.json")
                self.assertEqual(checkpoint["state"], "GRADED")
                self.assertEqual(checkpoint["last_error"]["type"], "ValidationError")

        with self.subTest(failure="invalid-transition"), local_test_directory("invalid-transition") as directory:
            renderer = Mock()
            pipeline = Pipeline(Path(directory), "invalid-transition", renderer=renderer)
            original_advance = pipeline._advance

            def skip_required_state(job, checkpoint, _target):
                return original_advance(job, checkpoint, State.IDENTIFIED)

            with patch.object(pipeline, "_advance", side_effect=skip_required_state), self.assertRaises(ValueError):
                pipeline.run_mock(self.source)
            job = pipeline.jobs_dir / job_key(self.identity)
            self.assertFalse(pipeline.excel.path.exists())
            self.assertFalse((job / "output").exists())
            renderer.render.assert_not_called()
            checkpoint = read_json(job / "student_record.json")
            self.assertEqual(checkpoint["state"], "NEW")
            self.assertEqual(checkpoint["last_error"]["type"], "ValueError")

    def test_resume_after_every_durable_state(self):
        for stop in list(State)[1:-1]:
            with self.subTest(stop=stop):
                pipeline = Pipeline(self.root, stop.value)
                original = pipeline._advance
                def interrupt(job, checkpoint, target):
                    original(job, checkpoint, target)
                    if target == stop:
                        raise RuntimeError("simulated stop")
                with patch.object(pipeline, "_advance", side_effect=interrupt), self.assertRaises(RuntimeError):
                    pipeline.run_mock(self.source)
                guard = AssertionError("Persisted grading must be reused") if list(State).index(stop) >= list(State).index(State.GRADED) else None
                if guard:
                    with patch.object(MockGrader, "grade", side_effect=guard):
                        result = Pipeline(self.root, stop.value).run_mock(self.source)
                else:
                    result = Pipeline(self.root, stop.value).run_mock(self.source)
                self.assertEqual(result["state"], "COMPLETE")
                book = load_workbook(result["workbook"])
                self.assertEqual(book[SHEET].max_row, 2)
                book.close()

    def test_resume_between_excel_write_and_checkpoint_preserves_teacher_edit(self):
        with patch.object(self.pipeline, "_advance", wraps=self.pipeline._advance) as advance:
            original = self.pipeline._advance._mock_wraps
            def stop_before(job, checkpoint, target):
                if target == State.REVIEWED:
                    raise RuntimeError("stop after Excel write")
                original(job, checkpoint, target)
            advance.side_effect = stop_before
            with self.assertRaises(RuntimeError):
                self.pipeline.run_mock(self.source)
        self.edit_excel("X2", "教师手动修改，必须保留。")
        result = self.pipeline.run_mock(self.source)
        self.assertEqual(read_json(result["render"])["approved_excel_values"]["teacher_comment"], "教师手动修改，必须保留。")

    def test_resume_after_grader_result_saved_before_checkpoint(self):
        original = self.pipeline._advance
        def stop(job, checkpoint, target):
            if target == State.GRADED:
                raise RuntimeError("stop after raw JSON saved")
            original(job, checkpoint, target)
        with patch.object(self.pipeline, "_advance", side_effect=stop), self.assertRaises(RuntimeError):
            self.pipeline.run_mock(self.source)
        with patch.object(MockGrader, "grade", side_effect=AssertionError("Do not repeat provider")):
            self.assertEqual(self.pipeline.run_mock(self.source)["state"], "COMPLETE")

    def test_rerender_uses_exact_excel_edits_without_grading(self):
        self.pipeline.run_mock(self.source)
        text = "  教师原话，保留空格。\n下一行。"
        self.edit_excel("X2", text)
        with patch.object(MockGrader, "grade", side_effect=AssertionError("Should not grade")):
            artifact = self.pipeline.rerender(self.identity)
        self.assertEqual(read_json(artifact)["approved_excel_values"]["teacher_comment"], text)

    def test_approved_imported_workbook_rerenders_without_production_checkpoint(self):
        renderer = Mock()
        renderer.render.return_value = self.root / "imported-output" / "student_card.png"
        pipeline = Pipeline(self.root, "imported-production", renderer=renderer)
        result = pipeline.validator.validate(
            MockGrader().grade(GradingInput(self.identity, self.source["essay"])),
            self.identity,
        )
        original_job = "calibration_synthetic_fixture_luna_xhigh_test"
        pipeline.excel.ensure_draft(
            result,
            self.identity,
            original_job,
            "approved-input-digest",
            "Q2",
            {"content_score": 22, "language_structure_score": 21, "total_score": 43},
        )
        pipeline.excel.set_status(original_job, self.identity, State.VALIDATED.value, "APPROVED")
        checkpoint = pipeline.jobs_dir / job_key(self.identity) / "student_record.json"
        self.assertFalse(checkpoint.exists())
        with patch.object(MockGrader, "grade", side_effect=AssertionError("Rerender must not grade")):
            artifact = pipeline.rerender(self.identity)
        self.assertEqual(artifact, renderer.render.return_value)
        record, _output_dir, audit = renderer.render.call_args.args
        self.assertEqual(record.topic, "Q2")
        self.assertEqual((record.content_score, record.language_structure_score, record.total_score), (22, 21, 43))
        self.assertEqual(record.review_status, "APPROVED")
        self.assertEqual(tuple(record.criteria), CRITERIA)
        self.assertEqual(audit["job_id"], original_job)
        self.assertEqual(audit["excel_row"], 2)
        self.assertEqual(audit["input_digest"], "approved-input-digest")
        self.assertEqual(audit["source_of_truth"], "approved_excel_row")

    def test_unapproved_invalid_and_formula_excel_edits_block_rerender(self):
        result = self.pipeline.run_mock(self.source)
        original = Path(result["render"]).read_bytes()
        for cell, value, restore in (("Y2", "PENDING", "APPROVED"), ("H2", "优秀", "可以更进一步"), ("X2", "=1+1", "恢复评语"), ("W2", "字" * 91, "恢复评语")):
            with self.subTest(cell=cell):
                self.edit_excel(cell, value)
                with self.assertRaises(ValidationError):
                    self.pipeline.rerender(self.identity)
                self.assertEqual(Path(result["render"]).read_bytes(), original)
                self.edit_excel(cell, restore)

    def test_student_isolation_same_number_different_classes(self):
        inputs, providers = [], []
        def factory():
            provider = MockGrader()
            providers.append(provider)
            old_grade = provider.grade
            def grade(request):
                inputs.append(request)
                return old_grade(request)
            provider.grade = grade
            return provider
        pipeline = Pipeline(self.root, grader_factory=factory)
        first = pipeline.run_mock(self.source)
        second_source = copy.deepcopy(self.source)
        second_source["identity"].update(class_name="MOCK-208", student_name="另一测试学生")
        second_source["essay"] = "完全不同的第二篇作文"
        second = pipeline.run_mock(second_source)
        self.assertNotEqual(first["job_dir"], second["job_dir"])
        self.assertIsNot(providers[0], providers[1])
        self.assertEqual(inputs[0].essay, self.source["essay"])
        self.assertEqual(inputs[1].essay, second_source["essay"])
        self.assertEqual(set(vars(inputs[0])), {"identity", "essay"})

    def test_changed_input_and_path_traversal_are_rejected(self):
        self.pipeline.run_mock(self.source)
        changed = {**self.source, "essay": "新的作文"}
        with self.assertRaises(ValidationError):
            self.pipeline.run_mock(changed)
        for batch in ("../outside", "C:/outside", "CON", "", "a/b"):
            with self.subTest(batch=batch), self.assertRaises(ValueError):
                Pipeline(self.root, batch)
        hostile = Identity("../../etc", "测试", "../class")
        self.assertNotIn("..", job_key(hostile))
        self.assertNotIn("/", job_key(hostile))

    def test_batch_lock_blocks_concurrent_writer(self):
        with batch_lock(self.pipeline.jobs_dir / ".pipeline.lock"):
            with self.assertRaises(RuntimeError):
                self.pipeline.run_mock(self.source)

    def test_excel_backups_are_made(self):
        self.pipeline.run_mock(self.source)
        self.assertGreaterEqual(len(list((self.pipeline.output_dir / "backups").glob("*.xlsx"))), 1)

    def test_roster_adapter_and_closed_set_confirmation(self):
        path = self.root / "roster.xlsx"
        book = Workbook()
        book.active.title = "作文诊断输入"
        book.active.append(["班号", "学生姓名", "班级"])
        book.active.append([16, "测试学生", 207])
        book.save(path)
        book.close()
        roster = read_roster(path)
        self.assertEqual(confirm_identity(roster, "207", "16"), Identity("16", "测试学生", "207"))
        with self.assertRaises(ValidationError):
            confirm_identity(roster, "207", "999")
        with self.assertRaises(ValidationError):
            confirm_identity(roster + roster, "207", "16")

    def test_failed_review_blocks_approval_and_rendering(self):
        with patch.object(self.pipeline.reviewer, "review", return_value=False), self.assertRaises(ValidationError):
            self.pipeline.run_mock(self.source)
        self.assertEqual(read_json(self.job / "student_record.json")["state"], "VALIDATED")
        with self.assertRaises(ValidationError):
            self.pipeline.excel.get(job_key(self.identity), self.identity, approved=True)
        self.assertFalse((self.job / "output").exists())
