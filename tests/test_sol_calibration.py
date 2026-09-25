import base64
import copy
import json
from pathlib import Path
import unittest

from openpyxl import load_workbook
from pypdf import PdfWriter

from excel.calibration_workbook import BLIND_SHEET, META_SHEET, REVIEW_STATUS
from excel.schema import HEADERS
from grading.schemas import CRITERIA, Identity, ValidationError
from grading.sol_grader import CredentialUnavailable, ModelCallError, ResponsesHttpTransport, SolGrader
from tests.local_temp import local_test_directory
from workflow.calibration_pipeline import CalibrationPipeline, sha256_file
from workflow.storage import read_json


def valid_output(identity):
    return {
        "reading_quality": {
            "complete_essay_legible": True,
            "uncertainties": [{"page_number": 2, "description": "一个字形略不清楚，但不影响句意。"}],
            "materially_affects_grade": False,
        },
        "grading": {
            **identity.to_dict(),
            "criteria": {
                name: {"rating": "可以更进一步", "short_comment": f"围绕{name}补充具体细节。"}
                for name in CRITERIA
            },
            "teacher_comment": "文章围绕一次校园义卖展开，过程完整，关键合作片段仍可写得更具体。",
        },
        "evidence": [{"judgment": "重点突出", "page_numbers": [2, 3], "rationale": "摊位互动集中在中段，但关键合作片段展开较短。"}],
    }


def raw_response(value):
    text = value if isinstance(value, str) else json.dumps(value, ensure_ascii=False)
    return {"id": "resp_test", "model": "gpt-5.6-sol", "status": "completed", "output_text": text}


class FakeTransport:
    def __init__(self, response=None, error=None):
        self.response = response
        self.error = error
        self.calls = []

    def ensure_ready(self):
        return None

    def __call__(self, payload, timeout_seconds):
        self.calls.append((payload, timeout_seconds))
        if self.error:
            raise self.error
        return copy.deepcopy(self.response)


class SolCalibrationTests(unittest.TestCase):
    def setUp(self):
        self.temp = local_test_directory(self._testMethodName)
        self.root = Path(self.temp.__enter__())
        self.addCleanup(self.temp.__exit__, None, None, None)
        self.identity = Identity("SYNTHETIC-16", "合成学生", "SYNTHETIC-CLASS")
        self.source = self.root / "selected.pdf"
        writer = PdfWriter()
        for _ in range(4):
            writer.add_blank_page(width=612, height=792)
        writer.write(self.source)

    def pipeline(self, transport):
        return CalibrationPipeline(self.root, grader_factory=lambda: SolGrader(transport=transport))

    def find_job(self):
        return next((self.root / "jobs").glob("calibration_*"))

    def assert_no_workbook(self):
        self.assertEqual(list((self.root / "output").rglob("*.xlsx")) if (self.root / "output").exists() else [], [])

    def test_valid_structured_output_writes_pending_workbook_and_stops_validated(self):
        transport = FakeTransport(raw_response(valid_output(self.identity)))
        result = self.pipeline(transport).run(self.source, self.identity, "一次校园义卖活动")
        self.assertEqual(result["state"], "VALIDATED")
        self.assertEqual(result["review_status"], REVIEW_STATUS)
        self.assertEqual(len(transport.calls), 1)
        checkpoint = read_json(self.find_job() / "student_record.json")
        self.assertEqual(checkpoint["history"], ["NEW", "SCANNED", "IDENTIFIED", "TRANSCRIBED", "GRADED", "VALIDATED"])
        self.assertEqual(checkpoint["live_request_attempts"], 1)
        book = load_workbook(result["workbook"], read_only=True, data_only=True)
        try:
            self.assertEqual(book.sheetnames, [BLIND_SHEET, META_SHEET])
            self.assertEqual([cell.value for cell in book[BLIND_SHEET][1]], HEADERS)
            blind = dict(zip(HEADERS, [cell.value for cell in book[BLIND_SHEET][2]]))
            self.assertEqual(blind["班级"], "SYNTHETIC-CLASS")
            self.assertEqual(str(blind["班号"]), "SYNTHETIC-16")
            self.assertEqual(blind["姓名"], "合成学生")
            self.assertEqual(blind["题目"], "一次校园义卖活动")
            self.assertIsNone(blind["内容分"])
            self.assertIsNone(blind["语文与结构分"])
            self.assertIsNone(blind["总分"])
            self.assertEqual(blind["审核状态"], REVIEW_STATUS)
            self.assertEqual(blind["审题扣题_评语"], valid_output(self.identity)["grading"]["criteria"]["审题扣题"]["short_comment"])
            self.assertNotIn("下次提升重点", blind)
            schema = read_json(self.find_job() / "schema_snapshot.json")
            grade_schema = schema["properties"]["grading"]["anyOf"][0]
            self.assertEqual(set(grade_schema["properties"]), {"student_id", "student_name", "class_name", "criteria", "teacher_comment"})
        finally:
            book.close()

    def test_invalid_responses_stop_before_workbook(self):
        def missing_criterion(value):
            value["grading"]["criteria"].pop(CRITERIA[0])

        def invalid_rating(value):
            value["grading"]["criteria"][CRITERIA[0]]["rating"] = "优秀"

        def empty_comment(value):
            value["grading"]["criteria"][CRITERIA[0]]["short_comment"] = "  \n"

        def mismatch(value):
            value["grading"]["student_name"] = "其他学生"

        def over_limit(value):
            value["grading"]["criteria"][CRITERIA[0]]["short_comment"] = "长" * 91

        cases = [("malformed-json", None)] + [(name, mutate) for name, mutate in (
            ("missing-criterion", missing_criterion), ("invalid-rating", invalid_rating),
            ("empty-comment", empty_comment), ("identity-mismatch", mismatch), ("over-limit", over_limit),
        )]
        for name, mutate in cases:
            with self.subTest(name=name), local_test_directory("sol-" + name) as directory:
                root = Path(directory)
                source = root / "selected.pdf"
                source.write_bytes(self.source.read_bytes())
                value = valid_output(self.identity)
                if mutate:
                    mutate(value)
                    response = raw_response(value)
                else:
                    response = raw_response("{not json")
                transport = FakeTransport(response)
                pipeline = CalibrationPipeline(root, grader_factory=lambda: SolGrader(transport=transport))
                with self.assertRaises((ValidationError, json.JSONDecodeError)):
                    pipeline.run(source, self.identity)
                job = next((root / "jobs").glob("calibration_*"))
                checkpoint = read_json(job / "student_record.json")
                self.assertEqual(checkpoint["state"], "GRADED")
                self.assertEqual(checkpoint["live_request_attempts"], 1)
                self.assertTrue((job / "raw_model_response.json").exists())
                self.assertEqual(read_json(job / "validation_report.json")["workbook_written"], False)
                self.assertEqual(list((root / "output").rglob("*.xlsx")) if (root / "output").exists() else [], [])

    def test_timeout_is_preserved_and_never_automatically_retried(self):
        transport = FakeTransport(error=TimeoutError("simulated timeout"))
        pipeline = self.pipeline(transport)
        with self.assertRaises(TimeoutError):
            pipeline.run(self.source, self.identity)
        job = self.find_job()
        checkpoint = read_json(job / "student_record.json")
        self.assertEqual(checkpoint["state"], "TRANSCRIBED")
        self.assertEqual(checkpoint["live_request_attempts"], 1)
        self.assertFalse((job / "raw_model_response.json").exists())
        self.assertEqual(read_json(job / "validation_report.json")["status"], "MODEL_CALL_FAILED")
        self.assert_no_workbook()
        with self.assertRaises(ModelCallError):
            pipeline.run(self.source, self.identity)
        self.assertEqual(len(transport.calls), 1)

    def test_second_attempt_requires_explicit_authorization_and_preserves_attempt_count(self):
        first = FakeTransport(error=TimeoutError("first attempt failed"))
        pipeline = self.pipeline(first)
        with self.assertRaises(TimeoutError):
            pipeline.run(self.source, self.identity)
        with self.assertRaises(ModelCallError):
            pipeline.run(self.source, self.identity)
        second = FakeTransport(raw_response(valid_output(self.identity)))
        authorized = self.pipeline(second)
        result = authorized.run(self.source, self.identity, allow_second_live_attempt=True)
        self.assertEqual(result["state"], "VALIDATED")
        checkpoint = read_json(self.find_job() / "student_record.json")
        self.assertEqual(checkpoint["live_request_attempts"], 2)
        self.assertEqual(len(first.calls), 1)
        self.assertEqual(len(second.calls), 1)

    def test_materially_unreadable_preserves_reading_quality_without_fabricated_grade(self):
        value = {
            "reading_quality": {
                "complete_essay_legible": False,
                "uncertainties": [{"page_number": 3, "description": "本页大段文字无法辨认。"}],
                "materially_affects_grade": True,
            },
            "grading": None,
            "evidence": [],
        }
        result = self.pipeline(FakeTransport(raw_response(value))).run(self.source, self.identity)
        self.assertEqual(result["state"], "GRADED")
        self.assertIsNone(result["workbook"])
        job = self.find_job()
        parsed = read_json(job / "parsed_result.json")
        self.assertIsNone(parsed["grading"])
        self.assertEqual(parsed["reading_quality"]["uncertainties"][0]["page_number"], 3)
        self.assertEqual(read_json(job / "validation_report.json")["status"], "MATERIAL_READING_UNCERTAINTY")
        self.assert_no_workbook()

    def test_request_contains_only_selected_pdf_and_allowed_context_with_no_tools(self):
        marker = "OTHER_STUDENT_SECRET_MARKER"
        (self.root / "other_student.pdf").write_text(marker, encoding="utf-8")
        transport = FakeTransport(raw_response(valid_output(self.identity)))
        self.pipeline(transport).run(self.source, self.identity)
        payload = transport.calls[0][0]
        self.assertEqual(set(payload), {"model", "reasoning", "store", "input", "text", "metadata"})
        self.assertEqual(payload["model"], "gpt-5.6-sol")
        self.assertIs(payload["store"], False)
        serialized = json.dumps(payload, ensure_ascii=False)
        for forbidden in (marker, "other_student.pdf", "历史教师", "优秀", "达标", "previous_response_id", "conversation", "tools"):
            self.assertNotIn(forbidden, serialized)
        content = payload["input"][0]["content"]
        self.assertEqual([item["type"] for item in content], ["input_text", "input_file"])
        encoded = content[1]["file_data"].split(",", 1)[1]
        self.assertEqual(base64.b64decode(encoded), self.source.read_bytes())
        request = self.find_job() / "grading_input_manifest.json"
        manifest = read_json(request)
        self.assertEqual(manifest["reference_version"], "reference-example-v1")
        self.assertIn("reference_example.md", manifest["snapshot_sha256"])
        prompt = (self.find_job() / "prompt_snapshot.txt").read_text(encoding="utf-8")
        self.assertIn("reference_example.md", prompt)
        self.assertIn("不得复制或套用范例", prompt)
        self.assertEqual(manifest["source_sha256"], sha256_file(self.source))
        for name, digest in manifest["snapshot_sha256"].items():
            self.assertEqual(digest, sha256_file(self.find_job() / name))

    def test_credential_preflight_blocks_without_consuming_attempt(self):
        pipeline = CalibrationPipeline(self.root, grader_factory=lambda: SolGrader(transport=ResponsesHttpTransport(api_key="")))
        with self.assertRaises(CredentialUnavailable):
            pipeline.run(self.source, self.identity)
        job = self.find_job()
        checkpoint = read_json(job / "student_record.json")
        self.assertEqual(checkpoint["state"], "TRANSCRIBED")
        self.assertEqual(checkpoint["live_request_attempts"], 0)
        self.assertEqual(read_json(job / "validation_report.json")["model_request_attempted"], False)
        self.assert_no_workbook()

    def test_source_and_configuration_snapshots_are_immutable_on_resume(self):
        pipeline = self.pipeline(FakeTransport(raw_response(valid_output(self.identity))))
        original_bytes = self.source.read_bytes()
        job, _, checkpoint = pipeline.prepare(self.source, self.identity)
        self.assertEqual(checkpoint["state"], "TRANSCRIBED")
        isolated_hash = sha256_file(job / "source.pdf")
        replacement = PdfWriter()
        for _ in range(5):
            replacement.add_blank_page(width=612, height=792)
        replacement.write(self.source)
        replacement_job, _, _ = pipeline.prepare(self.source, self.identity)
        self.assertNotEqual(replacement_job, job)
        self.assertEqual(sha256_file(job / "source.pdf"), isolated_hash)
        self.source.write_bytes(original_bytes)
        snapshot = job / "criteria_snapshot.json"
        snapshot.write_text("{}\n", encoding="utf-8")
        with self.assertRaises(ValidationError):
            pipeline.prepare(self.source, self.identity)


if __name__ == "__main__":
    unittest.main()
