"""AI-suggested section scores and question number flow into the teacher workbook."""
from pathlib import Path
import unittest

from openpyxl import load_workbook

from excel.schema import SHEET
from excel.workbook import ExcelStore
from grading.calibration_schema import CalibrationValidator, response_json_schema
from grading.grader import GradingInput
from grading.mock_grader import MockGrader
from grading.schemas import Identity, ValidationError, Validator
from tests.local_temp import local_test_directory
from tests import test_real_batch_integration as batch_fixture
from tests.test_sol_calibration import valid_output
from workflow.storage import read_json

TEXT_LIMITS = read_json(Path(__file__).resolve().parents[1] / "config" / "text_limits.json")


def scored(value, question_number="Q2", content=22, language=21):
    return {**value, "question_number": question_number,
            "scores": {"content_score": content, "language_structure_score": language}}


class ScoreValidationTests(unittest.TestCase):
    def setUp(self):
        self.identity = Identity("SYNTHETIC-16", "合成学生", "SYNTHETIC-CLASS")
        self.validator = CalibrationValidator(TEXT_LIMITS, page_count=4)

    def test_schema_requires_question_number_and_two_scores_only(self):
        schema = response_json_schema(TEXT_LIMITS)
        self.assertIn("question_number", schema["required"])
        self.assertIn("scores", schema["required"])
        score_schema = schema["properties"]["scores"]["anyOf"][0]
        self.assertEqual(set(score_schema["properties"]), {"content_score", "language_structure_score"})

    def test_valid_scores_and_question_number_are_kept(self):
        result = self.validator.validate(scored(valid_output(self.identity)), self.identity)
        self.assertEqual(result.question_number, "Q2")
        self.assertEqual(result.scores, {"content_score": 22, "language_structure_score": 21})

    def test_question_number_may_be_null_when_not_found(self):
        result = self.validator.validate(scored(valid_output(self.identity), question_number=None), self.identity)
        self.assertIsNone(result.question_number)

    def test_invalid_scores_and_question_numbers_are_rejected(self):
        for label, value in {
            "above-max": scored(valid_output(self.identity), content=31),
            "negative": scored(valid_output(self.identity), language=-1),
            "half-mark": scored(valid_output(self.identity), content=22.5),
            "boolean": scored(valid_output(self.identity), content=True),
            "free-text-question": scored(valid_output(self.identity), question_number="第二题"),
            "lowercase-question": scored(valid_output(self.identity), question_number="q2"),
            "missing-scores": {**scored(valid_output(self.identity)), "scores": None},
            "total-supplied": {**scored(valid_output(self.identity)), "scores": {
                "content_score": 22, "language_structure_score": 21, "total_score": 43}},
        }.items():
            with self.subTest(label), self.assertRaises(ValidationError):
                self.validator.validate(value, self.identity)

    def test_unreadable_essay_must_not_carry_scores(self):
        value = scored(valid_output(self.identity))
        value["reading_quality"]["materially_affects_grade"] = True
        value["grading"] = None
        with self.assertRaises(ValidationError):
            self.validator.validate(value, self.identity)
        value["scores"] = None
        self.assertIsNone(self.validator.validate(value, self.identity).grading)

    def test_response_from_older_schema_snapshot_still_validates(self):
        result = self.validator.validate(valid_output(self.identity), self.identity)
        self.assertIsNone(result.scores)


class TotalFormulaTests(unittest.TestCase):
    def setUp(self):
        self.temp = local_test_directory("ai-score-formula")
        self.root = Path(self.temp.__enter__())
        self.addCleanup(self.temp.__exit__, None, None, None)
        self.identity = Identity("016", "测试学生", "MOCK-207")
        self.validator = Validator()
        self.result = self.validator.validate(MockGrader().grade(GradingInput(self.identity, "测试作文")), self.identity)
        self.store = ExcelStore(self.root / "results.xlsx", self.validator)
        self.store.ensure_draft(self.result, self.identity, "job", "digest", "Q2",
                                {"content_score": 22, "language_structure_score": 21})

    def test_total_is_written_as_formula_and_read_as_sum(self):
        book = load_workbook(self.store.path)
        try:
            self.assertEqual(book[SHEET]["G2"].value, "=E2+F2")
        finally:
            book.close()
        record = self.store.get("job", self.identity)
        self.assertEqual((record.topic, record.content_score, record.language_structure_score, record.total_score),
                         ("Q2", 22, 21, 43))

    def test_teacher_score_edit_updates_total(self):
        book = load_workbook(self.store.path)
        book[SHEET]["E2"] = 25
        book.save(self.store.path)
        book.close()
        self.assertEqual(self.store.get("job", self.identity).total_score, 46)

    def test_formulas_outside_total_or_other_sums_are_rejected(self):
        for cell, formula in (("G2", "=E2+F2+1"), ("E2", "=10+12")):
            with self.subTest(cell=cell, formula=formula):
                book = load_workbook(self.store.path)
                book[SHEET]["E2"], book[SHEET]["G2"] = 22, "=E2+F2"
                book[SHEET][cell] = formula
                book.save(self.store.path)
                book.close()
                with self.assertRaises(ValidationError):
                    self.store.get("job", self.identity)


class ScoredTransport:
    def __init__(self, identities):
        self.identities = identities
        self.calls = []

    def ensure_ready(self):
        return None

    def __call__(self, payload, timeout_seconds):
        self.calls.append(payload)
        prompt = payload["input"][0]["content"][0]["text"]
        identity = next(item for item in self.identities if item.student_name in prompt)
        return batch_fixture.raw_response(scored(batch_fixture.valid_response(identity)))


class ScoredBatchTests(unittest.TestCase):
    setUp = batch_fixture.RealBatchIntegrationTests.setUp
    pipeline_factory = batch_fixture.RealBatchIntegrationTests.pipeline_factory
    make_pipeline = batch_fixture.RealBatchIntegrationTests.make_pipeline

    def test_ai_scores_and_question_number_reach_teacher_workbook(self):
        pipeline = self.make_pipeline(ScoredTransport(self.identities))
        result = pipeline.run_real_batch(self.records, self.source_dir, limit=1, model_profile="luna_xhigh")
        self.assertEqual(result["failed"], [])
        book = load_workbook(result["workbook"])
        try:
            row = [cell.value for cell in book[SHEET][2]]
        finally:
            book.close()
        self.assertEqual(row[3:7], ["Q2", 22, 21, "=E2+F2"])


if __name__ == "__main__":
    unittest.main()
