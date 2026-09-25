import copy
from dataclasses import replace
import hashlib
import json
import os
from pathlib import Path
import unittest

from openpyxl import load_workbook

from excel.workbook import ExcelStore, roster_text
from grading.schemas import CRITERIA, RATINGS, CriterionResult, Identity, ROOT, Validator
from rendering.renderer import (
    LAYOUT_PATH,
    REQUIRED_FIELDS,
    PillowRenderer,
    RenderLayoutError,
    load_layout,
)
from tests.local_temp import local_test_directory
from tests.local_fixtures import require_local_fixtures
from workflow.storage import atomic_json


MOCK_RENDER_WORKBOOK = Path(os.environ.get(
    "PHASE1_RENDER_MOCK_WORKBOOK", ROOT / "output/feedback_excel_v2_verification/results.xlsx"
))
TEACHER_RENDER_WORKBOOK = Path(os.environ.get(
    "PHASE1_RENDER_TEACHER_WORKBOOK", ROOT / "output/local_render_regression/teacher_summary.xlsx"
))


class RenderingTests(unittest.TestCase):
    def setUp(self):
        self.layout = load_layout()
    @property
    def record(self):
        if not hasattr(self, "_record_cache"):
            workbook_path = MOCK_RENDER_WORKBOOK
            require_local_fixtures(self, [workbook_path])
            book = load_workbook(workbook_path, read_only=True, data_only=True)
            try:
                rows = book["Phase1批改结果"].iter_rows(min_row=2, max_row=2, values_only=True)
                audit_rows = book["Pipeline审计"].iter_rows(min_row=2, max_row=2, values_only=True)
                values = next(rows, None)
                audit = next(audit_rows, None)
                if values is None or audit is None:
                    self.skipTest("local regression fixture not installed")
                identity = Identity(roster_text(values[1]), roster_text(values[2]), roster_text(values[0]))
                job_id = audit[0]
            finally:
                book.close()
            store = ExcelStore(workbook_path, Validator())
            self._record_cache = store.get(job_id, identity, approved=True)
        return self._record_cache

    @property
    def real_record(self):
        if not hasattr(self, "_real_record_cache"):
            workbook_path = TEACHER_RENDER_WORKBOOK
            require_local_fixtures(self, [workbook_path])
            book = load_workbook(workbook_path, read_only=True, data_only=True)
            try:
                values = next(
                    book["Phase1批改结果"].iter_rows(min_row=2, max_row=2, values_only=True),
                    None,
                )
                if values is None:
                    self.skipTest("local regression fixture not installed")
                identity = Identity(roster_text(values[1]), roster_text(values[2]), roster_text(values[0]))
            finally:
                book.close()
            store = ExcelStore(workbook_path, Validator())
            self._real_record_cache, _audit = store.get_render_source("local-regression-fixture", identity)
        return self._real_record_cache

    def test_master_and_layout_are_pinned(self):
        master = Path(self.layout["_master_path"])
        self.assertEqual(self.layout["layout_version"], "layout-v1.2")
        self.assertEqual((self.layout["canvas"]["width"], self.layout["canvas"]["height"]), (1024, 1536))
        self.assertEqual(self.layout["master"]["sha256"], hashlib.sha256(master.read_bytes()).hexdigest())
        self.assertEqual(set(self.layout["fields"]), REQUIRED_FIELDS)

    def test_layout_v1_2_preserves_prior_fixed_geometry(self):
        prior = json.loads((ROOT / "rendering/template/layout_v1_1.json").read_text(encoding="utf-8"))
        for name in REQUIRED_FIELDS:
            self.assertEqual(self.layout["fields"][name]["box"], prior["fields"][name]["box"], name)
        self.assertEqual(self.layout["master"], prior["master"])

    def test_approved_excel_values_reach_receipt_unchanged(self):
        with local_test_directory("render-approved") as directory:
            receipt_path = PillowRenderer().render(self.record, Path(directory), {"excel_row": 2})
            receipt = json.loads(Path(receipt_path).read_text(encoding="utf-8"))
            self.assertEqual(receipt["approved_excel_values"], self.record.to_dict())
            self.assertEqual(receipt["source"]["excel_row"], 2)
            book = load_workbook(MOCK_RENDER_WORKBOOK, read_only=True, data_only=True)
            try:
                self.assertIsNone(book["Phase1批改结果"]["D2"].value)
            finally:
                book.close()
            self.assertEqual(self.record.topic, "")
            self.assertEqual(receipt["fields"]["topic"]["source_value"], "")
            self.assertTrue(all(item["fits"] for item in receipt["fields"].values()))
            self.assertEqual(receipt["font"]["family"], "Noto Sans SC")
            self.assertEqual(receipt["font"]["weight"], 600)
            self.assertEqual(receipt["font"]["rating_weight"], 700)

    def test_criterion_comment_blocks_are_centered_from_their_own_boxes(self):
        with local_test_directory("render-centering") as directory:
            receipt_path = PillowRenderer().render(self.record, Path(directory))
            receipt = json.loads(Path(receipt_path).read_text(encoding="utf-8"))
            top_spaces = []
            bottom_spaces = []
            for criterion in self.layout["criteria_order"]:
                name = f"criteria.{criterion}.short_comment"
                spec = self.layout["fields"][name]
                box_x, box_y, box_w, box_h = spec["box"]
                pad_left, pad_top, pad_right, pad_bottom = spec["padding"]
                available_height = box_h - pad_top - pad_bottom
                field = receipt["fields"][name]
                expected_top = box_y + pad_top + (available_height - field["block_height"]) // 2
                self.assertEqual(field["block_top"], expected_top, name)
                self.assertEqual(field["block_height"], field["line_count"] * spec["line_height"], name)
                self.assertEqual(field["block_bottom"], field["block_top"] + field["block_height"], name)
                top_spaces.append(field["free_space_top"])
                bottom_spaces.append(field["free_space_bottom"])
            self.assertLessEqual(max(top_spaces) - min(top_spaces), 1)
            self.assertLessEqual(max(bottom_spaces) - min(bottom_spaces), 1)

    def test_rating_controls_all_legal_states_are_centered_and_clear_of_arrow(self):
        renderer = PillowRenderer()
        with local_test_directory("render-rating-states") as directory:
            for rating in RATINGS:
                criteria = {
                    name: CriterionResult(rating, self.record.criteria[name].short_comment)
                    for name in CRITERIA
                }
                record = replace(self.record, criteria=criteria)
                output = Path(directory) / rating
                receipt_path = renderer.render(record, output)
                receipt = json.loads(Path(receipt_path).read_text(encoding="utf-8"))
                for criterion in CRITERIA:
                    name = f"criteria.{criterion}.rating"
                    spec = self.layout["fields"][name]
                    field = receipt["fields"][name]
                    self.assertEqual(field["source_value"], rating, name)
                    self.assertEqual(field["line_count"], 1, name)
                    self.assertEqual(field["font_weight"], 700, name)
                    self.assertGreaterEqual(field["arrow_clearance"], self.layout["rating_control"]["min_arrow_clearance"], name)
                    box_x, box_y, box_w, box_h = spec["box"]
                    expected_x = box_x + (box_w - field["line_widths"][0]) / 2
                    self.assertEqual(field["line_positions"][0]["x"], expected_x, name)
                    self.assertLessEqual(abs(field["free_space_top"] - field["free_space_bottom"]), 1, name)

    def test_top_values_and_scores_are_glyph_centered_from_their_own_boxes(self):
        names = (
            "class_name", "student_id", "student_name", "topic",
            "content_score", "language_structure_score", "total_score",
        )
        with local_test_directory("render-top-glyph-centering") as directory:
            receipt_path = PillowRenderer().render(self.real_record, Path(directory))
            receipt = json.loads(Path(receipt_path).read_text(encoding="utf-8"))
            for name in names:
                field = receipt["fields"][name]
                self.assertEqual(field["vertical_alignment"], "glyph_center", name)
                self.assertEqual(field["line_count"], 1, name)
                self.assertLessEqual(
                    abs(field["visual_free_space_top"] - field["visual_free_space_bottom"]), 1,
                    name,
                )

    def test_real_teacher_summary_has_positive_v1_2_safety_margin(self):
        self.assertEqual(len(self.real_record.teacher_comment), 172)
        with local_test_directory("render-real-teacher-margin") as directory:
            receipt_path = PillowRenderer().render(self.real_record, Path(directory))
            receipt = json.loads(Path(receipt_path).read_text(encoding="utf-8"))
            field = receipt["fields"]["teacher_comment"]
            self.assertEqual(field["source_value"], self.real_record.teacher_comment)
            self.assertEqual(field["rendered_text"], self.real_record.teacher_comment)
            self.assertEqual(field["font_size"], 20)
            self.assertEqual(field["line_height"], 24)
            self.assertLessEqual(field["line_count"], 5)
            self.assertGreater(field["free_space_top"], 0)
            self.assertGreater(field["free_space_bottom"], 0)
            self.assertGreater(field["visual_free_space_top"], 0)
            self.assertGreater(field["visual_free_space_bottom"], 0)

    def test_explicit_capacity_tiers_cover_criterion_and_teacher_limits(self):
        long_comment = "测" * 90
        criteria = {
            name: CriterionResult(self.record.criteria[name].rating, long_comment)
            for name in CRITERIA
        }
        teacher_comment = "教" * 160
        record = replace(self.record, criteria=criteria, teacher_comment=teacher_comment)
        with local_test_directory("render-capacity-tiers") as directory:
            receipt_path = PillowRenderer().render(record, Path(directory))
            receipt = json.loads(Path(receipt_path).read_text(encoding="utf-8"))
            for criterion in CRITERIA:
                field = receipt["fields"][f"criteria.{criterion}.short_comment"]
                self.assertEqual(field["source_value"], long_comment)
                self.assertEqual(field["rendered_text"], long_comment)
                self.assertEqual(field["font_tier"], 1)
                self.assertLessEqual(field["line_count"], 3)
                self.assertTrue(field["fits"])
            field = receipt["fields"]["teacher_comment"]
            self.assertEqual(field["source_value"], teacher_comment)
            self.assertEqual(field["rendered_text"], teacher_comment)
            self.assertEqual(field["font_tier"], 0)
            self.assertEqual(field["font_size"], 20)
            self.assertEqual(field["line_height"], 24)
            self.assertLessEqual(field["line_count"], 5)
            self.assertGreater(field["free_space_top"], 0)
            self.assertGreater(field["free_space_bottom"], 0)
            self.assertTrue(field["fits"])

    def test_same_approved_input_is_pixel_deterministic(self):
        with local_test_directory("render-deterministic") as directory:
            first = Path(directory) / "first"
            second = Path(directory) / "second"
            PillowRenderer().render(self.record, first)
            PillowRenderer().render(self.record, second)
            self.assertEqual((first / "student_card.png").read_bytes(), (second / "student_card.png").read_bytes())

    def test_unapproved_record_cannot_render(self):
        with local_test_directory("render-unapproved") as directory:
            with self.assertRaises(RenderLayoutError):
                PillowRenderer().render(replace(self.record, review_status="PENDING"), Path(directory))
            self.assertFalse((Path(directory) / "student_card.png").exists())

    def test_overflow_fails_without_truncation(self):
        with local_test_directory("render-overflow") as directory:
            layout = copy.deepcopy(self.layout)
            layout.pop("_path", None)
            layout.pop("_master_path", None)
            layout.pop("_font_path", None)
            layout["fields"]["teacher_comment"]["max_lines"] = 1
            layout["fields"]["teacher_comment"].pop("font_tiers", None)
            layout_path = Path(directory) / "layout_overflow.json"
            atomic_json(layout_path, layout)
            with self.assertRaises(RenderLayoutError):
                PillowRenderer(layout_path).render(self.record, Path(directory) / "output")
            self.assertFalse((Path(directory) / "output" / "student_card.png").exists())


if __name__ == "__main__":
    unittest.main()
