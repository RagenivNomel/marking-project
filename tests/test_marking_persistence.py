"""The marking persistence rule: an essay is done only when the task workbook
holds its valid result row and audit row. Anything short of that is not done,
and its app-generated working state is discarded before it is marked again."""
import copy
import csv
import json
from pathlib import Path
import threading
import unittest

from openpyxl import Workbook, load_workbook
from pypdf import PdfReader, PdfWriter

from application import Action, AssignmentSource, WorkflowController
from excel.schema import HEADERS, SHEET
from grading.schemas import Identity, ROOT
from grading.sol_grader import SolGrader
from tests.local_temp import local_test_directory
from tests.test_real_batch_integration import raw_response, valid_response, working_state
from workflow.calibration_pipeline import CalibrationPipeline
from workflow.pipeline import Pipeline, job_key
from workflow.storage import exclusive_lock


CONFIG_DIR = ROOT / "config"


class ScriptedTransport:
    """Offline model transport; behaviour is scripted per student name."""

    def __init__(self, identities, *, fail_once=(), unreadable_once=()):
        self.identities = list(identities)
        self.fail_once = set(fail_once)
        self.unreadable_once = set(unreadable_once)
        self.calls = []
        self.lock = threading.Lock()

    def ensure_ready(self):
        return None

    def __call__(self, payload, timeout_seconds):
        prompt = payload["input"][0]["content"][0]["text"]
        identity = next(item for item in self.identities if item.student_name in prompt)
        with self.lock:
            self.calls.append(identity.student_id)
            fail = identity.student_name in self.fail_once
            self.fail_once.discard(identity.student_name)
            unreadable = identity.student_name in self.unreadable_once
            self.unreadable_once.discard(identity.student_name)
        if fail:
            raise RuntimeError("synthetic model failure")
        response = valid_response(identity)
        if unreadable:
            response = copy.deepcopy(response)
            response["reading_quality"] = {
                "complete_essay_legible": False,
                "uncertainties": ["第二页大部分字迹无法辨认"],
                "materially_affects_grade": True,
            }
        return raw_response(response)


def grader_pipeline_factory(transport):
    profiles = json.loads((CONFIG_DIR / "calibration_model.json").read_text(encoding="utf-8"))["profiles"]

    def factory(project_dir, model_profile=None):
        profile = profiles[model_profile]
        return CalibrationPipeline(
            project_dir,
            config_dir=CONFIG_DIR,
            model_profile=model_profile,
            grader_factory=lambda: SolGrader(
                model=profile["model"],
                reasoning_effort=profile["reasoning_effort"],
                transport=transport,
            ),
        )
    return factory


def write_legacy_attempt(essay_dir, identity, *, workbook, child=None, state="VALIDATED"):
    """Recreate the partial-recovery files the old pipeline left behind."""
    essay_dir.mkdir(parents=True, exist_ok=True)
    key = job_key(identity)
    checkpoint = {
        "schema_version": 2, "mode": "REAL_PDF_BATCH", "job_id": key,
        "identity": identity.to_dict(), "input_digest": "0" * 64, "source_sha256": "1" * 64,
        "source_page_count": 1, "model_profile": "luna_xhigh", "execution_namespace": "old",
        "workbook": str(workbook), "state": state,
        "history": ["NEW", "SCANNED", "IDENTIFIED", "TRANSCRIBED"], "last_error": None,
    }
    (essay_dir / "student_record.json").write_text(json.dumps(checkpoint), encoding="utf-8")
    (essay_dir / "validated_result.json").write_text("{\"partial\": true}", encoding="utf-8")
    (essay_dir / "workbook_persistence.json").write_text(
        json.dumps({"status": "PENDING", "workbook": str(workbook)}), encoding="utf-8")
    (essay_dir / "real_grading_attempts.json").write_text(
        json.dumps({"attempts": [{"attempt": 1, "state": "IN_PROGRESS"}]}), encoding="utf-8")
    (essay_dir / "real_grading_bridge.json").write_text(json.dumps({
        "calibration_job": str(child or essay_dir / "missing-child"),
        "execution_namespace": "old", "status": "IN_PROGRESS",
    }), encoding="utf-8")
    # An interrupted model call inside the working folder: the old
    # single-attempt guard would refuse to call the model again.
    stale = essay_dir / "work" / "jobs" / f"calibration_{key}"
    stale.mkdir(parents=True)
    (stale / "student_record.json").write_text(json.dumps({
        **checkpoint, "state": "GRADED", "live_request_attempts": 1,
    }), encoding="utf-8")
    (stale / "raw_model_response.json").write_text("{\"truncated\":", encoding="utf-8")


class MarkingPersistenceTests(unittest.TestCase):
    def setUp(self):
        self.temp = local_test_directory("marking-persistence")
        self.root = Path(self.temp.__enter__())
        self.addCleanup(self.temp.__exit__, None, None, None)
        self.source_dir = self.root / "essays"
        self.source_dir.mkdir()
        self.results_path = self.root / "Results" / "results.xlsx"
        self.batch_id = "persistence"
        self.jobs = self.root / "jobs" / self.batch_id
        self.identities = [Identity(f"{index:02d}", f"持久学生{index:02d}", "207") for index in range(1, 4)]
        self.records = []
        for sequence, identity in enumerate(self.identities, 1):
            name = f"submission_{sequence:03}.pdf"
            writer = PdfWriter()
            writer.add_blank_page(width=612, height=792)
            writer.add_metadata({"/StudentMarker": identity.student_name})
            with (self.source_dir / name).open("wb") as stream:
                writer.write(stream)
            self.records.append({"sequence": sequence, "source_pdf": name, **identity.to_dict(),
                                 "match_status": "STRONG_ROSTER_MATCH"})

    def pipeline(self, transport):
        return Pipeline(self.root, self.batch_id, config_dir=CONFIG_DIR, allow_real=True,
                        workbook_path=self.results_path,
                        real_pipeline_factory=grader_pipeline_factory(transport))

    def run_batch(self, transport, records=None, **kwargs):
        return self.pipeline(transport).run_real_batch(
            self.records if records is None else records, self.source_dir,
            model_profile="luna_xhigh", **kwargs)

    def essay_dir(self, identity):
        return self.jobs / job_key(identity)

    def rows(self):
        book = load_workbook(self.results_path, read_only=True, data_only=True)
        try:
            return list(book[SHEET].iter_rows(min_row=2, values_only=True))
        finally:
            book.close()

    def test_completed_essays_are_saved_and_leave_no_working_state(self):
        transport = ScriptedTransport(self.identities)
        result = self.run_batch(transport)
        self.assertEqual([item["identity"]["student_id"] for item in result["validated"]], ["01", "02", "03"])
        self.assertEqual(result["failed"], [])
        self.assertEqual([row[1] for row in self.rows()], ["01", "02", "03"])
        for identity in self.identities:
            self.assertEqual(working_state(self.essay_dir(identity)), [], identity.student_id)

    def test_failed_model_call_is_not_done_and_leaves_nothing_to_resume(self):
        transport = ScriptedTransport(self.identities, fail_once={self.identities[1].student_name})
        first = self.run_batch(transport)
        self.assertEqual([item["identity"]["student_id"] for item in first["failed"]], ["02"])
        self.assertEqual([row[1] for row in self.rows()], ["01", "03"])
        self.assertEqual(working_state(self.essay_dir(self.identities[1])), [])

        second = self.run_batch(transport)
        self.assertEqual(second["already_complete"], 2)
        self.assertEqual([item["identity"]["student_id"] for item in second["validated"]], ["02"])
        self.assertEqual(transport.calls.count("02"), 2)
        self.assertEqual(transport.calls.count("01"), 1)
        self.assertEqual([row[1] for row in self.rows()], ["01", "02", "03"])

    def test_unreadable_essay_is_not_done_and_is_marked_again_next_run(self):
        transport = ScriptedTransport(self.identities, unreadable_once={self.identities[0].student_name})
        first = self.run_batch(transport)
        self.assertEqual([item["identity"]["student_id"] for item in first["failed"]], ["01"])
        self.assertNotIn("01", [row[1] for row in self.rows()])
        self.assertEqual(working_state(self.essay_dir(self.identities[0])), [])

        second = self.run_batch(transport)
        self.assertEqual([item["identity"]["student_id"] for item in second["validated"]], ["01"])
        self.assertIn("01", [row[1] for row in self.rows()])

    def test_interrupted_attempt_is_discarded_and_rerun_from_scratch(self):
        target = self.identities[0]
        essay_dir = self.essay_dir(target)
        write_legacy_attempt(essay_dir, target, workbook=self.root / "elsewhere" / "results.xlsx")
        (essay_dir / "output").mkdir()
        (essay_dir / "output" / "card-evidence.txt").write_text("stage 3c", encoding="utf-8")
        outside = self.root / "jobs" / f"calibration_{self.batch_id}_{job_key(target)}"
        outside.mkdir(parents=True)
        (outside / "student_record.json").write_text("{}", encoding="utf-8")

        transport = ScriptedTransport(self.identities)
        result = self.run_batch(transport)

        self.assertEqual(result["failed"], [])
        self.assertEqual(sorted(transport.calls), ["01", "02", "03"])
        self.assertEqual([row[1] for row in self.rows()], ["01", "02", "03"])
        self.assertEqual(working_state(essay_dir), [])
        # Only this essay's own folder is recovery scope; card output and
        # anything outside the task's essay folder are left alone.
        self.assertEqual((essay_dir / "output" / "card-evidence.txt").read_text(encoding="utf-8"), "stage 3c")
        self.assertTrue((outside / "student_record.json").is_file())
        log = (self.jobs / "discarded_attempts.log").read_text(encoding="utf-8")
        self.assertIn(job_key(target), log)

    def test_saved_row_is_never_regraded_or_changed(self):
        transport = ScriptedTransport(self.identities)
        self.run_batch(transport, records=self.records[:1])
        book = load_workbook(self.results_path)
        try:
            sheet = book[SHEET]
            sheet.cell(2, HEADERS.index("教师总评") + 1).value = "教师修改后的总评"
            sheet.cell(2, HEADERS.index("总分") + 1).value = 36
            sheet.cell(2, HEADERS.index("审核状态") + 1).value = "APPROVED"
            book.save(self.results_path)
        finally:
            book.close()
        edited = self.rows()[0]
        # Leftover or corrupt working files never make a saved essay undone.
        essay_dir = self.essay_dir(self.identities[0])
        write_legacy_attempt(essay_dir, self.identities[0], workbook=self.root / "old.xlsx")
        (essay_dir / "student_record.json").write_text("{", encoding="utf-8")
        transport.calls.clear()

        result = self.run_batch(transport)

        self.assertEqual(result["already_complete"], 1)
        self.assertNotIn("01", transport.calls)
        rows = self.rows()
        self.assertEqual(rows[0], edited)
        self.assertEqual(rows[0][HEADERS.index("审核状态")], "APPROVED")
        self.assertEqual([row[1] for row in rows], ["01", "02", "03"])

    def test_invalid_existing_row_is_not_done_but_is_never_overwritten(self):
        transport = ScriptedTransport(self.identities)
        self.run_batch(transport, records=self.records[:1])
        book = load_workbook(self.results_path)
        try:
            book[SHEET].cell(2, HEADERS.index("审核状态") + 1).value = "MAYBE"
            book.save(self.results_path)
        finally:
            book.close()
        damaged = self.rows()[0]
        result = self.run_batch(transport, records=self.records[:1])
        self.assertEqual(result["already_complete"], 0)
        self.assertEqual([item["identity"]["student_id"] for item in result["failed"]], ["01"])
        self.assertEqual(self.rows(), [damaged])

    def test_discard_log_is_not_read_back_as_recovery_state(self):
        self.jobs.mkdir(parents=True)
        (self.jobs / "discarded_attempts.log").write_text("\x00 not a log {", encoding="utf-8")
        result = self.run_batch(ScriptedTransport(self.identities))
        self.assertEqual(len(result["validated"]), 3)

    def test_task_lock_keeps_a_second_run_out_until_the_first_finishes(self):
        transport = ScriptedTransport(self.identities)
        outcome = {}
        with exclusive_lock(self.jobs / ".real-batch.lock"):
            worker = threading.Thread(target=lambda: outcome.update(result=self.run_batch(transport)))
            worker.start()
            worker.join(0.5)
            self.assertTrue(worker.is_alive())
            self.assertEqual(transport.calls, [])
        worker.join(30)
        self.assertEqual(len(outcome["result"]["validated"]), 3)


class Pile1ReproductionTests(unittest.TestCase):
    """Pile 1: every essay had old checkpoints bound to a workbook path that no
    longer matches the task, plus IN_PROGRESS bridges into old calibration
    folders. Marking silently rejected all of them in a few milliseconds."""

    def setUp(self):
        self.temp = local_test_directory("pile1-reproduction")
        self.root = Path(self.temp.__enter__())
        self.addCleanup(self.temp.__exit__, None, None, None)
        self.pile = self.root / "pile1_test"
        self.pile.mkdir()
        self.identities = [Identity(f"{n}", f"一号堆学生{n}", "207") for n in (1, 2, 3)]
        continuous = PdfWriter()
        for index, _identity in enumerate(self.identities):
            # Rotation gives each blank page a distinct page digest.
            continuous.add_blank_page(width=612, height=792).rotate(90 * index)
        with (self.pile / "_continuous.pdf").open("wb") as stream:
            continuous.write(stream)
        rows, decisions = [], []
        for index, page in enumerate(PdfReader(self.pile / "_continuous.pdf").pages):
            name = f"submission_{index + 1:03}.pdf"
            writer = PdfWriter()
            writer.add_page(page)
            with (self.pile / name).open("wb") as stream:
                writer.write(stream)
            identity = self.identities[index]
            rows.append({"start_page_idx": index, "pages": 1, "raw_ocr_name": "", "name_preview": ""})
            decisions.append({"source_pdf": name, "class_name": identity.class_name,
                              "student_id": identity.student_id,
                              "match_status": "STRONG_ROSTER_MATCH", "evidence": "fixture"})
        with (self.pile / "_manifest.csv").open("w", newline="", encoding="utf-8") as stream:
            writer = csv.DictWriter(stream, fieldnames=("start_page_idx", "pages", "raw_ocr_name", "name_preview"))
            writer.writeheader()
            writer.writerows(rows)
        self.roster = self.root / "roster.xlsx"
        book = Workbook()
        sheet = book.active
        sheet.title = "作文诊断输入"
        sheet.append(["班号", "学生姓名", "班级"])
        for identity in self.identities:
            sheet.append([identity.student_id, identity.student_name, identity.class_name])
        book.save(self.roster)
        book.close()
        self.decisions = self.root / "pile1_identity_decisions.json"
        self.decisions.write_text(json.dumps(decisions, ensure_ascii=False), encoding="utf-8")
        self.source = AssignmentSource(split_pile=self.pile, roster=self.roster,
                                       identity_decisions=self.decisions, config_dir=CONFIG_DIR)
        self.transport = ScriptedTransport(self.identities)

        def pipeline_factory(project_dir, batch_id, config_dir, *, workbook_path=None):
            return Pipeline(project_dir, batch_id, config_dir=config_dir, allow_real=True,
                            workbook_path=workbook_path,
                            real_pipeline_factory=grader_pipeline_factory(self.transport))
        self.controller = WorkflowController(project_dir=self.root, pipeline_factory=pipeline_factory)

        # The stale state found in the real pile 1 task.
        self.task = self.controller.workflow.marking_batch_id(self.source)
        self.task_jobs = self.root / "jobs" / self.task
        self.old_workbook = self.root / "output" / self.task / "results.xlsx"
        self.old_children = []
        for identity in self.identities:
            child = self.root / "jobs" / f"calibration_{self.task}_{job_key(identity)}"
            child.mkdir(parents=True)
            (child / "student_record.json").write_text(json.dumps({"state": "TRANSCRIBED",
                                                                   "live_request_attempts": 1}),
                                                       encoding="utf-8")
            self.old_children.append(child)
            write_legacy_attempt(self.task_jobs / job_key(identity), identity,
                                 workbook=self.old_workbook, child=child, state="TRANSCRIBED")
        (self.task_jobs / ".real-batch.lock").write_text("0", encoding="utf-8")

    def test_stale_attempts_do_not_block_or_hide_marking(self):
        before = self.controller.inspect(self.source)
        self.assertEqual(before.summary["blocking_errors"], 0)
        for item in before.submissions:
            self.assertIn(Action.RUN_MARKING, item.next_actions, item.submission_id)
            self.assertEqual(item.attention, (), item.submission_id)

    def test_every_essay_reaches_the_grader_and_is_saved(self):
        outcome = self.controller.execute(Action.RUN_MARKING, self.source)
        result = outcome["result"]

        self.assertEqual(result["failed"], [])
        self.assertEqual(sorted(self.transport.calls), ["1", "2", "3"])
        self.assertEqual(len(result["validated"]), 3)
        workbook = outcome["source"].workbook
        self.assertIsNotNone(workbook)
        self.assertNotEqual(Path(workbook).resolve(), self.old_workbook.resolve())
        book = load_workbook(workbook, read_only=True, data_only=True)
        try:
            ids = sorted(row[1] for row in book[SHEET].iter_rows(min_row=2, values_only=True))
        finally:
            book.close()
        self.assertEqual(ids, ["1", "2", "3"])
        for identity in self.identities:
            self.assertEqual(working_state(self.task_jobs / job_key(identity)), [])
        for child in self.old_children:
            self.assertTrue((child / "student_record.json").is_file())

        after = self.controller.inspect(outcome["source"])
        self.assertEqual(after.summary["blocking_errors"], 0)
        for item in after.submissions:
            self.assertTrue(item.workbook_valid, item.submission_id)
            self.assertEqual(item.review_status, "PENDING")
            self.assertNotIn(Action.RUN_MARKING, item.next_actions)
            self.assertIn(Action.REVIEW_EXCEL, item.next_actions)


if __name__ == "__main__":
    unittest.main()
