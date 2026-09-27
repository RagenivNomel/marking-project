"""Offline Stage 3C approved-Excel feedback-card execution tests."""
from __future__ import annotations

from contextlib import ExitStack, contextmanager
import hashlib
import json
from pathlib import Path
import shutil
import sys
import threading
import time
import unittest
import uuid
from unittest.mock import patch


ROOT = Path(__file__).resolve().parents[1]
LOCAL_DEPS = ROOT / ".desktop-deps"
if LOCAL_DEPS.is_dir():
    sys.path.insert(0, str(LOCAL_DEPS))

from openpyxl import load_workbook
from PIL import Image
from PySide6.QtCore import QCoreApplication, QEventLoop

from application import Action, AssignmentSource, WorkflowController
from desktop.bridge import DesktopBridge
from desktop.projection import project
from desktop.teacher_flow import derive_teacher_flow
from excel.schema import AUDIT_SHEET, HEADERS, SHEET
from excel.workbook import ExcelStore
from grading.codex_sol_grader import CodexSolGrader
from grading.mock_grader import MockGrader
from grading.schemas import CRITERIA, CriterionResult, GradingResult, Identity, ROOT as APP_ROOT, Validator
from grading.sol_grader import SolGrader
from rendering.renderer import PillowRenderer
from application.workflows.feedback_rendering import ApprovedFeedbackRenderer
from tests.local_temp import local_test_directory
from workflow.calibration_pipeline import CalibrationPipeline
from workflow.pipeline import Pipeline, job_key


@contextmanager
def _forbid_grading_and_model_calls():
    """Make any Stage 3B/model entry from a Stage 3C test fail immediately."""
    with ExitStack() as stack:
        for owner, name in (
            (MockGrader, "grade"), (CodexSolGrader, "grade"), (SolGrader, "grade"),
            (Pipeline, "run_mock"), (Pipeline, "_grade_essay"), (Pipeline, "run_real_batch"),
            (CalibrationPipeline, "prepare"), (CalibrationPipeline, "run"),
        ):
            stack.enter_context(patch.object(owner, name, side_effect=AssertionError(f"{owner.__name__}.{name} must not run in Stage 3C")))
        yield


def _digest(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def _wait_for_bridge(bridge, app, timeout=8.0):
    deadline = time.monotonic() + timeout
    while bridge.busy and time.monotonic() < deadline:
        app.processEvents(QEventLoop.AllEvents, 30)
        time.sleep(0.003)
    app.processEvents(QEventLoop.AllEvents, 30)
    if bridge.busy:
        raise AssertionError("Feedback render worker did not finish")


class Stage3CRenderTests(unittest.TestCase):
    def setUp(self):
        self.temp = local_test_directory(self._testMethodName)
        self.root = Path(self.temp.__enter__())
        self.addCleanup(self.temp.__exit__, None, None, None)
        self.project = self.root / "project"
        self.batch_id = "stage3c-offline"
        self.jobs_root = self.project / "jobs" / self.batch_id
        self.jobs_root.mkdir(parents=True)
        output = self.project / "output" / self.batch_id
        output.mkdir(parents=True)
        self.config_dir = self.project / "config"
        self.config_dir.mkdir(parents=True)
        for config in (APP_ROOT / "config").glob("*.json"):
            shutil.copy2(config, self.config_dir / config.name)
        self.validator = Validator(self.config_dir)
        self.workbook = output / "results.xlsx"
        self.source = AssignmentSource(
            workbook=self.workbook, job_roots=(self.jobs_root,),
            receipt_roots=(self.jobs_root,), config_dir=self.config_dir,
        )
        self.controller = WorkflowController(project_dir=self.project)
        self._rows = {}

    def _identity(self, student_id, name=None):
        return Identity(str(student_id), name or f"学生{student_id}", "207")

    def _record(self, identity):
        return GradingResult(
            student_id=identity.student_id,
            student_name=identity.student_name,
            class_name=identity.class_name,
            criteria={name: CriterionResult(
                rating="可以更进一步", short_comment=f"{identity.student_id}号原始评语{index}。"
            ) for index, name in enumerate(CRITERIA, 1)},
            teacher_comment=f"{identity.student_id}号原始教师总评。",
        )

    def _add_row(self, identity, status="APPROVED", *, job_id=None, teacher_edits=None,
                 checkpoint=False, source_pdf=False):
        job_id = job_id or job_key(identity)
        digest = f"fixture-digest-{job_id}"
        result = self._record(identity)
        store = ExcelStore(self.workbook, self.validator)
        store.ensure_draft(result, identity, job_id, digest, topic="Q2 作文题目")
        store.set_status(job_id, identity, "VALIDATED", status)
        if teacher_edits:
            book = load_workbook(self.workbook)
            try:
                sheet = book[SHEET]
                row_number = store.row_number(job_id, identity)
                for header, value in teacher_edits.items():
                    sheet.cell(row_number, HEADERS.index(header) + 1).value = value
                book.save(self.workbook)
            finally:
                book.close()
        self._rows[job_id] = (identity, digest, result)
        if checkpoint:
            job_dir = self.jobs_root / job_key(identity)
            job_dir.mkdir(parents=True, exist_ok=True)
            pdf_path = job_dir / "source.pdf"
            if source_pdf:
                pdf_path.write_bytes(b"unchanged offline Stage 3C PDF fixture")
            checkpoint_data = {
                "schema_version": 2, "mode": "REAL_PDF_BATCH", "job_id": job_id,
                "identity": identity.to_dict(), "input_digest": digest,
                "workbook": str(self.workbook), "state": "VALIDATED",
                "history": ["NEW", "SCANNED", "IDENTIFIED", "TRANSCRIBED", "GRADED", "VALIDATED"],
                "last_error": None,
            }
            if source_pdf:
                checkpoint_data["source_sha256"] = _digest(pdf_path)
                checkpoint_data["source_page_count"] = 1
            (job_dir / "student_record.json").write_text(
                json.dumps(checkpoint_data, ensure_ascii=False, indent=2), encoding="utf-8"
            )
            (job_dir / "validated_result.json").write_text(
                json.dumps(result.to_dict(), ensure_ascii=False, indent=2), encoding="utf-8"
            )
        return store

    def _render_pipeline_factory(self, renderer=None):
        if renderer is None:
            renderer = PillowRenderer()
        wrapped = ApprovedFeedbackRenderer(renderer)

        def factory(*, project_dir, batch_id, source):
            return Pipeline(
                project_dir=project_dir, batch_id=batch_id,
                config_dir=source.config_dir or APP_ROOT / "config",
                renderer=wrapped, allow_real=True,
            )
        return factory

    def _controller_with_renderer(self, renderer):
        return WorkflowController(
            project_dir=self.project,
            render_pipeline_factory=self._render_pipeline_factory(renderer),
        )

    def _approved_action(self, controller=None, source=None):
        controller = controller or self.controller
        source = source or self.source
        inspection = controller.inspect(source)
        action = next((entry for entry in inspection.available_actions
                       if entry.action == Action.RENDER_APPROVED), None)
        return inspection, action

    def test_approved_teacher_edited_row_uses_existing_renderer_and_exact_filename(self):
        identity = self._identity("26", "王晨溪")
        edits = {
            "内容分": 20, "语文与结构分": 25, "总分": 45,
            "审题扣题_等级": "做得很好",
            "审题扣题_评语": "教师修改的审题意见。",
            "教师总评": "教师确认后的总评，保留原意。",
        }
        self._add_row(identity, teacher_edits=edits, checkpoint=True, source_pdf=True)
        book_before = _digest(self.workbook)
        source_pdf = self.jobs_root / job_key(identity) / "source.pdf"
        pdf_before = _digest(source_pdf)
        inspection, action = self._approved_action()
        self.assertEqual(inspection.summary["approved"], 1)
        self.assertEqual(inspection.summary["ready_to_render"], 1)
        self.assertEqual(action.submission_ids, (job_key(identity),))
        self.assertTrue(action.execution_wired)

        progress = []
        with _forbid_grading_and_model_calls():
            outcome = self.controller.execute(
                Action.RENDER_APPROVED, self.source, progress_callback=progress.append
            )

        result = outcome["result"]
        self.assertEqual((result["requested"], result["generated"], len(result["failed"])), (1, 1, 0))
        card_dir = self.jobs_root / job_key(identity) / "output"
        card = card_dir / "207-26-王晨溪-作文体检卡.png"
        receipt_path = card_dir / "render_receipt.json"
        self.assertTrue(card.is_file())
        self.assertTrue(receipt_path.is_file())
        self.assertEqual(Path(result["successes"][0]["path"]), card.resolve())
        self.assertEqual(len(list(card_dir.glob("*-作文体检卡.png"))), 1)
        receipt = json.loads(receipt_path.read_text(encoding="utf-8"))
        self.assertEqual(receipt["source"]["source_of_truth"], "approved_excel_row")
        self.assertEqual(receipt["approved_excel_values"]["criteria"][CRITERIA[0]]["rating"], "做得很好")
        self.assertEqual(receipt["approved_excel_values"]["criteria"][CRITERIA[0]]["short_comment"], edits["审题扣题_评语"])
        self.assertEqual(receipt["approved_excel_values"]["teacher_comment"], edits["教师总评"])
        self.assertEqual((receipt["approved_excel_values"]["content_score"],
                          receipt["approved_excel_values"]["language_structure_score"],
                          receipt["approved_excel_values"]["total_score"]), (20, 25, 45))
        self.assertEqual(tuple(receipt["fields"][f"criteria.{name}.rating"]["source_value"]
                               for name in CRITERIA), tuple(
                                   edits["审题扣题_等级"] if name == CRITERIA[0] else "可以更进一步"
                                   for name in CRITERIA
                               ))
        self.assertEqual(len([name for name in receipt["fields"] if name.endswith(".rating")]), 8)
        with Image.open(card) as image:
            self.assertEqual(image.format, "PNG")
            self.assertEqual(image.size, (1024, 1536))
            image.verify()
        self.assertEqual(_digest(self.workbook), book_before)
        self.assertEqual(_digest(source_pdf), pdf_before)
        self.assertEqual(progress[0]["completed"], 0)
        self.assertEqual(progress[-1]["completed"], 1)
        self.assertEqual(progress[-1]["operation"], "render")
        fresh = self.controller.inspect(outcome["source"])
        self.assertEqual(fresh.summary["rendered"], 1)
        self.assertEqual(fresh.summary["ready_to_render"], 0)

        card_hash = _digest(card)
        with _forbid_grading_and_model_calls():
            with self.assertRaises(ValueError):
                self.controller.execute(Action.RENDER_APPROVED, outcome["source"])
        self.assertEqual(_digest(card), card_hash)
        self.assertEqual(len(list(card_dir.glob("*-作文体检卡.png"))), 1)

    def test_pending_row_is_not_rendered_and_teacher_must_approve_in_excel(self):
        identity = self._identity("27")
        self._add_row(identity, status="PENDING")
        inspection, action = self._approved_action()
        self.assertEqual(inspection.summary["approved"], 0)
        self.assertEqual(inspection.summary["ready_to_render"], 0)
        self.assertIsNone(action)
        with _forbid_grading_and_model_calls():
            with self.assertRaises(ValueError):
                self.controller.execute(Action.RENDER_APPROVED, self.source)
        self.assertFalse((self.jobs_root / job_key(identity) / "output").exists())

    def test_real_validated_checkpoint_state_is_supported_without_changing_grading(self):
        identity = self._identity("28")
        self._add_row(identity, checkpoint=True)
        checkpoint_path = self.jobs_root / job_key(identity) / "student_record.json"
        before = checkpoint_path.read_bytes()
        with _forbid_grading_and_model_calls():
            self.controller.execute(Action.RENDER_APPROVED, self.source)
        self.assertEqual(checkpoint_path.read_bytes(), before)
        self.assertTrue((self.jobs_root / job_key(identity) / "output" / "207-28-学生28-作文体检卡.png").is_file())

    def test_failure_keeps_successful_cards_and_retry_touches_only_remaining_row(self):
        first = self._identity("31", "同名学生")
        second = self._identity("32", "同名学生")
        self._add_row(first)
        self._add_row(second)
        original = PillowRenderer()
        second_failed = {"value": False}

        class FailSecondOnce:
            def render(inner_self, record, output_dir, audit=None):
                if record.student_id == second.student_id and not second_failed["value"]:
                    second_failed["value"] = True
                    raise RuntimeError("offline injected renderer failure")
                return original.render(record, output_dir, audit)

        controller = self._controller_with_renderer(FailSecondOnce())
        book_before = _digest(self.workbook)
        with _forbid_grading_and_model_calls():
            first_outcome = controller.execute(Action.RENDER_APPROVED, self.source)
        first_result = first_outcome["result"]
        self.assertEqual((first_result["requested"], first_result["generated"], len(first_result["failed"])), (2, 1, 1))
        first_card = self.jobs_root / job_key(first) / "output" / "207-31-同名学生-作文体检卡.png"
        first_hash = _digest(first_card)
        second_output = self.jobs_root / job_key(second) / "output"
        attempt = json.loads((second_output / "render_attempt.json").read_text(encoding="utf-8"))
        self.assertEqual(attempt["status"], "FAILED")
        self.assertTrue((first_card.parent / "render_receipt.json").is_file())
        self.assertFalse((second_output / "207-32-同名学生-作文体检卡.png").exists())
        retry_inspection, retry_action = self._approved_action(controller)
        self.assertEqual(retry_inspection.summary["rendered"], 1)
        self.assertEqual(retry_action.submission_ids, (job_key(second),))

        with _forbid_grading_and_model_calls():
            retry = controller.execute(Action.RENDER_APPROVED, first_outcome["source"])
        self.assertEqual((retry["result"]["requested"], retry["result"]["generated"]), (1, 1))
        self.assertEqual(_digest(first_card), first_hash)
        second_card = second_output / "207-32-同名学生-作文体检卡.png"
        self.assertTrue(second_card.is_file())
        self.assertEqual(len(list(self.jobs_root.rglob("*-作文体检卡.png"))), 2)
        self.assertEqual(_digest(self.workbook), book_before)

    def _relabel_row(self, job_id, identity):
        """Give an existing row another student's identity, as in a workbook
        written before ExcelStore refused a second row for the same student."""
        book = load_workbook(self.workbook)
        try:
            audit = next(row for row in book[AUDIT_SHEET].iter_rows(min_row=2) if row[0].value == job_id)
            audit[1].value, audit[2].value = identity.class_name, identity.student_id
            row = book[SHEET][audit[3].value]
            row[0].value, row[1].value, row[2].value = (
                identity.class_name, identity.student_id, identity.student_name)
            book.save(self.workbook)
        finally:
            book.close()
        _old_identity, digest, result = self._rows[job_id]
        self._rows[job_id] = (identity, digest, result)

    def test_duplicate_submission_identity_is_stopped_before_a_filename_can_overwrite(self):
        identity = self._identity("40", "重复提交学生")
        self._add_row(identity, job_id="essay-row-a")
        self._add_row(self._identity("41", "重复提交学生"), job_id="essay-row-b")
        self._relabel_row("essay-row-b", identity)
        with _forbid_grading_and_model_calls():
            with self.assertRaisesRegex(ValueError, "same feedback-card filename"):
                self.controller.execute(Action.RENDER_APPROVED, self.source)
        self.assertFalse(list(self.jobs_root.rglob("*-作文体检卡.png")))

    def test_render_runs_in_qt_worker_and_progress_waits_for_real_png_validation(self):
        identity = self._identity("50", "异步学生")
        self._add_row(identity)
        app = QCoreApplication.instance() or QCoreApplication([])
        started, release = threading.Event(), threading.Event()
        original = PillowRenderer()

        class SlowRenderer:
            def render(inner_self, record, output_dir, audit=None):
                started.set()
                if not release.wait(5):
                    raise RuntimeError("test render gate was not released")
                return original.render(record, output_dir, audit)

        bridge = DesktopBridge(controller=self._controller_with_renderer(SlowRenderer()))
        bridge._language = "en"
        bridge._source = self.source
        bridge._inspection = bridge._controller.inspect(self.source)
        bridge._state = project(bridge._inspection, self.source, "Stage 3C fixture", "en")
        bridge._state["teacherFlow"] = derive_teacher_flow(bridge._state, "en")
        self.assertEqual(bridge._state["teacherFlow"]["primaryAction"], Action.RENDER_APPROVED.value)
        self.assertTrue(bridge._state["teacherFlow"]["primaryActionEnabled"])
        before = time.monotonic()
        try:
            with _forbid_grading_and_model_calls():
                bridge.generateFeedback()
                self.assertLess(time.monotonic() - before, 0.5)
                self.assertTrue(started.wait(3), "worker did not reach the renderer")
                self.assertTrue(bridge.busy)
                self.assertEqual(bridge.state["progress"]["operation"], "render")
                self.assertTrue(bridge.state["progress"]["running"])
                release.set()
                _wait_for_bridge(bridge, app)
        finally:
            release.set()
        state = bridge.state
        self.assertEqual(state["progress"]["completed"], 1)
        self.assertFalse(state["progress"]["running"])
        self.assertEqual(state["teacherFlow"]["step"], "complete")
        self.assertEqual(state["renderResult"]["generated"], 1)
        bridge.shutdown()

    def test_bridge_surfaces_render_failure_and_keeps_generate_action_retryable(self):
        identity = self._identity("51", "重试学生")
        self._add_row(identity)
        app = QCoreApplication.instance() or QCoreApplication([])
        original = PillowRenderer()
        failed = {"value": False}

        class FailOnceRenderer:
            def render(inner_self, record, output_dir, audit=None):
                if not failed["value"]:
                    failed["value"] = True
                    raise RuntimeError("offline injected render failure")
                return original.render(record, output_dir, audit)

        controller = self._controller_with_renderer(FailOnceRenderer())
        bridge = DesktopBridge(controller=controller)
        bridge._language = "en"
        bridge._source = self.source
        bridge._inspection = controller.inspect(self.source)
        bridge._state = project(bridge._inspection, self.source, "Stage 3C retry fixture", "en")
        bridge._state["teacherFlow"] = derive_teacher_flow(bridge._state, "en")

        with _forbid_grading_and_model_calls():
            bridge.generateFeedback()
            _wait_for_bridge(bridge, app)
            failed_state = bridge.state
            self.assertEqual(len(failed_state["renderResult"]["failed"]), 1)
            self.assertEqual(failed_state["progress"]["completed"], 0)
            self.assertEqual(failed_state["progress"]["attention"], 1)
            self.assertTrue(failed_state["teacherFlow"]["primaryActionEnabled"])
            self.assertEqual(failed_state["teacherFlow"]["primaryAction"], Action.RENDER_APPROVED.value)
            bridge.generateFeedback()
            _wait_for_bridge(bridge, app)

        final = bridge.state
        self.assertEqual(final["teacherFlow"]["step"], "complete")
        self.assertEqual(final["progress"]["completed"], 1)
        self.assertEqual(final["renderResult"]["generated"], 1)
        card_dir = self.jobs_root / job_key(identity) / "output"
        self.assertTrue((card_dir / "207-51-重试学生-作文体检卡.png").is_file())
        attempt = json.loads((card_dir / "render_attempt.json").read_text(encoding="utf-8"))
        self.assertEqual(attempt["status"], "PASS")
        bridge.shutdown()


if __name__ == "__main__":
    unittest.main()
