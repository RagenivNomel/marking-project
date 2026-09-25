"""Execute Phase 1 acceptance scenarios using fictional data only."""
import argparse
import copy
import hashlib
import json
from pathlib import Path
from unittest.mock import patch

from openpyxl import load_workbook

from excel.schema import HEADERS, SHEET
from grading.mock_grader import MockGrader
from grading.schemas import CRITERIA, RATINGS, ROOT, Identity, Rating, ValidationError
from rendering.renderer import MockRenderer
from tests.local_fixtures import local_fixtures_available, report_local_fixture_skip
from workflow.pipeline import Pipeline, job_key
from workflow.state import State
from workflow.storage import atomic_json, read_json


def sha256(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


class TrackingRenderer(MockRenderer):
    def __init__(self):
        self.calls = 0

    def render(self, record, output_dir, audit=None):
        self.calls += 1
        return super().render(record, output_dir, audit)


def failure_mutators():
    def invalid_rating(result):
        result["criteria"][CRITERIA[0]]["rating"] = "优秀"

    def missing_criterion(result):
        result["criteria"].pop(CRITERIA[0])

    def empty_comment(result):
        result["criteria"][CRITERIA[0]]["short_comment"] = " \n"

    def over_limit(result):
        result["criteria"][CRITERIA[0]]["short_comment"] = "长" * 51

    def identity_mismatch(result):
        result["student_name"] = "另一个学生"

    return {
        "invalid-rating": invalid_rating,
        "missing-criterion": missing_criterion,
        "empty-comment": empty_comment,
        "over-limit": over_limit,
        "identity-mismatch": identity_mismatch,
    }


def run_success(run_root, source):
    identity = Identity.from_dict(source["identity"])
    renderer = TrackingRenderer()
    pipeline = Pipeline(run_root / "success", "synthetic-success", renderer=renderer)
    result = pipeline.run_mock(source)
    checkpoint = read_json(Path(result["job_dir"]) / "student_record.json")
    expected_history = [state.value for state in State]
    assert checkpoint["history"] == expected_history
    approved = pipeline.excel.get(job_key(identity), identity, approved=True).to_dict()
    receipt = read_json(result["render"])
    assert receipt["approved_excel_values"] == approved
    validated = read_json(Path(result["job_dir"]) / "validated_result.json")
    assert approved == validated

    edit_renderer = TrackingRenderer()
    edit_pipeline = Pipeline(run_root / "rerender-edit", "synthetic-rerender-edit", renderer=edit_renderer)
    edit_result = edit_pipeline.run_mock(source)
    exact_teacher_edit = "  教师原话，保留前后空格。\n第二行也必须保留。  "
    exact_criterion_edit = "  具体修改：雨声、冷风与脚步。\n"
    book = load_workbook(edit_pipeline.excel.path)
    try:
        sheet = book[SHEET]
        assert [cell.value for cell in sheet[1]] == HEADERS
        sheet["F2"] = exact_criterion_edit
        sheet["X2"] = exact_teacher_edit
        book.save(edit_pipeline.excel.path)
    finally:
        book.close()
    with patch.object(MockGrader, "grade", side_effect=AssertionError("rerender must not call a grader")):
        artifact = edit_pipeline.rerender(identity)
    rerendered = read_json(artifact)["approved_excel_values"]
    reread = edit_pipeline.excel.get(job_key(identity), identity, approved=True).to_dict()
    assert rerendered == reread
    assert rerendered["teacher_comment"] == exact_teacher_edit
    assert rerendered["criteria"][CRITERIA[0]]["short_comment"] == exact_criterion_edit
    return {
        "state": result["state"],
        "history": checkpoint["history"],
        "workbook": result["workbook"],
        "render": result["render"],
        "renderer_calls": renderer.calls,
        "final_excel_equals_validated_json": approved == validated,
        "initial_render_equals_approved_excel": receipt["approved_excel_values"] == approved,
        "headers": HEADERS,
        "independent_rerender_edit_proof": {
            "workbook": str(edit_pipeline.excel.path),
            "initial_render": edit_result["render"],
            "rerender": str(artifact),
            "renderer_calls": edit_renderer.calls,
            "grader_calls_during_rerender": 0,
            "rerender_equals_fresh_approved_excel": rerendered == reread,
            "exact_teacher_edit_preserved": rerendered["teacher_comment"] == exact_teacher_edit,
            "exact_criterion_edit_preserved": rerendered["criteria"][CRITERIA[0]]["short_comment"] == exact_criterion_edit,
        },
    }


def run_failures(run_root, source):
    identity = Identity.from_dict(source["identity"])
    evidence = {}
    for name, mutate in failure_mutators().items():
        scenario_root = run_root / name
        renderer = TrackingRenderer()

        class MutatingGrader:
            def grade(self, request):
                result = MockGrader().grade(request)
                mutate(result)
                return result

        pipeline = Pipeline(scenario_root, "synthetic-failure", grader_factory=MutatingGrader, renderer=renderer)
        error = None
        try:
            pipeline.run_mock(source)
        except ValidationError as exc:
            error = {"type": type(exc).__name__, "message": str(exc)}
        assert error is not None
        job = pipeline.jobs_dir / job_key(identity)
        checkpoint = read_json(job / "student_record.json")
        excel_files = [str(path) for path in scenario_root.rglob("*.xlsx")]
        render_files = [str(path) for path in scenario_root.rglob("mock_render.json")]
        assert not excel_files and not render_files and renderer.calls == 0
        evidence[name] = {
            "error": error,
            "durable_state": checkpoint["state"],
            "excel_files_after": excel_files,
            "render_files_after": render_files,
            "renderer_calls": renderer.calls,
            "no_excel_write": True,
            "no_render": True,
        }

    name = "invalid-transition"
    scenario_root = run_root / name
    renderer = TrackingRenderer()
    pipeline = Pipeline(scenario_root, "synthetic-failure", renderer=renderer)
    original_advance = pipeline._advance

    def skip_required_state(job, checkpoint, _target):
        return original_advance(job, checkpoint, State.IDENTIFIED)

    error = None
    try:
        with patch.object(pipeline, "_advance", side_effect=skip_required_state):
            pipeline.run_mock(source)
    except ValueError as exc:
        error = {"type": type(exc).__name__, "message": str(exc)}
    assert error is not None and "NEW -> IDENTIFIED" in error["message"]
    job = pipeline.jobs_dir / job_key(identity)
    checkpoint = read_json(job / "student_record.json")
    excel_files = [str(path) for path in scenario_root.rglob("*.xlsx")]
    render_files = [str(path) for path in scenario_root.rglob("mock_render.json")]
    assert not excel_files and not render_files and renderer.calls == 0
    evidence[name] = {
        "error": error,
        "durable_state": checkpoint["state"],
        "excel_files_after": excel_files,
        "render_files_after": render_files,
        "renderer_calls": renderer.calls,
        "no_excel_write": True,
        "no_render": True,
    }
    return evidence


def run_restarts(run_root, source):
    evidence = {}
    for stop in list(State)[1:-1]:
        scenario_root = run_root / f"restart-{stop.value.lower()}"
        pipeline = Pipeline(scenario_root, "synthetic-restart")
        original = pipeline._advance

        def interrupt(job, checkpoint, target, *, stop=stop):
            original(job, checkpoint, target)
            if target == stop:
                raise RuntimeError("simulated interruption")

        try:
            with patch.object(pipeline, "_advance", side_effect=interrupt):
                pipeline.run_mock(source)
        except RuntimeError as exc:
            assert str(exc) == "simulated interruption"
        resumed = Pipeline(scenario_root, "synthetic-restart").run_mock(source)
        checkpoint = read_json(Path(resumed["job_dir"]) / "student_record.json")
        book = load_workbook(resumed["workbook"], read_only=True)
        try:
            row_count = book[SHEET].max_row
        finally:
            book.close()
        assert resumed["state"] == "COMPLETE" and checkpoint["history"] == [state.value for state in State]
        assert row_count == 2
        evidence[stop.value] = {"resumed_state": resumed["state"], "history": checkpoint["history"], "excel_rows_including_header": row_count}
    return evidence


def run_isolation(run_root, source):
    provider_ids = []
    requests = []

    class CapturingGrader(MockGrader):
        def grade(self, request):
            requests.append(request)
            return super().grade(request)

    def factory():
        provider = CapturingGrader()
        provider_ids.append(id(provider))
        return provider

    pipeline = Pipeline(run_root / "isolation", "synthetic-isolation", grader_factory=factory)
    first = pipeline.run_mock(source)
    second_source = copy.deepcopy(source)
    second_source["identity"].update(class_name="MOCK-208", student_name="另一测试学生")
    second_source["essay"] = "完全不同的第二篇虚构作文。"
    second = pipeline.run_mock(second_source)
    assert first["job_dir"] != second["job_dir"]
    assert len(provider_ids) == 2 and provider_ids[0] != provider_ids[1]
    assert requests[0].essay == source["essay"] and requests[1].essay == second_source["essay"]
    assert set(vars(requests[0])) == {"identity", "essay"}
    return {
        "distinct_job_directories": [first["job_dir"], second["job_dir"]],
        "fresh_grader_per_student": True,
        "grader_input_fields": sorted(vars(requests[0])),
        "different_essay_inputs": True,
    }


def preserved_assets():
    backup = ROOT / "backups" / "phase1_originals"
    evidence = []
    for item in json.loads((backup / "manifest.json").read_text(encoding="utf-8-sig")):
        original = ROOT / item["path"]
        copy_path = backup / item["path"]
        original_hash = sha256(original).upper()
        copy_hash = sha256(copy_path).upper()
        assert original_hash == copy_hash == item["sha256"]
        evidence.append({"path": item["path"], "original_sha256": original_hash, "backup_sha256": copy_hash, "unchanged": True})
    return evidence


def main(argv=None):
    parser = argparse.ArgumentParser()
    parser.add_argument("--run-root", type=Path, required=True)
    parser.add_argument("--report", type=Path, required=True)
    args = parser.parse_args(argv)
    backup = ROOT / "backups" / "phase1_originals"
    manifest_path = backup / "manifest.json"
    if report_local_fixture_skip([manifest_path]):
        return 0
    entries = json.loads(manifest_path.read_text(encoding="utf-8-sig"))
    if not entries:
        print("SKIPPED: local regression fixture not installed")
        return 0
    required_assets = [path for item in entries for path in (ROOT / item["path"], backup / item["path"])]
    if not local_fixtures_available(required_assets):
        print("SKIPPED: local regression fixture not installed")
        return 0
    args.run_root.mkdir(parents=True, exist_ok=False)
    source = read_json(ROOT / "examples" / "mock_student.json")
    report = {
        "verdict": "PASS",
        "scope": "fictional mock data only; no real submissions; no AI backend; JSON placeholder renderer only",
        "criteria": list(CRITERIA),
        "rating_enum": {member.name: member.value for member in Rating},
        "ratings": list(RATINGS),
        "success": run_success(args.run_root, source),
        "failures": run_failures(args.run_root, source),
        "restart_resume": run_restarts(args.run_root, source),
        "student_isolation": run_isolation(args.run_root, source),
        "preserved_assets": preserved_assets(),
    }
    atomic_json(args.report, report)
    print(json.dumps(report, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
