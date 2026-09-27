import json
import hashlib
from pathlib import Path
import threading
import time
import unittest
from types import SimpleNamespace
from concurrent.futures import ThreadPoolExecutor

from openpyxl import load_workbook
from PIL import Image
from pypdf import PdfWriter

from excel.schema import AUDIT_SHEET, HEADERS, SHEET
from excel.workbook import ExcelStore
from grading.codex_sol_grader import CodexSolGrader
from grading.schemas import Identity, ROOT, Validator
from tests.local_temp import local_test_directory
from tests.test_real_batch_integration import valid_response, working_state
from workflow.calibration_pipeline import CalibrationPipeline
from workflow.pipeline import MAX_CONCURRENT_GRADERS, Pipeline, job_key
from workflow.storage import read_json


class OfflineCodexRunner:
    """Exercise the real isolated grader with a fake, offline CLI process."""

    def __init__(self, identities, *, hold_first_three=False, delays=None,
                 fail_once=(), invalid=(), hold_names=()):
        self.identities = list(identities)
        self.by_name = {item.student_name: item for item in identities}
        self.hold_first_three = hold_first_three
        self.delays = dict(delays or {})
        self.fail_once = set(fail_once)
        self.invalid = set(invalid)
        self.hold_names = set(hold_names)
        self.calls = []
        self.rendered = []
        self.graders = []
        self.lock = threading.Lock()
        self.active = 0
        self.maximum_active = 0
        self.first_three_entered = threading.Event()
        self.release = threading.Event()
        if not hold_first_three and not self.hold_names:
            self.release.set()

    def page_renderer(self, pdf_bytes, workspace, _dpi):
        path = workspace / "page_001.png"
        Image.new("RGB", (8, 8), "white").save(path, format="PNG")
        with self.lock:
            self.rendered.append((str(workspace.resolve()), str(path.resolve()),
                                  hashlib.sha256(pdf_bytes).hexdigest()))
        return [path]

    def grader_factory(self, project_dir):
        grader = CodexSolGrader(
            model="gpt-5.6-luna",
            reasoning_effort="xhigh",
            process_runner=self.process_runner,
            page_renderer=self.page_renderer,
            workspace_parent=Path(project_dir) / "jobs" / ".codex-workspaces",
        )
        with self.lock:
            self.graders.append(grader)
        return grader

    def process_runner(self, command, **kwargs):
        command = list(command)
        if "login" in command:
            return SimpleNamespace(returncode=0, stdout="Logged in using ChatGPT", stderr="")
        if "debug" in command:
            catalog = {"models": [{
                "slug": "gpt-5.6-luna",
                "input_modalities": ["text", "image"],
                "supported_reasoning_levels": [{"effort": "xhigh"}],
            }]}
            return SimpleNamespace(returncode=0, stdout=json.dumps(catalog), stderr="")

        prompt = kwargs["input"]
        identity = next(item for item in self.identities if item.student_name in prompt)
        with self.lock:
            invocation = len(self.calls) + 1
            self.active += 1
            active_at_start = self.active
            self.maximum_active = max(self.maximum_active, self.active)
            if self.active >= 3:
                self.first_three_entered.set()
            call = {
                "identity": identity,
                "prompt": prompt,
                "cwd": str(Path(kwargs["cwd"]).resolve()),
                "images": [command[i + 1] for i, arg in enumerate(command[:-1]) if arg == "--image"],
                "command": command,
                "active_at_start": active_at_start,
                "invocation": invocation,
                "started_at": time.monotonic(),
            }
            self.calls.append(call)
            fail_this_call = identity.student_name in self.fail_once
            if fail_this_call:
                self.fail_once.remove(identity.student_name)

        should_hold = ((self.hold_first_three and invocation <= 3)
                       or identity.student_name in self.hold_names)
        try:
            if should_hold:
                if not self.release.wait(10):
                    raise TimeoutError("offline grading gate was not released")
            delay = self.delays.get(identity.student_name, 0)
            if delay:
                time.sleep(delay)
            if fail_this_call:
                raise RuntimeError("synthetic offline grader failure")
            response = valid_response(identity)
            if identity.student_name in self.invalid:
                response["grading"]["student_id"] = "wrong-student"
            output_path = Path(command[command.index("--output-last-message") + 1])
            output_path.write_text(json.dumps(response, ensure_ascii=False), encoding="utf-8")
            return SimpleNamespace(returncode=0, stdout="", stderr="")
        finally:
            with self.lock:
                call["finished_at"] = time.monotonic()
                self.active -= 1


class RealBatchParallelTests(unittest.TestCase):
    def setUp(self):
        self.temp = local_test_directory("real-batch-parallel")
        self.root = Path(self.temp.__enter__())
        self.addCleanup(self.temp.__exit__, None, None, None)
        self.config_dir = ROOT / "config"
        self.source_dir = self.root / "essays"
        self.source_dir.mkdir()
        self.results_path = self.root / "Results" / "results.xlsx"
        self.batch_id = "parallel-offline"
        self.identities = [
            Identity(f"{index:02d}", f"离线学生{index:02d}", "207")
            for index in range(1, 6)
        ]
        self.records = []
        for sequence, identity in enumerate(self.identities, 1):
            source_name = f"submission_{sequence:03}.pdf"
            writer = PdfWriter()
            writer.add_blank_page(width=612, height=792)
            writer.add_metadata({"/StudentMarker": f"essay-{identity.student_id}"})
            with (self.source_dir / source_name).open("wb") as stream:
                writer.write(stream)
            self.records.append({
                "sequence": sequence,
                "source_pdf": source_name,
                **identity.to_dict(),
                "match_status": "STRONG_ROSTER_MATCH",
            })

    def make_pipeline(self, offline, batch_id=None):
        def calibration_factory(project_dir, model_profile=None):
            return CalibrationPipeline(
                project_dir,
                config_dir=self.config_dir,
                model_profile=model_profile,
                grader_factory=lambda: offline.grader_factory(project_dir),
            )

        return Pipeline(
            self.root,
            batch_id or self.batch_id,
            config_dir=self.config_dir,
            allow_real=True,
            workbook_path=self.results_path,
            real_pipeline_factory=calibration_factory,
        )

    def read_rows(self):
        book = load_workbook(self.results_path, read_only=True, data_only=True)
        try:
            return list(book[SHEET].iter_rows(min_row=2, values_only=True))
        finally:
            book.close()

    def identity_names_in_calls(self, offline):
        return [call["identity"].student_name for call in offline.calls]

    def test_three_overlap_fourth_waits_and_students_remain_isolated(self):
        self.assertEqual(MAX_CONCURRENT_GRADERS, 3)
        offline = OfflineCodexRunner(self.identities[:4], hold_first_three=True)
        pipeline = self.make_pipeline(offline)
        with ThreadPoolExecutor(max_workers=1) as caller:
            batch = caller.submit(
                pipeline.run_real_batch, self.records[:4], self.source_dir,
                model_profile="luna_xhigh",
            )
            self.assertTrue(offline.first_three_entered.wait(10), "Three fake graders did not overlap")
            with offline.lock:
                self.assertEqual(offline.active, 3)
                self.assertEqual(len(offline.calls), 3)
                self.assertEqual(offline.maximum_active, 3)
                self.assertNotIn(self.identities[3].student_name, self.identity_names_in_calls(offline))
            offline.release.set()
            result = batch.result(timeout=30)

        self.assertEqual(len(result["validated"]), 4)
        self.assertEqual(result["failed"], [])
        self.assertEqual(offline.maximum_active, 3)
        self.assertEqual(len(offline.graders), 4)
        self.assertEqual(len({id(grader) for grader in offline.graders}), 4)
        fourth = next(call for call in offline.calls if call["identity"] == self.identities[3])
        self.assertLessEqual(fourth["active_at_start"], 2)
        self.assertEqual(len({call["cwd"] for call in offline.calls}), 4)
        self.assertTrue(all(len(call["images"]) == 1 for call in offline.calls))
        self.assertEqual(len({call["images"][0] for call in offline.calls}), 4)
        self.assertEqual(len(offline.rendered), 4)
        rendered_hashes = {workspace: source_hash for workspace, _image, source_hash in offline.rendered}
        self.assertEqual(len(rendered_hashes), 4)
        for call in offline.calls:
            self.assertIn(call["identity"].student_name, call["prompt"])
            for other in self.identities[:4]:
                if other != call["identity"]:
                    self.assertNotIn(other.student_name, call["prompt"])
            self.assertIn("--ephemeral", call["command"])
            self.assertEqual(call["command"][call["command"].index("--model") + 1], "gpt-5.6-luna")
            self.assertIn('model_reasoning_effort="xhigh"', call["command"])
            record = next(item for item in self.records
                          if item["student_id"] == call["identity"].student_id)
            expected_hash = hashlib.sha256((self.source_dir / record["source_pdf"]).read_bytes()).hexdigest()
            self.assertEqual(rendered_hashes[call["cwd"]], expected_hash)

        jobs_root = self.root / "jobs" / self.batch_id
        for identity in self.identities[:4]:
            # Each essay had its own working folder, cleared once its row was saved.
            self.assertEqual(working_state(jobs_root / job_key(identity)), [])
        rows = self.read_rows()
        self.assertEqual([row[1] for row in rows], ["01", "02", "03", "04"])
        self.assertEqual(len({row[1] for row in rows}), 4)
        workbook = load_workbook(self.results_path, read_only=True, data_only=True)
        try:
            self.assertEqual(workbook[AUDIT_SHEET].max_row, 5)
        finally:
            workbook.close()
        evidence = read_json(jobs_root / "parallel_timing.json")
        self.assertEqual(evidence["max_simultaneous_graders"], 3)
        self.assertEqual(evidence["max_simultaneous_workbook_writers"], 1)
        self.assertEqual(len(evidence["students"]), 4)
        self.assertTrue(all({"grading_started_at", "grading_finished_at", "validation_finished_at",
                             "persistence_finished_at"}.issubset(events)
                            for events in evidence["students"].values()))

    def test_out_of_order_finishes_single_writer_teacher_edits_and_backups(self):
        seed = OfflineCodexRunner(self.identities[:1])
        self.make_pipeline(seed).run_real_batch(
            self.records[:1], self.source_dir, model_profile="luna_xhigh"
        )
        workbook = load_workbook(self.results_path)
        try:
            sheet = workbook[SHEET]
            sheet.cell(2, HEADERS.index("内容分") + 1).value = 12
            sheet.cell(2, HEADERS.index("语文与结构分") + 1).value = 23
            sheet.cell(2, HEADERS.index("总分") + 1).value = 35
            sheet.cell(2, HEADERS.index("教师总评") + 1).value = "教师已编辑的总评"
            sheet.cell(2, HEADERS.index("审核状态") + 1).value = "APPROVED"
            workbook.save(self.results_path)
        finally:
            workbook.close()

        offline = OfflineCodexRunner(
            self.identities[1:4],
            delays={self.identities[1].student_name: 0.25,
                    self.identities[2].student_name: 0.02,
                    self.identities[3].student_name: 0.01},
        )
        pipeline = self.make_pipeline(offline)
        original_ensure = pipeline.excel.ensure_draft
        writer_lock = threading.Lock()
        writers = 0
        maximum_writers = 0
        seen_before_write = []

        def counted_ensure(*args, **kwargs):
            nonlocal writers, maximum_writers
            with writer_lock:
                writers += 1
                maximum_writers = max(maximum_writers, writers)
            try:
                current_ids = []
                if self.results_path.exists():
                    current = load_workbook(self.results_path, read_only=True, data_only=True)
                    try:
                        current_ids = [row[1] for row in current[SHEET].iter_rows(min_row=2, values_only=True)]
                    finally:
                        current.close()
                seen_before_write.append(tuple(current_ids))
                time.sleep(0.01)
                return original_ensure(*args, **kwargs)
            finally:
                with writer_lock:
                    writers -= 1

        pipeline.excel.ensure_draft = counted_ensure
        result = pipeline.run_real_batch(self.records[:4], self.source_dir, model_profile="luna_xhigh")
        self.assertEqual(result["already_complete"], 1)
        self.assertEqual([item["identity"]["student_id"] for item in result["validated"]], ["02", "03", "04"])
        self.assertEqual(maximum_writers, 1)
        self.assertEqual(seen_before_write, [("01",), ("01", "02"), ("01", "02", "03")])
        call_finish = {call["identity"].student_id: call["finished_at"] for call in offline.calls}
        self.assertLess(call_finish["03"], call_finish["02"])
        self.assertLess(call_finish["04"], call_finish["02"])

        rows = self.read_rows()
        self.assertEqual([row[1] for row in rows], ["01", "02", "03", "04"])
        self.assertEqual(len({row[1] for row in rows}), 4)
        self.assertEqual(rows[0][4:7], (12, 23, 35))
        self.assertEqual(rows[0][23], "教师已编辑的总评")
        self.assertEqual(rows[0][24], "APPROVED")
        self.assertEqual(set(self.identity_names_in_calls(offline)),
                         {item.student_name for item in self.identities[1:4]})
        backups = list((self.root / "jobs" / self.batch_id / "workbook_backups").glob("*.xlsx"))
        self.assertGreaterEqual(len(backups), 3)
        for backup in backups:
            saved = load_workbook(backup, read_only=True, data_only=True)
            try:
                self.assertIn(SHEET, saved.sheetnames)
                self.assertIn(AUDIT_SHEET, saved.sheetnames)
            finally:
                saved.close()
        evidence = read_json(self.root / "jobs" / self.batch_id / "parallel_timing.json")
        self.assertGreaterEqual(evidence["max_simultaneous_graders"], 1)
        self.assertLessEqual(evidence["max_simultaneous_graders"], 3)
        self.assertEqual(evidence["max_simultaneous_workbook_writers"], 1)

    def test_failure_isolation_and_single_continue_operation(self):
        offline = OfflineCodexRunner(
            self.identities[:4],
            fail_once={self.identities[1].student_name},
            invalid={self.identities[2].student_name},
        )
        pipeline = self.make_pipeline(offline)
        first_progress = []
        first = pipeline.run_real_batch(
            self.records[:4], self.source_dir, model_profile="luna_xhigh",
            progress_callback=lambda item: first_progress.append(dict(item)),
        )
        self.assertEqual([item["identity"]["student_id"] for item in first["validated"]], ["01", "04"])
        self.assertEqual([item["identity"]["student_id"] for item in first["failed"]], ["02", "03"])
        self.assertEqual([row[1] for row in self.read_rows()], ["01", "04"])
        # Failed work is attached to its own identity in the returned batch;
        # the UI attention count is a batch total, never another student's state.
        self.assertEqual({item["identity"]["student_id"] for item in first["failed"]}, {"02", "03"})
        self.assertEqual(first_progress[-1]["completed"], 2)
        self.assertEqual(first_progress[-1]["attention"], 2)
        self.assertEqual(first_progress[-1]["active"], 0)

        retry_pipeline = self.make_pipeline(offline)
        retry = retry_pipeline.run_real_batch(
            [self.records[1]], self.source_dir, model_profile="luna_xhigh",
        )
        self.assertEqual(len(retry["validated"]), 1)
        self.assertEqual(retry["validated"][0]["identity"]["student_id"], "02")
        self.assertEqual(self.identity_names_in_calls(offline).count(self.identities[0].student_name), 1)
        self.assertEqual(self.identity_names_in_calls(offline).count(self.identities[1].student_name), 2)
        self.assertEqual(self.identity_names_in_calls(offline).count(self.identities[2].student_name), 1)
        self.assertEqual(self.identity_names_in_calls(offline).count(self.identities[3].student_name), 1)
        self.assertEqual([row[1] for row in self.read_rows()], ["01", "02", "04"])

    def test_persistence_failure_keeps_workbook_readable_and_continue_regrades_uncommitted_result(self):
        offline = OfflineCodexRunner(self.identities[:2])
        pipeline = self.make_pipeline(offline)
        original_save = pipeline.excel._save
        failed_once = False

        def fail_second_student_once(book):
            nonlocal failed_once
            rows = list(book[SHEET].iter_rows(min_row=2, values_only=True))
            if not failed_once and rows and rows[-1][1] == "02":
                failed_once = True
                raise OSError("synthetic persistence failure")
            return original_save(book)

        pipeline.excel._save = fail_second_student_once
        first = pipeline.run_real_batch(self.records[:2], self.source_dir, model_profile="luna_xhigh")
        self.assertEqual([item["identity"]["student_id"] for item in first["validated"]], ["01"])
        self.assertEqual([item["identity"]["student_id"] for item in first["failed"]], ["02"])
        self.assertEqual([row[1] for row in self.read_rows()], ["01"])

        resumed = self.make_pipeline(offline).run_real_batch(
            self.records[:2], self.source_dir, model_profile="luna_xhigh"
        )
        self.assertEqual(resumed["already_complete"], 1)
        self.assertEqual([item["identity"]["student_id"] for item in resumed["validated"]], ["02"])
        self.assertEqual(len(offline.calls), 3)
        self.assertEqual([row[1] for row in self.read_rows()], ["01", "02"])
        self.assertTrue(failed_once)

    def test_authoritative_commit_survives_transient_corruption_and_profile_change(self):
        offline = OfflineCodexRunner(self.identities[:1])
        pipeline = self.make_pipeline(offline)
        first = pipeline.run_real_batch(
            self.records[:1], self.source_dir, model_profile="luna_xhigh"
        )
        self.assertEqual(len(first["validated"]), 1)
        job_dir = self.root / "jobs" / self.batch_id / job_key(self.identities[0])
        job_dir.mkdir(parents=True, exist_ok=True)
        (job_dir / "student_record.json").write_text("{", encoding="utf-8")
        (job_dir / "validated_result.json").write_text("not-json", encoding="utf-8")
        offline.calls.clear()

        reopened = self.make_pipeline(offline).run_real_batch(
            self.records[:1], self.source_dir, model_profile="sol_medium"
        )
        self.assertEqual(reopened["already_complete"], 1)
        self.assertEqual(reopened["validated"], [])
        self.assertEqual(reopened["failed"], [])
        self.assertEqual(offline.calls, [])
        self.assertEqual([row[1] for row in self.read_rows()], ["01"])

    def test_progress_stays_in_progress_until_authoritative_save_finishes(self):
        offline = OfflineCodexRunner(self.identities[:1])
        pipeline = self.make_pipeline(offline)
        original_ensure = pipeline.excel.ensure_draft
        writer_started = threading.Event()
        allow_write = threading.Event()
        progress = []

        def held_ensure(*args, **kwargs):
            writer_started.set()
            if not allow_write.wait(10):
                raise TimeoutError("test did not release persistence")
            return original_ensure(*args, **kwargs)

        pipeline.excel.ensure_draft = held_ensure
        with ThreadPoolExecutor(max_workers=1) as caller:
            future = caller.submit(
                pipeline.run_real_batch, self.records[:1], self.source_dir,
                model_profile="luna_xhigh", progress_callback=lambda item: progress.append(dict(item)),
            )
            self.assertTrue(writer_started.wait(10), "Persistence did not reach its controlled boundary")
            waiting_for_write = progress[-1]
            self.assertEqual(waiting_for_write["completed"], 0)
            self.assertEqual(waiting_for_write["active"], 1)
            self.assertEqual(waiting_for_write["attention"], 0)
            self.assertFalse(future.done())
            allow_write.set()
            result = future.result(timeout=20)

        self.assertEqual(len(result["validated"]), 1)
        self.assertEqual(progress[-1]["completed"], 1)
        self.assertEqual(progress[-1]["active"], 0)
        self.assertEqual(progress[-1]["attention"], 0)
        self.assertLessEqual(max(item["completed"] for item in progress), 1)
        self.assertEqual(progress[-1]["completed"], len(result["validated"]))
        increases = [max(0, right["completed"] - left["completed"])
                     for left, right in zip(progress, progress[1:])]
        self.assertEqual(sum(increases), 1)

    def test_partial_close_resume_skips_persisted_students(self):
        offline = OfflineCodexRunner(
            self.identities,
            hold_names={self.identities[1].student_name, self.identities[2].student_name},
        )
        pipeline = self.make_pipeline(offline)
        stop = threading.Event()
        progress = []

        def stop_after_first_saved(item):
            progress.append(dict(item))
            if item["completed"] >= 1 and not stop.is_set():
                stop.set()
                offline.release.set()

        with ThreadPoolExecutor(max_workers=1) as caller:
            first_future = caller.submit(
                pipeline.run_real_batch, self.records, self.source_dir,
                model_profile="luna_xhigh", progress_callback=stop_after_first_saved,
                cancelled=stop.is_set,
            )
            first = first_future.result(timeout=30)
        self.assertTrue(first["interrupted"])
        self.assertEqual(len(first["validated"]), 3)
        self.assertEqual(set(self.identity_names_in_calls(offline)),
                         {item.student_name for item in self.identities[:3]})
        self.assertEqual([row[1] for row in self.read_rows()], ["01", "02", "03"])

        reopened = self.make_pipeline(offline).run_real_batch(
            self.records, self.source_dir, model_profile="luna_xhigh"
        )
        self.assertEqual(reopened["already_complete"], 3)
        self.assertEqual([item["identity"]["student_id"] for item in reopened["validated"]], ["04", "05"])
        self.assertEqual(len(offline.calls), 5)
        self.assertEqual([row[1] for row in self.read_rows()], ["01", "02", "03", "04", "05"])
        self.assertTrue(progress[-1]["interrupted"])

    def test_excel_store_serializes_separate_instances_for_one_workbook(self):
        validator = Validator(self.config_dir)
        stores = [ExcelStore(self.results_path, validator), ExcelStore(self.results_path, validator)]
        validated = [
            validator.validate(valid_response(identity)["grading"], identity)
            for identity in self.identities[:2]
        ]
        original_save = ExcelStore._save
        counter_lock = threading.Lock()
        active = 0
        maximum = 0

        def monitored_save(store, book):
            nonlocal active, maximum
            with counter_lock:
                active += 1
                maximum = max(maximum, active)
            try:
                time.sleep(0.04)
                return original_save(store, book)
            finally:
                with counter_lock:
                    active -= 1

        ExcelStore._save = monitored_save
        try:
            with ThreadPoolExecutor(max_workers=2) as writers:
                futures = [
                    writers.submit(store.ensure_draft, result, identity, job_key(identity), f"digest-{i}")
                    for i, (store, result, identity) in enumerate(zip(stores, validated, self.identities[:2]))
                ]
                for future in futures:
                    future.result(timeout=20)
        finally:
            ExcelStore._save = original_save
        self.assertEqual(maximum, 1)
        self.assertEqual([row[1] for row in self.read_rows()], ["01", "02"])


if __name__ == "__main__":
    unittest.main()
