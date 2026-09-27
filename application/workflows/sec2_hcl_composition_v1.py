"""Stage 1 inspection plus Stage 3B/3C execution adapters for Section 2 HCL essays.

Inspection remains a read-only projection over explicit evidence. RUN_MARKING
rechecks the assignment and delegates to the production batch pipeline.
RENDER_APPROVED rechecks teacher-approved Excel and delegates to the existing
deterministic renderer; it does not approve teacher feedback or rerun grading.
"""
from __future__ import annotations

from collections import Counter, defaultdict
from dataclasses import dataclass, field, replace
from datetime import datetime, timezone
from hashlib import sha256
import os
from pathlib import Path
import shutil
from uuid import uuid4

from openpyxl import load_workbook

from application.models import (
    Action, AssignmentSource, Attention, AvailableAction, Inspection, Stage,
    SubmissionState,
)
from excel.schema import AUDIT_HEADERS, AUDIT_SHEET, HEADERS, SHEET
from excel.workbook import ExcelStore, read_roster, roster_text
from grading.schemas import Identity, ROOT, ValidationError, Validator
from scanning.intake import apply_identity_decisions, read_existing_split
from scanning.split_by_student import split_continuous_anonymous
from workflow.pipeline import job_key
from workflow.storage import atomic_json, batch_lock, lock_is_active, read_json
from .task_storage import content_digest, safe_name, task_id, task_output_paths


WORKFLOW_ID = "sec2_hcl_composition_v1"
DISPLAY_NAME = "中二高华作文批改"
MARKING_PROFILE = "luna_xhigh"
STAGES = (
    Stage("prepare", "准备作文"), Stage("confirm", "确认学生作文"),
    Stage("mark", "独立 AI 批改"), Stage("validate", "验证批改结果"),
    Stage("write_excel", "写入审核工作簿"),
    Stage("teacher_review", "等待教师审核并保存 Excel", human_gate=True),
    Stage("refresh_review", "重新读取审核结果", execution_wired=True),
    Stage("render", "生成学生体检卡"), Stage("deliver", "查看反馈输出"),
)


def _path(value, base=None):
    path = Path(value)
    return ((base / path) if base is not None and not path.is_absolute() else path).resolve()


def _hash(path):
    return sha256(Path(path).read_bytes()).hexdigest()


@dataclass
class _Entry:
    """Transient workflow-specific evidence; never serialized as source of truth."""
    key: str
    identity: Identity | None = None
    confirmed: bool = False
    grading: bool = False
    record: object = None
    audit: dict | None = None
    rendered: bool = False
    evidence: list = field(default_factory=list)
    attention: list = field(default_factory=list)
    source_paths: set = field(default_factory=set)

    def issue(self, code, message, details="", severity="error"):
        self.attention.append(Attention(code, severity, message, str(details), self.key))


class CompositionWorkflow:
    workflow_id = WORKFLOW_ID
    display_name = DISPLAY_NAME
    stages = STAGES

    def __init__(self, project_dir=ROOT, pipeline_factory=None, render_pipeline_factory=None):
        self.project_dir = Path(project_dir).resolve()
        self.pipeline_factory = pipeline_factory
        self.render_pipeline_factory = render_pipeline_factory

    @staticmethod
    def _scan_workdir(scan: Path) -> Path:
        return scan.parent / f".{safe_name(scan.stem)}.essay-work-{content_digest(scan)[:12]}"

    def prepare_source(self, source: AssignmentSource) -> AssignmentSource:
        """Turn a teacher-selected continuous PDF into an app-owned split pile."""
        if source.continuous_scan is None or source.split_pile is not None:
            return source
        scan = _path(source.continuous_scan)
        if not scan.is_file() or scan.suffix.lower() != ".pdf":
            raise ValidationError("Choose a continuous-scan PDF")

        workdir = self._scan_workdir(scan)
        if not workdir.exists():
            self._build_workdir(scan, workdir)
        else:
            self._check_workdir(scan, workdir)
        return replace(source, split_pile=workdir)

    @staticmethod
    def _check_workdir(scan: Path, workdir: Path) -> None:
        try:
            master = workdir / "_continuous.pdf"
            if not master.is_file() or _hash(master) != _hash(scan):
                raise ValidationError("The generated continuous-PDF copy does not match the selected scan")
            read_existing_split(workdir)
        except Exception as exc:
            raise ValidationError(
                f"The app-owned scan workspace is incomplete or invalid: {workdir}: {exc}"
            ) from exc

    @classmethod
    def _build_workdir(cls, scan: Path, workdir: Path) -> None:
        """Split into a temporary folder and rename it only once it is complete.

        An interrupted split leaves only a ``.partial-`` folder, which the next
        attempt removes, so ``workdir`` never exists half-built.
        """
        for stale in workdir.parent.glob(f"{workdir.name}.partial-*"):
            shutil.rmtree(stale, ignore_errors=True)
        partial = workdir.with_name(f"{workdir.name}.partial-{os.getpid()}-{uuid4().hex[:8]}")
        try:
            split_continuous_anonymous(scan, partial)
            if _hash(partial / "_continuous.pdf") != _hash(scan):
                raise ValidationError("The generated continuous-PDF copy does not match the selected scan")
            read_existing_split(partial)
            try:
                partial.rename(workdir)
            except OSError:
                if not workdir.exists():
                    raise
                # Another window finished the same scan first; use its folder.
                cls._check_workdir(scan, workdir)
        finally:
            shutil.rmtree(partial, ignore_errors=True)

    @staticmethod
    def marking_batch_id(source: AssignmentSource) -> str:
        """The task's identity: the content hash of its continuous scan."""
        return task_id(source)

    @staticmethod
    def _task_paths(source: AssignmentSource):
        """This task's output paths, or None while its scan is unreadable."""
        try:
            return task_output_paths(source)
        except (OSError, ValueError):
            return None

    def marking_source(self, source: AssignmentSource) -> AssignmentSource:
        """Bind a prepared task to its own jobs folder and results workbook.

        The workbook is derived from the task alone; saved job files never
        choose or redirect it. Explicitly attached job roots stay a separate
        import view and are not mixed with the managed workbook.
        """
        paths = self._task_paths(source) if source.split_pile is not None else None
        if paths is None:
            # Without a readable scan there is no task yet; intake reports why.
            return source
        jobs = (self.project_dir / "jobs" / self.marking_batch_id(source)).resolve()
        workbook = source.workbook
        if workbook is None and not source.job_roots:
            workbook = paths.workbook
        if (workbook is not None and _path(workbook) == paths.workbook.resolve()
                and not paths.workbook.is_file()):
            # The task workbook is an output: attach it only once it exists.
            workbook = None
        job_roots = source.job_roots or ((jobs,) if jobs.is_dir() else ())
        return replace(
            source,
            workbook=workbook,
            job_roots=job_roots,
            results_directory=paths.results_directory,
            feedback_cards_directory=paths.feedback_cards_directory,
        )

    def inspect(self, source: AssignmentSource) -> Inspection:
        source = self.prepare_source(source)
        source = self.marking_source(source)
        paths = self._task_paths(source) if source.split_pile is not None else None
        if paths is not None and source.workbook is not None:
            expected = paths.workbook.resolve()
            if _path(source.workbook) != expected:
                return self._finish({}, [Attention(
                    "SOURCE_ASSOCIATION_MISMATCH", "error",
                    "所选工作簿与作文资料不属于同一任务，请分别读取。",
                    f"Selected workbook: {_path(source.workbook)}; expected for this assignment: {expected}",
                )], False)
        entries, issues = {}, []
        try:
            validator = Validator(source.config_dir or ROOT / "config")
        except Exception as exc:
            return self._finish({}, [Attention("CONFIG_INVALID", "error", "批改配置暂时无法读取。", str(exc))], False)
        workbook = _path(source.workbook) if source.workbook is not None else None
        before = None
        if workbook is not None and workbook.is_file():
            try:
                before = _hash(workbook)
                self._workbook(workbook, validator, entries, issues)
            except Exception as exc:
                issues.append(Attention("WORKBOOK_UNREADABLE", "error", "无法读取审核工作簿，请检查文件是否可用。", str(exc)))
        for root in source.job_roots:
            lock = _path(root) / ".pipeline.lock"
            if lock_is_active(lock):
                issues.append(Attention("BATCH_LOCKED", "error", "任务可能仍在运行，请先检查，不要直接重试。", str(lock)))
        self._intake(source, entries, issues)
        self._cards(source, workbook, entries, issues)
        if before is not None:
            try:
                unchanged = before == _hash(workbook)
            except OSError:
                unchanged = False
            if not unchanged:
                issues.append(Attention("WORKBOOK_CHANGED", "error", "审核文件在读取期间发生变化，请保存后重新读取。"))
                for entry in entries.values():
                    entry.record = None
                    entry.rendered = False
        # Suppress actions when assignment-wide evidence is incomplete or ambiguous.
        return self._finish(entries, issues, workbook is not None and workbook.is_file())

    @staticmethod
    def _workbook(path, validator, entries, issues):
        book = load_workbook(path, read_only=True, data_only=False)
        try:
            if (SHEET not in book.sheetnames or AUDIT_SHEET not in book.sheetnames
                    or [c.value for c in book[SHEET][1]] != HEADERS
                    or [c.value for c in book[AUDIT_SHEET][1]] != AUDIT_HEADERS):
                raise ValueError("Unexpected workflow workbook schema")
            rows = {row[0].row: row for row in book[SHEET].iter_rows(min_row=2)
                    if any(c.value is not None for c in row)}
            audits = [row for row in book[AUDIT_SHEET].iter_rows(min_row=2)
                      if any(c.value is not None for c in row)]
            job_counts = Counter(str(row[0].value) for row in audits)
            # Result rows and audit entries pair up by 班级 + 班号, not by row
            # position, so a sorted or filtered teacher sheet stays valid.
            def student(cells):
                return roster_text(cells[0].value), roster_text(cells[1].value)
            student_counts = Counter(student(row) for row in rows.values())
            by_student = defaultdict(list)
            for audit in audits:
                key = (roster_text(audit[1].value), roster_text(audit[2].value))
                if key not in student_counts:
                    issues.append(Attention("AUDIT_ROW_INVALID", "error", "工作簿的审核记录与作文行不一致。", f"Audit row {audit[0].row}: no result row for student {key!r}"))
                else:
                    by_student[key].append(audit)
            store = ExcelStore(path, validator)
            for number, row in rows.items():
                candidates = by_student[student(row)]
                audit = candidates[0] if len(candidates) == 1 else None
                job = audit[0].value if audit else None
                valid_job = (isinstance(job, str) and bool(job.strip()) and job_counts[job] == 1
                             and student_counts[student(row)] == 1)
                key = job if valid_job else f"excel-row:{number}"
                entry = entries[key] = _Entry(key, evidence=[str(path) + f"#row={number}"])
                try:
                    entry.identity = Identity.from_dict(dict(zip(
                        ("class_name", "student_id", "student_name"),
                        (roster_text(c.value) for c in row[:3]),
                    )))
                    if audit is None or not valid_job:
                        raise ValueError("Missing or duplicate audit job/student mapping")
                    if any(c.data_type == "f" for c in audit):
                        raise ValueError("Audit requires literal values, not formulas")
                    if not isinstance(audit[4].value, str) or not audit[4].value.strip():
                        raise ValueError("Audit input digest is missing")
                    record = store.get(job, entry.identity)
                    if record.review_status not in ("PENDING", "APPROVED"):
                        raise ValueError("Unknown approval status")
                    if record.review_status == "APPROVED":
                        record, metadata = store.get_render_source(job, entry.identity)
                    else:
                        metadata = dict(job_id=job, excel_row=number, input_digest=audit[4].value, pipeline_status=audit[5].value)
                    entry.record, entry.audit = record, metadata
                    entry.confirmed = entry.grading = True
                except Exception as exc:
                    entry.issue("WORKBOOK_ROW_INVALID", "这份作文的审核数据需要检查，原文没有被修改。", exc)
        finally:
            book.close()

    @staticmethod
    def _find(roots, filename, issues):
        paths = set()
        for root_value in roots:
            try:
                root = _path(root_value)
                if not root.is_dir():
                    raise FileNotFoundError(f"Declared artifact directory missing: {root}")
                paths.update(p.resolve() for p in root.rglob(filename))
            except Exception as exc:
                issues.append(Attention("ARTIFACT_ROOT_UNREADABLE", "error", "部分任务资料无法读取，请检查文件夹。", str(exc)))
        return sorted(paths)

    @staticmethod
    def _intake(source, entries, issues):
        if source.split_pile is None:
            if source.roster is not None or source.identity_decisions is not None:
                issues.append(Attention("INTAKE_SOURCE_MISSING", "error", "请同时指定作文资料文件夹。"))
            return
        try:
            pile = _path(source.split_pile)
            submissions = read_existing_split(pile)
        except Exception as exc:
            issues.append(Attention("INTAKE_INVALID", "error", "作文资料不完整或页序需要检查。", str(exc)))
            return
        decisions = {}
        if source.roster is not None and source.identity_decisions is not None:
            try:
                records = apply_identity_decisions(
                    submissions,
                    read_roster(source.roster, source.roster_sheet),
                    read_json(source.identity_decisions),
                    require_complete=False,
                )
                decisions = {r["source_pdf"]: r for r in records}
            except Exception as exc:
                issues.append(Attention("IDENTITY_DECISIONS_INVALID", "error", "学生确认记录需要检查。", str(exc)))
        for submission in submissions:
            pdf = (pile / submission.source_pdf).resolve()
            record = decisions.get(submission.source_pdf)
            identity = (Identity.from_dict({k: record[k] for k in ("class_name", "student_id", "student_name")})
                        if record else None)
            confirmed = bool(record and record["match_status"] == "STRONG_ROSTER_MATCH")
            # A confirmed essay's saved result is the workbook row keyed by the
            # same student; decisions allow each student only once per task.
            key = job_key(identity) if confirmed else f"submission:{pdf}"
            entry = entries.setdefault(key, _Entry(key))
            entry.evidence.append(str(pdf))
            entry.source_paths.add(pdf)
            if confirmed:
                if entry.identity is not None and identity != entry.identity:
                    entry.issue("IDENTITY_MISMATCH", "学生确认信息与保存记录不一致。", str(pdf))
                else:
                    entry.identity, entry.confirmed = identity, True
                # Another workbook row for this student that is not this
                # essay's audited row cannot be overwritten, so marking stops.
                student = (identity.class_name, identity.student_id)
                unlinked = [other.key for other in entries.values() if other is not entry
                            and other.identity is not None
                            and (other.identity.class_name, other.identity.student_id) == student]
                if unlinked:
                    entry.issue("WORKBOOK_ROW_INVALID", "这份作文的审核数据需要检查，原文没有被修改。",
                                f"Workbook rows for this student not linked to this essay: {unlinked}")
            else:
                if identity is not None:
                    entry.identity = identity
                entry.confirmed = False
                entry.issue("IDENTITY_UNRESOLVED", "请确认这份作文对应的学生。", str(pdf), "warning")

    def _cards(self, source, workbook, entries, issues):
        from .card_evidence import inspect_card
        from .feedback_rendering import feedback_card_filename
        roots = tuple(dict.fromkeys((*source.receipt_roots, *source.job_roots)))
        paths = self._find(roots, "render_receipt.json", issues)
        found_for = set()
        for path in paths:
            try:
                receipt = read_json(path)
                key = receipt["source"]["job_id"]
                entry = entries.get(key)
                if entry is None:
                    # A directory name can localize an error, never establish provenance.
                    folder_key = path.parent.parent.name if path.parent.name == "output" else path.parent.name
                    candidate = entries.get(folder_key)
                    if candidate:
                        candidate.issue("RECEIPT_JOB_MISMATCH", "体检卡生成记录的作文标识不一致。", f"Receipt job mismatch: {path}: {key}")
                    else:
                        issues.append(Attention("RECEIPT_UNRESOLVED", "error", "部分体检卡无法对应到已确认的审核记录。", f"Receipt audit job mismatch: {path}: {key}"))
                    continue
                found_for.add(entry.key)
                entry.evidence.append(str(path))
                if entry.record is None or entry.record.review_status != "APPROVED" or workbook is None:
                    entry.issue("CARD_SOURCE_UNAPPROVED", "已有体检卡，但当前审核数据尚未确认有效。", str(path))
                    continue
                expected_teacher_output = None
                paths = self._task_paths(source) if source.split_pile is not None else None
                if paths is not None and workbook == paths.workbook.resolve():
                    expected_teacher_output = paths.feedback_cards_directory / feedback_card_filename(entry.record)
                okay, code, detail = inspect_card(
                    path, workbook, entry.key,
                    entry.audit["input_digest"], entry.record.to_dict(),
                    expected_teacher_output=expected_teacher_output,
                )
                if okay:
                    entry.rendered = True
                else:
                    entry.issue(code, "体检卡与当前审核资料不一致或文件缺失，请检查。", f"{path}: {detail}")
            except Exception as exc:
                issues.append(Attention("RECEIPT_INVALID", "error", "部分体检卡的生成记录无法读取。", f"{path}: {exc}"))
        pngs = set(self._find(roots, "student_card.png", issues))
        pngs.update(self._find(roots, "*-作文体检卡.png", issues))
        for png in sorted(pngs):
            if (png.parent / "render_receipt.json").is_file():
                continue
            key = png.parent.parent.name if png.parent.name == "output" else png.parent.name
            entry = entries.get(key)
            if entry:
                entry.issue("RECEIPT_MISSING", "已找到体检卡，但缺少对应生成记录。", str(png))
            else:
                issues.append(Attention("ORPHAN_CARD", "warning", "部分图片缺少可核对的生成记录。", str(png)))
        for entry in entries.values():
            # Conflicting receipts need an explicit choice; don't advertise a mixed result.
            if any(a.code.startswith(("CARD_", "RECEIPT_")) for a in entry.attention):
                entry.rendered = False

    @staticmethod
    def _finish(entries, issues, has_workbook):
        issues = list(dict.fromkeys(issues))
        global_error = any(a.severity == "error" for a in issues)
        submissions = []
        targets = defaultdict(list)
        for entry in entries.values():
            blocked = global_error or any(a.severity == "error" for a in entry.attention)
            review = entry.record.review_status if entry.record else None
            ready = bool(entry.record and review == "APPROVED" and not entry.rendered and not blocked)
            actions = []
            if not entry.confirmed:
                actions.append(Action.CONFIRM_SUBMISSIONS)
            if entry.confirmed and entry.record is None and not blocked:
                actions.append(Action.RUN_MARKING)
            if review == "PENDING":
                actions.append(Action.REVIEW_EXCEL)
            if ready:
                actions.append(Action.RENDER_APPROVED)
            if entry.rendered:
                actions.append(Action.VIEW_OUTPUTS)
            if entry.attention:
                actions.append(Action.RESOLVE_ATTENTION)
            for action in actions:
                targets[action].append(entry.key)
            identity = entry.identity
            submissions.append(SubmissionState(
                submission_id=entry.key,
                student_name=identity.student_name if identity else None,
                class_name=identity.class_name if identity else None,
                student_id=identity.student_id if identity else None,
                identity_confirmed=entry.confirmed, grading_available=entry.grading,
                workbook_valid=entry.record is not None, review_status=review,
                rendered=entry.rendered, ready_to_render=ready,
                evidence=tuple(dict.fromkeys(entry.evidence)), attention=tuple(entry.attention), next_actions=tuple(actions),
            ))
        if has_workbook:
            targets[Action.REFRESH_REVIEW_STATUS] = [e.key for e in entries.values() if e.record is not None]
        if issues:
            targets.setdefault(Action.RESOLVE_ATTENTION, [])
        if not entries and not issues and not has_workbook:
            targets[Action.PREPARE_SUBMISSIONS] = []
        return Inspection(WORKFLOW_ID, DISPLAY_NAME, STAGES, tuple(submissions), tuple(issues), tuple(
            AvailableAction(action, tuple(targets[action]), action in (
                Action.REFRESH_REVIEW_STATUS, Action.RUN_MARKING, Action.RENDER_APPROVED
            ))
            for action in Action if action in targets
        ))

    def _feedback_source(self, source):
        """Bind rendering to this app's matching output/jobs batch only."""
        source = self.prepare_source(source)
        source = self.marking_source(source)
        if source.workbook is None:
            raise ValidationError("Read the approved results workbook before generating feedback")
        workbook = _path(source.workbook)
        if source.split_pile is not None:
            batch_id = self.marking_batch_id(source)
            expected_workbook = task_output_paths(source).workbook.resolve()
            if workbook != expected_workbook:
                raise ValidationError("The selected workbook does not match this task's managed results location")
        else:
            output_root = (self.project_dir / "output").resolve()
            try:
                relative = workbook.relative_to(output_root)
            except ValueError as exc:
                raise ValidationError("Feedback output is limited to this project's own results workbook") from exc
            if len(relative.parts) != 2 or relative.parts[1].casefold() != "results.xlsx":
                raise ValidationError("The selected workbook is not in this project's task output folder")
            batch_id = relative.parts[0]
            expected_workbook = (output_root / batch_id / "results.xlsx").resolve()
        expected_jobs = (self.project_dir / "jobs" / batch_id).resolve()
        if workbook != expected_workbook:
            raise ValidationError("The selected workbook does not match the task output location")
        feedback_dir = source.feedback_cards_directory
        if feedback_dir is None:
            feedback_dir = workbook.parent / "Feedback Cards"
        if expected_jobs.is_dir():
            # This task's deterministic Stage 3B job tree is the only receipt
            # root used for a Stage 3C write. Unrelated chosen roots cannot be
            # used to redirect card output.
            source = replace(source, workbook=workbook, job_roots=(expected_jobs,),
                             receipt_roots=(expected_jobs,),
                             results_directory=workbook.parent,
                             feedback_cards_directory=feedback_dir)
        else:
            source = replace(source, workbook=workbook, job_roots=(), receipt_roots=(),
                             results_directory=workbook.parent,
                             feedback_cards_directory=feedback_dir)
        return source, batch_id, expected_jobs

    def _feedback_pipeline(self, source, batch_id):
        if self.render_pipeline_factory is not None:
            return self.render_pipeline_factory(
                project_dir=self.project_dir, batch_id=batch_id, source=source
            )
        from .feedback_rendering import ApprovedFeedbackRenderer
        from rendering.renderer import PillowRenderer
        from workflow.pipeline import Pipeline

        publish_dir = source.feedback_cards_directory if source.split_pile is not None else None
        return Pipeline(
            project_dir=self.project_dir,
            batch_id=batch_id,
            config_dir=source.config_dir or ROOT / "config",
            renderer=ApprovedFeedbackRenderer(PillowRenderer(), publish_dir),
            allow_real=True,
            workbook_path=(source.workbook or
                           ((source.results_directory / "results.xlsx") if source.results_directory else None)),
        )

    def execute_feedback_generation(self, source: AssignmentSource, *, progress_callback=None):
        """Render only fresh, approved Excel rows through the proven renderer."""
        source, batch_id, jobs_root = self._feedback_source(source)
        preliminary = self.inspect(source)
        if preliminary.summary["blocking_errors"]:
            raise ValidationError("Resolve the task evidence warnings before generating feedback")
        preliminary_action = next((item for item in preliminary.available_actions
                                   if item.action == Action.RENDER_APPROVED), None)
        if preliminary_action is None or not preliminary_action.submission_ids:
            raise ValidationError("There are no approved essays ready for feedback generation")

        # Serialize app-level batches, then re-inspect while holding the lock.
        # A second window that acted on stale readiness cannot overwrite a card.
        with batch_lock(jobs_root / ".stage3c-render.lock"):
            source = replace(source, job_roots=(jobs_root,), receipt_roots=(jobs_root,))
            return self._execute_feedback_generation_locked(
                source, batch_id, jobs_root, progress_callback=progress_callback
            )

    def _execute_feedback_generation_locked(self, source, batch_id, jobs_root, *, progress_callback=None):
        from .feedback_rendering import feedback_card_filename_for_identity
        from workflow.pipeline import job_key

        inspection = self.inspect(source)
        if inspection.summary["blocking_errors"]:
            raise ValidationError("Resolve the task evidence warnings before generating feedback")
        action = next((item for item in inspection.available_actions
                       if item.action == Action.RENDER_APPROVED), None)
        if action is None or not action.submission_ids:
            raise ValidationError("There are no approved essays ready for feedback generation")
        by_id = {item.submission_id: item for item in inspection.submissions}
        targets = []
        for submission_id in action.submission_ids:
            item = by_id.get(submission_id)
            if (item is None or not item.ready_to_render or item.review_status != "APPROVED"
                    or not item.class_name or not item.student_id or not item.student_name):
                raise ValidationError("An eligible approved Excel row no longer matches the task")
            identity = Identity.from_dict({
                "class_name": item.class_name,
                "student_id": item.student_id,
                "student_name": item.student_name,
            })
            targets.append((submission_id, identity))

        names = [feedback_card_filename_for_identity(identity)
                 for _submission_id, identity in targets]
        if len(names) != len(set(names)):
            raise ValidationError("Two approved submissions would use the same feedback-card filename")

        pipeline = self._feedback_pipeline(source, batch_id)
        completed_ids = set()
        successes, failures = [], []
        total = len(targets)

        def report(active=""):
            if progress_callback is None:
                return
            completed = len(completed_ids)
            progress_callback({
                "operation": "render",
                "total": total,
                "completed": completed,
                "active": 1 if active else 0,
                "waiting": max(0, total - completed - len(failures) - (1 if active else 0)),
                "attention": len(failures),
                "current": active,
                "interrupted": False,
            })

        report()
        for submission_id, identity in targets:
            report(f"{identity.class_name} · {identity.student_id} · {identity.student_name}")
            job_dir = jobs_root / job_key(identity)
            output_dir = job_dir / "output"
            try:
                receipt_path = Path(pipeline.rerender(identity)).resolve()
                receipt = read_json(receipt_path)
                artifact = Path(receipt["output"]["artifact"]).resolve()
                expected_filename = feedback_card_filename_for_identity(identity)
                expected_output = (output_dir / expected_filename).resolve()
                if artifact != expected_output or not artifact.is_file():
                    raise ValidationError("The renderer did not save the expected student card")
                if source.split_pile is not None:
                    expected_teacher_output = ((source.feedback_cards_directory or
                                                (source.results_directory / "Feedback Cards")) / expected_filename).resolve()
                    teacher_output = receipt.get("teacher_output")
                    output_record = receipt.get("output") or {}
                    should_publish = (
                        not isinstance(teacher_output, dict)
                        or not isinstance(teacher_output.get("artifact"), str)
                        or _path(teacher_output.get("artifact")) != expected_teacher_output
                        or not expected_teacher_output.is_file()
                        or teacher_output.get("sha256") != output_record.get("sha256")
                        or _hash(expected_teacher_output) != output_record.get("sha256")
                    )
                    if should_publish:
                        from .feedback_rendering import publish_card_receipt
                        publish_card_receipt(receipt_path, expected_teacher_output.parent, expected_filename)
                        receipt = read_json(receipt_path)
                else:
                    expected_teacher_output = artifact

                if jobs_root.is_dir() and jobs_root not in source.job_roots:
                    source = replace(source, job_roots=(jobs_root,), receipt_roots=(jobs_root,))
                fresh = self.inspect(source)
                rendered = next((row for row in fresh.submissions
                                 if row.submission_id == submission_id and row.rendered), None)
                if rendered is None:
                    detail = next((f"{attention.code}: {attention.technical_details}"
                                   for row in fresh.submissions if row.submission_id == submission_id
                                   for attention in row.attention), "The saved card did not pass Stage 1 verification")
                    raise ValidationError(detail)
                completed_ids.add(submission_id)
                success = {"submission_id": submission_id, "identity": identity.to_dict(),
                           "path": str(expected_teacher_output), "receipt": str(receipt_path)}
                successes.append(success)
                try:
                    atomic_json(output_dir / "render_attempt.json", {
                        "status": "PASS", "identity": identity.to_dict(),
                        "workbook": str(source.workbook), "artifact": str(artifact),
                        "receipt": str(receipt_path),
                        "completed_at": datetime.now(timezone.utc).isoformat(),
                    })
                except Exception as log_error:
                    # A verified receipt/PNG is the completion evidence; the
                    # optional attempt log cannot turn success into failure.
                    success["attempt_log_warning"] = str(log_error)
            except Exception as exc:
                failure = {
                    "submission_id": submission_id,
                    "identity": identity.to_dict(),
                    "error": f"{type(exc).__name__}: {exc}",
                    "message": "This card could not be generated. The approved Excel result is unchanged; you can retry after checking the details.",
                }
                failures.append(failure)
                try:
                    output_dir.mkdir(parents=True, exist_ok=True)
                    atomic_json(output_dir / "render_attempt.json", {
                        "status": "FAILED", "identity": identity.to_dict(),
                        "workbook": str(source.workbook), "error_type": type(exc).__name__,
                        "error": str(exc), "failed_at": datetime.now(timezone.utc).isoformat(),
                    })
                except Exception as log_error:
                    failure["error"] += f"; failure log could not be saved: {log_error}"
            report()

        return {"result": {
            "batch_id": batch_id,
            "requested": total,
            "generated": len(successes),
            "successes": successes,
            "failed": failures,
        }, "source": source}

    def execute_marking(self, source: AssignmentSource, *, progress_callback=None,
                        cancelled=None):
        """Run the existing one-student production batch behind the app seam."""
        from grading.schemas import ValidationError
        from scanning.intake import apply_identity_decisions, read_existing_split
        source = self.prepare_source(source)
        if source.split_pile is None or source.roster is None or source.identity_decisions is None:
            raise ValidationError("Choose a prepared composition folder, roster and confirmed student identities first")

        source = self.marking_source(source)
        inspection = self.inspect(source)
        if inspection.summary["blocking_errors"]:
            raise ValidationError("The selected task has a blocking file or identity issue. Read the task details before marking")
        if not inspection.submissions or any(not item.identity_confirmed for item in inspection.submissions):
            raise ValidationError("Confirm the student for every submission before marking")

        batch_id = self.marking_batch_id(source)
        expected_jobs = (self.project_dir / "jobs" / batch_id).resolve()
        expected_workbook = task_output_paths(source).workbook.resolve()
        if any(_path(root) != expected_jobs for root in source.job_roots):
            raise ValidationError("This assignment is linked to other saved marking records; read those records before starting a new run")
        if source.workbook is not None and _path(source.workbook) != expected_workbook:
            raise ValidationError("This assignment already names a different results workbook; the app will not replace it")
        execution_source = replace(source, workbook=expected_workbook, job_roots=(expected_jobs,))

        submissions = read_existing_split(source.split_pile)
        records = apply_identity_decisions(
            submissions,
            read_roster(source.roster, source.roster_sheet),
            read_json(source.identity_decisions),
            require_complete=True,
        )
        by_student = defaultdict(list)
        for item in inspection.submissions:
            if item.class_name and item.student_id:
                by_student[(item.class_name, item.student_id)].append(item)
        if any(len(items) != 1 for items in by_student.values()):
            raise ValidationError("Each student must be linked to only one submission before marking")

        eligible = []
        for record in records:
            items = by_student.get((record["class_name"], record["student_id"]), [])
            if not items:
                raise ValidationError("A confirmed submission is missing from the inspected task")
            item = items[0]
            # A valid workbook row is DONE. Everything else is ordinary
            # unfinished work, including interrupted/validated attempts.
            if not item.workbook_valid and Action.RUN_MARKING in item.next_actions:
                eligible.append(record)
        records = eligible

        if self.pipeline_factory is None:
            from workflow.pipeline import Pipeline
            pipeline = Pipeline(
                self.project_dir, batch_id, config_dir=source.config_dir or (ROOT / "config"),
                allow_real=True, workbook_path=expected_workbook,
            )
        else:
            pipeline = self.pipeline_factory(
                self.project_dir, batch_id, source.config_dir or (ROOT / "config"),
                workbook_path=expected_workbook,
            )
        result = pipeline.run_real_batch(
            records,
            source.split_pile,
            essay_question=None,
            model_profile=MARKING_PROFILE,
            progress_callback=progress_callback,
            cancelled=cancelled,
        )
        return {"result": result, "source": self.marking_source(execution_source)}
