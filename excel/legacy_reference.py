"""Read-only extraction of historical teacher records with cell provenance."""
from openpyxl import load_workbook

from grading.schemas import ValidationError


LEGACY_HEADERS = (
    "班号", "学生姓名", "班级", "题号", "内容分/30", "语文与结构/30", "总分/60",
    "审题扣题", "情节紧凑", "重点突出", "详略得当", "情景交融", "多感官", "修辞运用",
    "结尾照应/升华", "教师点评", "给Chat GPT的一些指示",
)


def read_legacy_reference(path, identities, sheet_name="作文诊断输入"):
    """Return exact historical values; no rating conversion or scan linkage."""
    formulas = load_workbook(path, read_only=True, data_only=False)
    cached = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = formulas[sheet_name]
        values_sheet = cached[sheet_name]
        headers = tuple(cell.value for cell in next(sheet.iter_rows(min_row=1, max_row=1, max_col=17)))
        if headers != LEGACY_HEADERS:
            raise ValidationError("Unexpected legacy teacher columns")
        wanted = {(item.class_name, item.student_id): item for item in identities}
        found = {}
        formula_rows = sheet.iter_rows(min_row=2, max_col=17, values_only=True)
        cached_rows = values_sheet.iter_rows(min_row=2, max_col=17, values_only=True)
        for row_number, (row_values, cached_values) in enumerate(zip(formula_rows, cached_rows), 2):
            row = list(row_values)
            key = (str(row[2]).removesuffix(".0"), str(row[0]).removesuffix(".0"))
            if key not in wanted:
                continue
            if key in found:
                raise ValidationError(f"Duplicate legacy teacher record: {key}")
            record = dict(zip(LEGACY_HEADERS, row))
            record["总分/60_显示值"] = cached_values[6]
            record["来源工作表"] = sheet_name
            record["来源行"] = row_number
            found[key] = record
        missing = set(wanted) - set(found)
        if missing:
            raise ValidationError(f"Missing legacy teacher records: {sorted(missing)}")
        return [found[key] for key in wanted]
    finally:
        formulas.close()
        cached.close()
