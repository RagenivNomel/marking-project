"""Batch-friendly Excel repository; approved wording is read back verbatim."""
from copy import copy
from dataclasses import asdict, dataclass
from datetime import datetime, timezone
from pathlib import Path
import math
import os
import shutil
import tempfile
import threading

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter

from grading.schemas import CRITERIA, RATINGS, CriterionResult, Identity, ValidationError
from workflow.storage import replace_file
from .schema import AUDIT_HEADERS, AUDIT_SHEET, HEADERS, SHEET


APPROVAL_STATUSES = ("PENDING", "APPROVED")
APPROVAL_VALIDATION_RANGE = "Y2:Y1048576"
TOTAL_COLUMN = HEADERS.index("总分") + 1


def _total_formula(row_index):
    return f"=E{row_index}+F{row_index}"

_WORKBOOK_LOCKS = {}
_WORKBOOK_LOCKS_GUARD = threading.Lock()


def _workbook_lock(path):
    key = str(Path(path).resolve()).casefold()
    with _WORKBOOK_LOCKS_GUARD:
        return _WORKBOOK_LOCKS.setdefault(key, threading.RLock())


def roster_text(value):
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip() if value is not None else ""


ROSTER_NAME_HEADERS = ("姓名", "学生姓名")


def _roster_columns(headers):
    """Map 班级/班号/姓名 to column indices, or None if this is not a roster sheet."""
    headers = [roster_text(value) for value in headers]
    names = [header for header in ROSTER_NAME_HEADERS if header in headers]
    if "班号" not in headers or not names:
        return None
    if len(names) > 1:
        raise ValidationError("名册的第一行同时有“姓名”和“学生姓名”，请只保留一个。")
    for header in ("班级", "班号", names[0]):
        if headers.count(header) > 1:
            raise ValidationError(f"名册的第一行有两个“{header}”列，请只保留一个。")
    return {
        "class_name": headers.index("班级") if "班级" in headers else None,
        "student_id": headers.index("班号"),
        "student_name": headers.index(names[0]),
    }


def read_roster(path, sheet_name=None, class_name=None):
    """Read 班级, 班号 and 姓名 (or 学生姓名) from any Excel workbook.

    Every sheet whose first row has these headings is read and every other
    column or sheet is ignored, so one sheet per class also works. A student
    listed twice with the same name counts once; the same 班级 + 班号 with a
    different name is refused. A sheet without a 班级 column takes the class
    the teacher gives once as class_name. sheet_name restricts the search.
    """
    class_name = roster_text(class_name) or None
    book = load_workbook(path, read_only=True, data_only=True)
    try:
        if sheet_name is not None and sheet_name not in book.sheetnames:
            raise ValidationError(f"名册里没有名为“{sheet_name}”的工作表。")
        found = []
        for sheet in ([book[sheet_name]] if sheet_name is not None else book.worksheets):
            header = next(sheet.iter_rows(max_row=1, values_only=True), ())
            columns = _roster_columns(header)
            if columns is not None:
                found.append((sheet, columns))
        if not found:
            raise ValidationError("名册里找不到学生名单。第一行需要有“班号”和“姓名”两列（“班级”列可选）。")
        without_class = [sheet.title for sheet, columns in found if columns["class_name"] is None]
        if without_class and class_name is None:
            raise ValidationError(f"名册的工作表“{'、'.join(without_class)}”没有“班级”列，请填写班级。")
        if not without_class and class_name is not None:
            raise ValidationError("名册已有“班级”列，请不要另外填写班级。")
        result, seen = [], {}
        for sheet, columns in found:
            for number, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), 2):
                values = {key: roster_text(row[index]) if index is not None and index < len(row) else ""
                          for key, index in columns.items()}
                if not any(values.values()):
                    continue
                if columns["class_name"] is None:
                    values["class_name"] = class_name
                place = f"工作表“{sheet.title}”第 {number} 行"
                missing = [label for key, label in (("class_name", "班级"), ("student_id", "班号"), ("student_name", "姓名"))
                           if not values[key]]
                if missing:
                    raise ValidationError(f"名册{place}缺少{'、'.join(missing)}。")
                identity = Identity.from_dict(values)
                key = (identity.class_name, identity.student_id)
                if key in seen:
                    earlier, earlier_place = seen[key]
                    if earlier.student_name != identity.student_name:
                        raise ValidationError(
                            f"名册里 {key[0]} 班 {key[1]} 号有两个不同的名字：{earlier_place}是“{earlier.student_name}”，"
                            f"{place}是“{identity.student_name}”。")
                    continue
                seen[key] = (identity, place)
                result.append(identity)
        if not result:
            raise ValidationError("名册里没有学生。")
        return result
    finally:
        book.close()


@dataclass(frozen=True)
class ExcelStudentRecord:
    student_id: str
    student_name: str
    class_name: str
    topic: str
    content_score: float | None
    language_structure_score: float | None
    total_score: float | None
    criteria: dict[str, CriterionResult]
    teacher_comment: str
    review_status: str

    def to_dict(self):
        return asdict(self)


class ExcelStore:
    def __init__(self, path, validator, *, backup_dir=None):
        self.path = Path(path)
        self.validator = validator
        self.backup_dir = Path(backup_dir) if backup_dir is not None else None
        # Separate stores targeting the same authoritative workbook share the
        # complete read/modify/save critical section.
        self._write_lock = _workbook_lock(self.path)

    def _open(self):
        if self.path.exists():
            book = load_workbook(self.path)
            valid = (SHEET in book.sheetnames and AUDIT_SHEET in book.sheetnames
                     and [c.value for c in book[SHEET][1]] == HEADERS
                     and [c.value for c in book[AUDIT_SHEET][1]] == AUDIT_HEADERS)
            if not valid:
                book.close()
                raise ValidationError("Unexpected results workbook schema")
            self._ensure_approval_validation(book[SHEET])
            return book
        book = Workbook()
        sheet = book.active
        sheet.title = SHEET
        sheet.append(HEADERS)
        sheet.freeze_panes = "H2"
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="285E61")
            cell.alignment = Alignment(wrap_text=True, vertical="center")
        sheet.row_dimensions[1].height = 42
        for index, header in enumerate(HEADERS, 1):
            wide = header.endswith("评语") or header == "教师总评" or header == "题目"
            sheet.column_dimensions[get_column_letter(index)].width = 34 if wide else 16
        for index in range(8, 24, 2):
            validation = DataValidation(type="list", formula1='"' + ','.join(RATINGS) + '"')
            validation.errorTitle = "等级无效"
            validation.error = "请选择指定的三个等级之一。"
            validation.showErrorMessage = True
            validation.errorStyle = "stop"
            sheet.add_data_validation(validation)
            validation.add(f"{get_column_letter(index)}2:{get_column_letter(index)}1048576")
        self._ensure_approval_validation(sheet)
        audit = book.create_sheet(AUDIT_SHEET)
        audit.append(AUDIT_HEADERS)
        audit.sheet_state = "hidden"
        return book

    @staticmethod
    def _ensure_approval_validation(sheet):
        formula = '"' + ','.join(APPROVAL_STATUSES) + '"'
        for validation in sheet.data_validations.dataValidation:
            if (validation.type == "list" and validation.formula1 == formula
                    and str(validation.sqref) == APPROVAL_VALIDATION_RANGE):
                return
        validation = DataValidation(type="list", formula1=formula, allow_blank=False)
        validation.errorTitle = "审核状态无效"
        validation.error = "请选择 PENDING 或 APPROVED。"
        validation.showErrorMessage = True
        validation.errorStyle = "stop"
        sheet.add_data_validation(validation)
        validation.add(APPROVAL_VALIDATION_RANGE)

    def ensure_validation_rules(self):
        """Persist current production validation rules without changing cell values."""
        with self._write_lock:
            return self._ensure_validation_rules_unlocked()

    def _ensure_validation_rules_unlocked(self):
        if not self.path.exists():
            raise ValidationError("Results workbook does not exist")
        book = self._open()
        try:
            self._save(book)
        finally:
            book.close()

    def _save(self, book):
        self.path.parent.mkdir(parents=True, exist_ok=True)
        fd, temporary = tempfile.mkstemp(dir=self.path.parent, suffix=".xlsx")
        os.close(fd)
        try:
            book.save(temporary)
            if self.path.exists():
                backups = self.backup_dir or (self.path.parent / "backups")
                backups.mkdir(parents=True, exist_ok=True)
                stamp = datetime.now(timezone.utc).strftime("%Y%m%dT%H%M%S%f")
                shutil.copy2(self.path, backups / f"{self.path.stem}-{stamp}.xlsx")
            replace_file(temporary, self.path)
        finally:
            Path(temporary).unlink(missing_ok=True)

    @staticmethod
    def _student_row(book, class_name, student_id):
        # Students are found by 班级 + 班号, never by position, so the teacher
        # sheet may be sorted or filtered freely. The Excel行 audit column is
        # kept only as a readable hint.
        matches = [row for row in book[SHEET].iter_rows(min_row=2)
                   if (roster_text(row[0].value), roster_text(row[1].value)) == (class_name, student_id)]
        if len(matches) > 1:
            raise ValidationError("Duplicate student rows in Excel results sheet")
        if not matches:
            raise ValidationError("Excel audit student has no result row")
        return matches[0]

    @classmethod
    def _row(cls, book, job_id):
        matches = [row for row in book[AUDIT_SHEET].iter_rows(min_row=2) if row[0].value == job_id]
        if len(matches) > 1:
            raise ValidationError("Duplicate job rows in Excel audit sheet")
        if not matches:
            return None, None
        audit = matches[0]
        return cls._student_row(book, roster_text(audit[1].value), roster_text(audit[2].value)), audit

    @staticmethod
    def _score(value, label):
        if value is None or value == "":
            return None
        if type(value) not in (int, float) or not math.isfinite(value) or value < 0:
            raise ValidationError(f"{label}: blank or finite nonnegative number required")
        return value

    def ensure_draft(self, result, identity, job_id, digest, topic="", teacher_scores=None):
        result = self.validator.validate(result.to_dict(), identity)
        if not isinstance(topic, str):
            raise ValidationError("题目 must be text")
        teacher_scores = {} if teacher_scores is None else teacher_scores
        score_keys = {"content_score", "language_structure_score", "total_score"}
        if not isinstance(teacher_scores, dict) or not set(teacher_scores).issubset(score_keys):
            raise ValidationError("teacher_scores: unexpected score field")
        content_score = self._score(teacher_scores.get("content_score"), "内容分")
        language_structure_score = self._score(teacher_scores.get("language_structure_score"), "语文与结构分")
        total_score = self._score(teacher_scores.get("total_score"), "总分")
        with self._write_lock:
            return self._ensure_draft_unlocked(
                result, identity, job_id, digest, topic,
                content_score, language_structure_score, total_score,
            )

    def _ensure_draft_unlocked(self, result, identity, job_id, digest, topic,
                               content_score, language_structure_score, total_score):
        book = self._open()
        try:
            sheet = book[SHEET]
            existing, audit_row = self._row(book, job_id)
            if existing:
                if audit_row[4].value != digest:
                    raise ValidationError("Existing Excel row belongs to different input")
                self._decode(existing, identity)
                return
            for row in sheet.iter_rows(min_row=2):
                if not any(cell.value is not None for cell in row):
                    continue
                if (roster_text(row[0].value), roster_text(row[1].value)) == (
                    identity.class_name, identity.student_id,
                ):
                    raise ValidationError("Existing Excel row for this student lacks matching audit evidence")
            for audit in book[AUDIT_SHEET].iter_rows(min_row=2):
                if not any(cell.value is not None for cell in audit):
                    continue
                if (roster_text(audit[1].value), roster_text(audit[2].value)) == (
                        identity.class_name, identity.student_id):
                    raise ValidationError("Existing Excel audit evidence conflicts with this result")
            row_index = sheet.max_row + 1
            if total_score is None and content_score is not None and language_structure_score is not None:
                total_score = _total_formula(row_index)
            values = [result.class_name, result.student_id, result.student_name, topic,
                      content_score, language_structure_score, total_score]
            for name in CRITERIA:
                values += [result.criteria[name].rating, result.criteria[name].short_comment]
            values += [result.teacher_comment, "PENDING"]
            for col, value in enumerate(values, 1):
                cell = sheet.cell(row_index, col, value)
                if isinstance(value, str) and col != TOTAL_COLUMN:
                    cell.data_type = "s"
                cell.alignment = Alignment(wrap_text=True, vertical="top")
            sheet.row_dimensions[row_index].height = 120
            sheet.auto_filter.ref = sheet.dimensions
            book[AUDIT_SHEET].append([job_id, identity.class_name, identity.student_id, row_index, digest, "VALIDATED"])
            self._save(book)
        finally:
            book.close()

    def _decode(self, row, identity):
        # Only 总分 may be a formula, and only the exact =E<row>+F<row> sum, which
        # is recomputed here because openpyxl does not evaluate formulas.
        values = [cell.value for cell in row]
        total_is_formula = False
        for cell in row:
            if cell.data_type != "f":
                continue
            if cell.column != TOTAL_COLUMN or str(cell.value).replace(" ", "").upper() != _total_formula(cell.row):
                raise ValidationError("Result rows require literal approved values; only 总分 may be =E+F")
            total_is_formula = True
        if total_is_formula:
            content = self._score(values[4], "内容分")
            language = self._score(values[5], "语文与结构分")
            values[6] = None if content is None or language is None else content + language
        grading = dict(
            class_name=roster_text(values[0]), student_id=roster_text(values[1]), student_name=roster_text(values[2]),
            criteria={name: {"rating": values[7 + 2*i], "short_comment": values[8 + 2*i]} for i, name in enumerate(CRITERIA)},
            teacher_comment=values[23],
        )
        validated = self.validator.validate(grading, identity)
        return ExcelStudentRecord(
            validated.student_id, validated.student_name, validated.class_name, values[3] or "",
            self._score(values[4], "内容分"), self._score(values[5], "语文与结构分"), self._score(values[6], "总分"),
            validated.criteria, validated.teacher_comment, values[24],
        )

    def get(self, job_id, identity, approved=False):
        if not self.path.exists():
            raise ValidationError("Results workbook does not exist")
        book = self._open()
        try:
            row, _ = self._row(book, job_id)
            if row is None:
                raise ValidationError("Student result not found")
            record = self._decode(row, identity)
            if approved and record.review_status != "APPROVED":
                raise ValidationError("Student result is not APPROVED in Excel")
            return record
        finally:
            book.close()

    def get_render_source(self, preferred_job_id, identity):
        """Return an approved row and its workbook audit metadata for rendering.

        Normal production rows are addressed by the pipeline job key. A workbook
        imported from an existing calibration job may retain that job ID instead;
        in that case an unambiguous class/student audit match is sufficient. The
        approved Excel row remains the sole feedback source in both cases.
        """
        if not self.path.exists():
            raise ValidationError("Results workbook does not exist")
        book = self._open()
        try:
            audit = book[AUDIT_SHEET]
            preferred = [row for row in audit.iter_rows(min_row=2) if row[0].value == preferred_job_id]
            if len(preferred) > 1:
                raise ValidationError("Duplicate job rows in Excel audit sheet")
            candidates = preferred
            if not candidates:
                candidates = [
                    row for row in audit.iter_rows(min_row=2)
                    if roster_text(row[1].value) == identity.class_name
                    and roster_text(row[2].value) == identity.student_id
                ]
            if len(candidates) != 1:
                raise ValidationError("Approved Excel render source is not uniquely identified")
            audit_row = candidates[0]
            row = self._student_row(book, roster_text(audit_row[1].value), roster_text(audit_row[2].value))
            row_number = row[0].row
            record = self._decode(row, identity)
            if record.review_status != "APPROVED":
                raise ValidationError("Student result is not APPROVED in Excel")
            digest = audit_row[4].value
            if not isinstance(digest, str) or not digest:
                raise ValidationError("Excel audit input digest is missing")
            return record, {
                "job_id": audit_row[0].value,
                "excel_row": row_number,
                "input_digest": digest,
                "pipeline_status": audit_row[5].value,
            }
        finally:
            book.close()

    def row_number(self, job_id, identity):
        """Return the validated teacher-sheet row for render audit metadata."""
        if not self.path.exists():
            raise ValidationError("Results workbook does not exist")
        book = self._open()
        try:
            row, _ = self._row(book, job_id)
            if row is None:
                raise ValidationError("Student result not found")
            self._decode(row, identity)
            return row[0].row
        finally:
            book.close()

    def set_status(self, job_id, identity, pipeline_status, review_status=None):
        with self._write_lock:
            return self._set_status_unlocked(job_id, identity, pipeline_status, review_status)

    def _set_status_unlocked(self, job_id, identity, pipeline_status, review_status=None):
        book = self._open()
        try:
            row, audit_row = self._row(book, job_id)
            if row is None:
                raise ValidationError("Student result not found")
            self._decode(row, identity)
            if review_status is not None:
                if review_status not in ("PENDING", "APPROVED"):
                    raise ValidationError("Unknown review status")
                row[24].value = review_status
            audit_row[5].value = pipeline_status
            self._save(book)
        finally:
            book.close()

    def sort_rows(self):
        """Order the teacher sheet by 班级 then 班号, lowest first (1, 2, 3 …).

        Values, cell types, styles, comments and row heights move with each
        student; a 总分 =E+F sum is rewritten for its new row. Excel行 hints
        in the audit sheet are refreshed. Nothing is saved if already in order.
        """
        with self._write_lock:
            if not self.path.exists():
                return
            book = self._open()
            try:
                if self._sort_rows_unlocked(book):
                    self._save(book)
            finally:
                book.close()

    @staticmethod
    def _sort_rows_unlocked(book):
        sheet = book[SHEET]
        width = len(HEADERS)
        rows = []
        for row in sheet.iter_rows(min_row=2, max_col=width):
            if not any(cell.value is not None for cell in row):
                continue
            rows.append({
                "row": row[0].row,
                "key": tuple(roster_text(row[i].value) for i in (0, 1)),
                "height": sheet.row_dimensions[row[0].row].height,
                "cells": [(cell.value, cell.data_type, copy(cell._style), copy(cell.comment)) for cell in row],
            })
        ordered = sorted(rows, key=lambda item: tuple(_sort_key(value) for value in item["key"]))
        targets = [item["row"] for item in rows]
        if [item["row"] for item in ordered] == targets:
            return False
        targets.sort()
        new_rows = {}
        for target, item in zip(targets, ordered):
            new_rows[item["key"]] = target
            sheet.row_dimensions[target].height = item["height"]
            for col, (value, data_type, style, comment) in enumerate(item["cells"], 1):
                cell = sheet.cell(target, col)
                if (col == TOTAL_COLUMN and data_type == "f"
                        and str(value).replace(" ", "").upper() == _total_formula(item["row"])):
                    value = _total_formula(target)
                cell.value = value
                if value is not None:
                    cell.data_type = data_type
                cell._style = style
                cell.comment = comment
        for audit in book[AUDIT_SHEET].iter_rows(min_row=2):
            key = (roster_text(audit[1].value), roster_text(audit[2].value))
            if key in new_rows:
                audit[3].value = new_rows[key]
        return True


def _sort_key(text):
    # Numbers compare as numbers, so 班号 9 comes before 10; any non-numeric
    # value comes after every number.
    return (0, int(text), "") if text.isdigit() else (1, 0, text)
