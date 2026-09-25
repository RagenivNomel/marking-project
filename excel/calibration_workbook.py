"""Separate batch-shaped workbook repository for blind calibration output."""
from pathlib import Path
import json
import os
import tempfile

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter

from grading.calibration_schema import CalibrationResponse
from grading.schemas import CRITERIA, RATINGS, ValidationError
from .schema import HEADERS

BLIND_SHEET = "批改结果"
META_SHEET = "校准审计"
REVIEW_STATUS = "PENDING_HUMAN_REVIEW"


def _style_header(row):
    for cell in row:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="285E61")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


class CalibrationWorkbookRepository:
    def __init__(self, path):
        self.path = Path(path)

    def write_initial(self, result: CalibrationResponse, metadata: dict, topic: str):
        if result.grading is None:
            raise ValidationError("Unreadable result must not be written as a grading row")
        if self.path.exists():
            raise ValidationError("Calibration workbook already exists; initial blind output is immutable")
        self.path.parent.mkdir(parents=True, exist_ok=True)
        book = Workbook()
        sheet = book.active
        sheet.title = BLIND_SHEET
        sheet.append(HEADERS)
        grade = result.grading
        values = [grade.class_name, grade.student_id, grade.student_name, topic, None, None, None]
        for name in CRITERIA:
            values.extend([grade.criteria[name].rating, grade.criteria[name].short_comment])
        values.extend([grade.teacher_comment, REVIEW_STATUS])
        sheet.append(values)
        _style_header(sheet[1])
        sheet.freeze_panes = "H2"
        sheet.auto_filter.ref = sheet.dimensions
        sheet.row_dimensions[1].height = 42
        sheet.row_dimensions[2].height = 120
        for index, header in enumerate(HEADERS, 1):
            sheet.column_dimensions[get_column_letter(index)].width = 34 if header.endswith("评语") or header in ("题目", "教师总评") else 16
        for cell in sheet[2]:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            if isinstance(cell.value, str):
                cell.data_type = "s"
        for index in range(8, 24, 2):
            validation = DataValidation(type="list", formula1='"' + ','.join(RATINGS) + '"')
            validation.showErrorMessage = True
            validation.errorStyle = "stop"
            validation.errorTitle = "等级无效"
            validation.error = "请选择指定的三个等级之一。"
            sheet.add_data_validation(validation)
            validation.add(f"{get_column_letter(index)}2:{get_column_letter(index)}1048576")

        audit = book.create_sheet(META_SHEET)
        audit.append(["字段", "值"])
        for key, value in metadata.items():
            audit.append([key, str(value) if value is not None else ""])
        audit.append(["审核状态", REVIEW_STATUS])
        audit.append(["阅读疑点", "\n".join(f"第{item.page_number}页：{item.description}" for item in result.reading_quality.uncertainties)])
        audit.append(["判断依据", "\n".join(f"{item.judgment}｜第{','.join(map(str, item.page_numbers))}页｜{item.rationale}" for item in result.evidence)])
        audit.append(["结构化审计JSON", json.dumps({"reading_quality": result.to_dict()["reading_quality"], "evidence": result.to_dict()["evidence"]}, ensure_ascii=False, sort_keys=True)])
        _style_header(audit[1])
        audit.column_dimensions["A"].width = 30
        audit.column_dimensions["B"].width = 100
        for row in audit.iter_rows(min_row=2):
            row[1].alignment = Alignment(wrap_text=True, vertical="top")

        fd, temporary = tempfile.mkstemp(dir=self.path.parent, suffix=".xlsx")
        os.close(fd)
        try:
            book.save(temporary)
            os.replace(temporary, self.path)
        finally:
            book.close()
            Path(temporary).unlink(missing_ok=True)

    def verify_status(self):
        book = load_workbook(self.path, read_only=True, data_only=True)
        try:
            headers = [cell.value for cell in book[BLIND_SHEET][1]]
            if headers != HEADERS:
                raise ValidationError("Calibration workbook schema changed unexpectedly")
            values = dict(zip(headers, [cell.value for cell in book[BLIND_SHEET][2]]))
            if values.get("审核状态") != REVIEW_STATUS:
                raise ValidationError("Calibration workbook review status changed unexpectedly")
            return values
        finally:
            book.close()
